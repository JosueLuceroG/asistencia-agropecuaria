"""
Excel report exporter - extracted from main v4.py
Handles generating .xlsx attendance reports with full formatting.
Optimized with cached styles to avoid recreating objects per cell.
"""
import os
from datetime import datetime, date, timedelta
from config_manager import (
    PASS_EXCEL, cargar_configuracion, obtener_horarios_config,
    obtener_horario_empleado, obtener_horario_especial_empleado,
    obtener_mapa_dias_no_laborados, es_dia_no_laborado, es_dia_laboral,
    limpiar_nombre, normalizar_horario, formato_hora,
    formato_entrada_reporte, formato_salida_reporte, dias_sem, meses,
)
from attendance_calculator import (
    construir_reporte_asistencia, resumen_horarios_especiales,
    calcular_filas_dashboard_mensual, datos_memoria_crudos,
    checador_memoria, ultimo_reporte, titulo_actual,
)

# ─── LAZY OPENPYXL IMPORTS ─────────────────────────────────────
Workbook = Font = Alignment = PatternFill = Border = Side = get_column_letter = None

def _cargar_openpyxl():
    global Workbook, Font, Alignment, PatternFill, Border, Side, get_column_letter
    if Workbook is None:
        from openpyxl import Workbook as _Workbook
        from openpyxl.styles import Font as _Font, Alignment as _Alignment, PatternFill as _PatternFill, Border as _Border, Side as _Side
        from openpyxl.utils import get_column_letter as _get_column_letter
        Workbook = _Workbook
        Font = _Font
        Alignment = _Alignment
        PatternFill = _PatternFill
        Border = _Border
        Side = _Side
        get_column_letter = _get_column_letter

# ─── CACHED STYLES ────────────────────────────────────────────
# Pre-create all style objects once for reuse across cells
_FILLS = {}
_BORDERS = {}
_FONTS = {}
_ALIGNMENTS = {}


def _get_fill(color):
    """Get or create a cached PatternFill."""
    if color not in _FILLS:
        _FILLS[color] = PatternFill("solid", fgColor=color)
    return _FILLS[color]


def _get_font(bold=False, color=None, size=None, italic=False):
    """Get or create a cached Font."""
    key = (bold, color, size, italic)
    if key not in _FONTS:
        _FONTS[key] = Font(bold=bold, color=color, size=size, italic=italic)
    return _FONTS[key]


def _get_alignment(horizontal=None, vertical=None, wrap_text=False):
    """Get or create a cached Alignment."""
    key = (horizontal, vertical, wrap_text)
    if key not in _ALIGNMENTS:
        _ALIGNMENTS[key] = Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap_text)
    return _ALIGNMENTS[key]


# ─── EXCEL CONSTANTS ───────────────────────────────────────────
EXCEL_SABADO_EXTRA_FILL = "CCFBF1"
EXCEL_SABADO_EXTRA_FONT = "0F766E"

# Pre-bake thin_border for performance
_THIN_BORDER_CACHE = None


def _get_thin_border():
    global _THIN_BORDER_CACHE
    if _THIN_BORDER_CACHE is None:
        side = Side(style='thin', color="D9D9D9")
        _THIN_BORDER_CACHE = Border(left=side, right=side, top=side, bottom=side)
    return _THIN_BORDER_CACHE


def _stylize(ws, row, col, value, fill=None, font=None, alignment=None, border=None, number_format=None):
    """Fast cell write with pre-cached styles."""
    c = ws.cell(row=row, column=col, value=value)
    if fill:
        c.fill = fill
    if font:
        c.font = font
    if alignment:
        c.alignment = alignment
    if border:
        c.border = border
    if number_format:
        c.number_format = number_format
    return c


def agregar_resumen_horarios_especiales_excel(ws, fila_inicio, start_col, total_cols, horarios, empleados_reporte, thin_border):
    filas = resumen_horarios_especiales(horarios, empleados_reporte)
    if not filas:
        return fila_inicio
    fila = fila_inicio
    end_col = start_col + max(total_cols, 5) - 1
    ws.merge_cells(start_row=fila, start_column=start_col, end_row=fila, end_column=end_col)
    titulo = ws.cell(row=fila, column=start_col, value="Horarios especiales aplicados")
    titulo.font = _get_font(bold=True, color="FFFFFF")
    titulo.fill = _get_fill("1D4ED8")
    titulo.alignment = _get_alignment(horizontal="center")
    fila += 1
    headers = ["Horario especial", "Entrada", "Salida", "Tolerancia", "Persona"]
    for j, h in enumerate(headers):
        c = ws.cell(row=fila, column=start_col + j, value=h)
        c.font = _get_font(bold=True)
        c.fill = _get_fill("DBEAFE")
        c.border = thin_border
        c.alignment = _get_alignment(horizontal="center")
    fila += 1
    for item in filas:
        for empleado in item["empleados"]:
            vals = [item["nombre"], item["entrada"], item["salida"], f"{item['tolerancia']} min", empleado]
            for j, v in enumerate(vals):
                c = ws.cell(row=fila, column=start_col + j, value=v)
                c.border = thin_border
                c.fill = _get_fill("EFF6FF")
                c.alignment = _get_alignment(horizontal="center" if j < 4 else "left", wrap_text=True)
            fila += 1
    return fila + 1


def agregar_leyenda_colores_excel(ws, fila_inicio, start_col, thin_border, incluir_especial=True, incluir_no_laboral=True, incluir_no_laborado=True, incluir_sabado_extra=True):
    fila = fila_inicio
    ws.cell(row=fila, column=start_col, value="Leyenda").font = _get_font(bold=True, color="FFFFFF")
    ws.cell(row=fila, column=start_col).fill = _get_fill("334155")
    ws.cell(row=fila, column=start_col).alignment = _get_alignment(horizontal="center")
    ws.cell(row=fila, column=start_col).border = thin_border
    fila += 1
    items = [
        ("RETARDO", "FEF3C7", "Color amarillo: retardo"),
        ("FALTA / SIN SALIDA", "FFEBEE", "Color rojo: falta o falta de salida"),
    ]
    if incluir_sabado_extra:
        items.append(("SABADO EXTRA", EXCEL_SABADO_EXTRA_FILL, "Sabado no laboral: suma asistencia extra"))
    if incluir_no_laborado:
        items.append(("NO LABORADO", "E5E7EB", "Fecha capturada como no laborada: no castiga"))
    if incluir_no_laboral:
        items.append(("NO LABORAL", "E5E7EB", "Domingo/no laboral: solo informativo, no castiga"))
    if incluir_especial:
        items.append(("HORARIO ESPECIAL", "E0F2FE", "Color azul: persona con horario especial"))
    for etiqueta, color, descripcion in items:
        c1 = ws.cell(row=fila, column=start_col, value=etiqueta)
        c2 = ws.cell(row=fila, column=start_col + 1, value=descripcion)
        c1.fill = _get_fill(color)
        c2.fill = _get_fill("FFFFFF")
        c1.font = _get_font(bold=True, color="0F172A")
        c1.alignment = _get_alignment(horizontal="center")
        c2.alignment = _get_alignment(horizontal="left")
        c1.border = thin_border
        c2.border = thin_border
        fila += 1
    return fila + 1


def crear_dashboard_mensual(wb, reporte, fechas_periodo, horarios, f_inicio, f_fin, thin_border, col_emp):
    ws_dash = wb.create_sheet("Dashboard Mensual", 0)
    ws_dash.sheet_view.showGridLines = False
    ws_dash.protection.set_password(PASS_EXCEL)
    ws_dash.protection.sheet = True
    filas = calcular_filas_dashboard_mensual(reporte, fechas_periodo, horarios, col_emp)
    dias_no_laborados = obtener_mapa_dias_no_laborados()
    dias_laborables_mes = {d for d in fechas_periodo if es_dia_laboral(d, horarios, dias_no_laborados)}
    dias_no_laborados_mes = {d for d in fechas_periodo if d in dias_no_laborados}
    total_programados = sum(f["programados"] for f in filas)
    total_trabajados = sum(f["trabajados"] for f in filas)
    total_sabados_extra = sum(f["sabados_extra"] for f in filas)
    total_domingos_info = sum(f["domingos_info"] for f in filas)
    total_faltas = sum(f["faltas"] for f in filas)
    total_retardos = sum(f["retardos"] for f in filas)
    total_sin_salida = sum(f["sin_salida"] for f in filas)
    asistencia_pct = (sum(f["asistencia_pct"] for f in filas) / len(filas)) if filas else 0
    ws_dash.merge_cells("A1:K2")
    titulo = ws_dash["A1"]
    titulo.value = "DASHBOARD MENSUAL DE ASISTENCIA"
    titulo.font = _get_font(size=20, bold=True, color="FFFFFF")
    titulo.alignment = _get_alignment(horizontal="center", vertical="center")
    titulo.fill = _get_fill("1F4E78")
    ws_dash.merge_cells("A3:K3")
    subtitulo = ws_dash["A3"]
    subtitulo.value = f"Periodo: {f_inicio.strftime('%d/%m/%Y')} al {f_fin.strftime('%d/%m/%Y')} | Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    subtitulo.font = _get_font(italic=True, color="475569")
    subtitulo.alignment = _get_alignment(horizontal="center")
    cards = [
        ("Personas", len(filas), "DBEAFE"),
        ("Asistencia", f"{asistencia_pct:.0%}", "DCFCE7" if asistencia_pct >= 0.9 else "FEF3C7"),
        ("Dias laborables", len(dias_laborables_mes), "E0F2FE"),
        ("No laborados", len(dias_no_laborados_mes), "E5E7EB" if dias_no_laborados_mes else "F1F5F9"),
        ("Faltas", total_faltas, "FFEBEE" if total_faltas else "F1F5F9"),
        ("Retardos", total_retardos, "FEF3C7" if total_retardos else "F1F5F9"),
    ]
    for idx, (label, value, color) in enumerate(cards):
        row = 5 if idx < 3 else 8
        col = 1 + (idx % 3) * 3
        ws_dash.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 2)
        ws_dash.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 2)
        c_label = ws_dash.cell(row=row, column=col, value=label)
        c_value = ws_dash.cell(row=row + 1, column=col, value=value)
        for c in [c_label, c_value]:
            c.fill = _get_fill(color)
            c.alignment = _get_alignment(horizontal="center", vertical="center")
        c_label.font = _get_font(bold=True, color="334155")
        c_value.font = _get_font(size=18, bold=True, color="0F172A")
    fila = 12
    headers = ["Empleado", "Horario", "Dias laborables", "Dias trabajados", "Sabados extra", "% asistencia", "Retardos", "Faltas", "Sin salida", "Salidas temp.", "Domingos info"]
    for j, h in enumerate(headers, start=1):
        c = ws_dash.cell(row=fila, column=j, value=h)
        c.font = _get_font(bold=True, color="FFFFFF")
        c.fill = _get_fill("366092")
        c.border = thin_border
        c.alignment = _get_alignment(horizontal="center", wrap_text=True)
    fila += 1
    for item in filas:
        vals = [item["empleado"], item["horario"], item["programados"], item["trabajados"],
                item["sabados_extra"], item["asistencia_pct"], item["retardos"], item["faltas"],
                item["sin_salida"], item["salidas_temp"], item["domingos_info"]]
        bg = "EFF6FF" if item["horario"] != "General" else ("FFFFFF" if fila % 2 == 0 else "F8FAFC")
        for j, v in enumerate(vals, start=1):
            c = ws_dash.cell(row=fila, column=j, value=v)
            c.border = thin_border
            c.fill = _get_fill(bg)
            c.alignment = _get_alignment(horizontal="center" if j != 1 else "left", wrap_text=True)
            if j == 6:
                c.number_format = "0%"
                if item["asistencia_pct"] > 1:
                    c.fill = _get_fill("BBF7D0")
                    c.font = _get_font(color="166534", bold=True)
                elif item["asistencia_pct"] < 0.8:
                    c.fill = _get_fill("FFEBEE")
                    c.font = _get_font(color="C62828", bold=True)
                elif item["asistencia_pct"] < 0.9:
                    c.fill = _get_fill("FEF3C7")
                    c.font = _get_font(color="B45309", bold=True)
            if j == 5 and int(v or 0) > 0:
                c.font = _get_font(color="166534", bold=True)
            if j == 11 and int(v or 0) > 0:
                c.font = _get_font(color="475569", bold=True)
            if j in [7, 8, 9, 10] and int(v or 0) > 0:
                c.font = _get_font(color="B45309" if j == 7 else "C62828", bold=True)
        fila += 1
    agregar_resumen_horarios_especiales_excel(ws_dash, fila + 2, 1, 11, horarios, reporte[col_emp].unique(), thin_border)
    widths = [32, 18, 15, 15, 13, 13, 11, 11, 12, 14, 13]
    for idx, width in enumerate(widths, start=1):
        ws_dash.column_dimensions[get_column_letter(idx)].width = width
    ws_dash.freeze_panes = "A13"
    ws_dash.page_setup.orientation = ws_dash.ORIENTATION_LANDSCAPE
    ws_dash.page_setup.paperSize = ws_dash.PAPERSIZE_LETTER
    ws_dash.page_setup.fitToWidth = 1
    ws_dash.page_setup.fitToHeight = 0
    return ws_dash


def exportar_reporte(datos_memoria_crudos, f_inicio, f_fin, checador_info, checador_nombre, modo, carpeta):
    """Main entry point for exporting reports. Called from main.py"""
    import pandas as pd
    _cargar_openpyxl()

    horarios = obtener_horarios_config(checador_nombre)
    empleados_asignados = checador_info.get("empleados", []) if checador_info else []
    usar_primera_ultima = bool(checador_info and str(checador_info.get("marca", "")).strip().lower() == "hikvision")

    reporte, base_asistencia, fechas_periodo = construir_reporte_asistencia(
        datos_memoria_crudos, f_inicio, f_fin, horarios,
        incluir_faltas=True, empleados_asignados=empleados_asignados,
        usar_primera_ultima=usar_primera_ultima
    )

    if reporte.empty:
        return None

    col_emp = 'nombre'

    if modo == "Diario":
        nombre_fin = f"ASISTENCIA_{f_inicio.strftime('%d_%m_%Y')}.xlsx"
        titulo_actual = f"ASISTENCIA DEL DÍA: {f_inicio.strftime('%d/%m/%Y')}"
    else:
        nombre_fin = f"REPORTE_{modo.upper()}_{f_inicio.day}_AL_{f_fin.day}.xlsx"
        titulo_actual = f"REPORTE {modo.upper()}: {f_inicio.day} AL {f_fin.day} DE {meses[f_fin.month].upper()}"

    salida = os.path.join(carpeta, nombre_fin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Detalle Mensual" if modo == "Mensual" else "Reporte"
    ws.protection.set_password(PASS_EXCEL)
    ws.protection.sheet = True

    thin_border = _get_thin_border()
    unique_dates = fechas_periodo if modo in ["Semanal", "Mensual"] else [f_inicio]

    if modo == "Diario": total_cols = 7
    elif modo == "Semanal": total_cols = max(1 + (len(unique_dates) * 2), 5)
    else: total_cols = 10

    start_col = 2

    f_gen = ws.cell(row=1, column=start_col + total_cols - 1,
                    value=f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    f_gen.font = _get_font(italic=True, size=9, color="808080")
    f_gen.alignment = _get_alignment(horizontal="right")

    ws.merge_cells(start_row=2, start_column=start_col, end_row=4, end_column=start_col + total_cols - 1)
    t = ws.cell(row=2, column=start_col, value="ASISTENCIA")
    t.font = _get_font(size=26, bold=True, color="1F4E78")
    t.alignment = _get_alignment(horizontal="center", vertical="center")
    t.fill = _get_fill("DCE6F1")

    ws.merge_cells(start_row=5, start_column=start_col, end_row=5, end_column=start_col + total_cols - 1)
    s = ws.cell(row=5, column=start_col, value=titulo_actual)
    s.font = _get_font(bold=True, color="FFFFFF")
    s.alignment = _get_alignment(horizontal="center")
    s.fill = _get_fill("366092")

    current_row = 7

    if modo == "Diario":
        headers = ["Nombre Completo", "Horario", "Fecha", "Entrada", "Salida", "Retardo", "Tiempo Retardo"]
        for j, h in enumerate(headers):
            c = ws.cell(row=6, column=start_col + j, value=h)
            c.font = _get_font(bold=True)
            c.fill = _get_fill("F2F2F2")
            c.border = thin_border
            c.alignment = _get_alignment(horizontal="center")
        data_d = reporte[reporte["Fecha_D"] == f_inicio].sort_values(by=col_emp)
        for i, (_, row) in enumerate(data_d.iterrows()):
            current_row = 7 + i
            especial = obtener_horario_especial_empleado(row[col_emp], horarios)
            horario_txt = especial.get("nombre", "General") if especial else "General"
            bg = "E0F2FE" if especial else ("F9F9F9" if i % 2 != 0 else "FFFFFF")
            if row["Falta"]:
                vals = [row[col_emp], horario_txt, row["Fecha_D"].strftime("%d/%m/%Y"), "FALTA", "FALTA", "FALTA", "FALTA"]
            else:
                if bool(row.get("Dia_No_Laboral", False)):
                    estado_no_laboral = "NO"
                    tiempo_no_laboral = "0 min"
                    salida_txt = formato_salida_reporte(row["Entrada"], row["Salida"])
                    if salida_txt == "SIN SALIDA":
                        salida_txt = "NO APLICA"
                    vals = [row[col_emp], horario_txt, row["Fecha_D"].strftime("%d/%m/%Y"),
                           formato_entrada_reporte(row["Entrada"]), salida_txt, estado_no_laboral, tiempo_no_laboral]
                else:
                    salida_txt = formato_salida_reporte(row["Entrada"], row["Salida"])
                    vals = [row[col_emp], horario_txt, row["Fecha_D"].strftime("%d/%m/%Y"),
                           formato_entrada_reporte(row["Entrada"]), salida_txt, row["Retardo_SN"], row["Tiempo_Retardo"]]
            for j, v in enumerate(vals):
                c = ws.cell(row=current_row, column=start_col + j, value=v)
                c.border = thin_border
                c.alignment = _get_alignment(horizontal="center")
                c.fill = _get_fill(bg)
                if j == 1 and especial:
                    c.font = _get_font(color="1D4ED8", bold=True)
                if v == "SIN ENTRADA":
                    c.font = _get_font(color="C62828", bold=True)
                    c.fill = _get_fill("FFEBEE")
                if v == "SIN SALIDA":
                    c.font = _get_font(color="C62828", bold=True)
                    c.fill = _get_fill("FFEBEE")
                if v == "FALTA":
                    c.font = _get_font(color="C62828", bold=True)
                    c.fill = _get_fill("FFEBEE")
                if v == "SI":
                    c.font = _get_font(color="B45309", bold=True)
                    c.fill = _get_fill("FEF3C7")
                if v == "SABADO EXTRA":
                    c.font = _get_font(color=EXCEL_SABADO_EXTRA_FONT, bold=True)
                    c.fill = _get_fill(EXCEL_SABADO_EXTRA_FILL)
                if v == "NO LABORADO":
                    c.font = _get_font(color="475569", bold=True)
                    c.fill = _get_fill("E5E7EB")
        current_row += 1
        current_row = agregar_leyenda_colores_excel(ws, current_row + 1, start_col, thin_border,
                                                     incluir_especial=True, incluir_no_laboral=False,
                                                     incluir_no_laborado=False, incluir_sabado_extra=False)
        current_row = agregar_resumen_horarios_especiales_excel(ws, current_row, start_col, total_cols,
                                                                 horarios, data_d[col_emp].unique(), thin_border)

    elif modo == "Semanal":
        ws.merge_cells(start_row=6, start_column=start_col, end_row=7, end_column=start_col)
        ws.cell(row=6, column=start_col, value="Nombre Completo")
        c_idx = start_col + 1
        for d in unique_dates:
            ws.merge_cells(start_row=6, start_column=c_idx, end_row=6, end_column=c_idx + 1)
            ws.cell(row=6, column=c_idx, value=f"{dias_sem[d.weekday()]} {d.day}")
            ws.cell(row=7, column=c_idx, value="Entrada")
            ws.cell(row=7, column=c_idx + 1, value="Salida")
            c_idx += 2
        for r in [6, 7]:
            for c in range(start_col, start_col + total_cols):
                cell = ws.cell(row=r, column=c)
                cell.font = _get_font(bold=True)
                cell.fill = _get_fill("F2F2F2")
                cell.border = thin_border
                cell.alignment = _get_alignment(horizontal="center")
        empleados = sorted(reporte[col_emp].unique())
        current_row = 8
        for i, emp in enumerate(empleados):
            especial = obtener_horario_especial_empleado(emp, horarios)
            bg = "E0F2FE" if especial else ("F9F9F9" if i % 2 != 0 else "FFFFFF")
            nombre_txt = emp
            c_nombre = ws.cell(row=current_row, column=start_col, value=nombre_txt)
            c_nombre.fill = _get_fill(bg)
            c_nombre.border = thin_border
            c_nombre.alignment = _get_alignment(horizontal="left", vertical="center", wrap_text=True)
            if especial:
                c_nombre.font = _get_font(color="1D4ED8", bold=True)
            c_idx = start_col + 1
            for d in unique_dates:
                res = reporte[(reporte[col_emp] == emp) & (reporte["Fecha_D"] == d)]
                if res.empty:
                    ev, sv = ("", "")
                elif res.iloc[0]["Falta"]:
                    ev, sv = ("FALTA", "FALTA")
                elif bool(res.iloc[0].get("Dia_No_Laboral", False)):
                    if bool(res.iloc[0].get("No_Laborado", False)):
                        marca = "NO LABORADO"
                    elif bool(res.iloc[0].get("Extra_Sabado", False)):
                        marca = "SABADO EXTRA"
                    else:
                        marca = ""
                    entrada_txt = formato_entrada_reporte(res.iloc[0]["Entrada"])
                    ev = f"{entrada_txt}\n{marca}" if marca else entrada_txt
                    salida_txt = formato_salida_reporte(res.iloc[0]["Entrada"], res.iloc[0]["Salida"])
                    sv = "" if salida_txt == "SIN SALIDA" else salida_txt
                else:
                    ev = formato_entrada_reporte(res.iloc[0]["Entrada"])
                    sv = formato_salida_reporte(res.iloc[0]["Entrada"], res.iloc[0]["Salida"])
                    if str(res.iloc[0]["Retardo"]).startswith("Si"):
                        ev = f"{ev}\nRETARDO"
                    if str(res.iloc[0]["Salida_Temprana"]).startswith("Si"):
                        sv = f"{sv}\nSALIDA TEMP."
                for j, v in enumerate([ev, sv]):
                    c = ws.cell(row=current_row, column=c_idx + j, value=v)
                    c.fill = _get_fill(bg)
                    c.border = thin_border
                    c.alignment = _get_alignment(horizontal="center", wrap_text=True)
                    if v == "SIN ENTRADA" or (isinstance(v, str) and "SIN ENTRADA" in v):
                        c.font = _get_font(color="C62828", bold=True)
                        c.fill = _get_fill("FFEBEE")
                    if v == "SIN SALIDA":
                        c.font = _get_font(color="C62828", bold=True)
                        c.fill = _get_fill("FFEBEE")
                    if v == "FALTA":
                        c.font = _get_font(color="C62828", bold=True)
                        c.fill = _get_fill("FFEBEE")
                    if isinstance(v, str) and "RETARDO" in v:
                        c.font = _get_font(color="B45309", bold=True)
                        c.fill = _get_fill("FEF3C7")
                    if isinstance(v, str) and "SALIDA TEMP." in v:
                        c.font = _get_font(color="C2410C", bold=True)
                        c.fill = _get_fill("FFEDD5")
                    if isinstance(v, str) and "SABADO EXTRA" in v:
                        c.font = _get_font(color=EXCEL_SABADO_EXTRA_FONT, bold=True)
                        c.fill = _get_fill(EXCEL_SABADO_EXTRA_FILL)
                    if isinstance(v, str) and "NO LABORADO" in v:
                        c.font = _get_font(color="475569", bold=True)
                        c.fill = _get_fill("E5E7EB")
                c_idx += 2
            current_row += 1
        current_row = agregar_leyenda_colores_excel(ws, current_row + 1, start_col, thin_border,
                                                     incluir_especial=True, incluir_no_laboral=False)
        current_row = agregar_resumen_horarios_especiales_excel(ws, current_row, start_col, total_cols,
                                                                 horarios, empleados, thin_border)

    elif modo == "Mensual":
        headers = ["Nombre Completo", "Horario", "Semana", "Lunes", "Martes", "Miercoles",
                   "Jueves", "Viernes", "Sabado", "Domingo"]
        for j, h in enumerate(headers):
            c = ws.cell(row=6, column=start_col + j, value=h)
            c.font = _get_font(bold=True, color="FFFFFF")
            c.fill = _get_fill("366092")
            c.border = thin_border
            c.alignment = _get_alignment(horizontal="center")
        ws.row_dimensions[6].height = 26
        dias_no_laborados = obtener_mapa_dias_no_laborados()
        semanas = []
        if unique_dates:
            curr_week = []
            curr_monday = unique_dates[0] - timedelta(days=unique_dates[0].weekday())
            for d in unique_dates:
                monday = d - timedelta(days=d.weekday())
                if monday == curr_monday:
                    curr_week.append(d)
                else:
                    semanas.append((curr_monday, curr_week))
                    curr_week = [d]
                    curr_monday = monday
            if curr_week:
                semanas.append((curr_monday, curr_week))
        empleados = sorted(reporte[col_emp].unique())
        current_row = 7
        for emp in empleados:
            especial = obtener_horario_especial_empleado(emp, horarios)
            horario_txt = especial.get("nombre", "General") if especial else "General"
            emp_data = reporte[reporte[col_emp] == emp].sort_values(by="Fecha_D")
            emp_por_fecha = {row["Fecha_D"]: row for _, row in emp_data.iterrows()}
            start_row_emp = current_row
            primera_semana_emp = True
            for num_sem, (monday, dates_in_week) in enumerate(semanas, start=1):
                if not any(d in emp_por_fecha or d in dias_no_laborados for d in dates_in_week):
                    continue
                start_str = dates_in_week[0].strftime("%d/%m")
                end_str = dates_in_week[-1].strftime("%d/%m")
                vals_base = [emp if primera_semana_emp else "", horario_txt if primera_semana_emp else "",
                            f"Sem {num_sem}\n{start_str}-{end_str}"]
                bg_base = "E0F2FE" if especial else ("FFFFFF" if current_row % 2 == 0 else "F8FAFC")
                for j, v in enumerate(vals_base):
                    c = ws.cell(row=current_row, column=start_col + j, value=v)
                    c.border = thin_border
                    c.fill = _get_fill(bg_base)
                    c.alignment = _get_alignment(horizontal="left" if j == 0 else "center", vertical="center", wrap_text=True)
                    if j == 1 and especial:
                        c.font = _get_font(color="1D4ED8", bold=True)
                for dia_idx in range(7):
                    d = monday + timedelta(days=dia_idx)
                    c = ws.cell(row=current_row, column=start_col + 3 + dia_idx)
                    c.border = thin_border
                    c.alignment = _get_alignment(horizontal="center", vertical="center", wrap_text=True)
                    c.fill = _get_fill(bg_base)
                    if d not in unique_dates:
                        c.value = ""
                        continue
                    row = emp_por_fecha.get(d)
                    if row is None:
                        if d in dias_no_laborados:
                            motivo = dias_no_laborados.get(d, "")
                            c.value = f"{d.day}\nNO LABORADO" + (f"\n{motivo}" if motivo else "")
                            c.font = _get_font(color="475569", bold=True)
                            c.fill = _get_fill("E5E7EB")
                        else:
                            c.value = ""
                        continue
                    if row.get("Falta", False):
                        c.value = f"{d.day}\nFALTA"
                        c.font = _get_font(color="C62828", bold=True)
                        c.fill = _get_fill("FFEBEE")
                    elif bool(row.get("Dia_No_Laboral", False)):
                        entrada_txt = formato_hora(row["Entrada"])
                        salida_txt = "" if row["Entrada"] == row["Salida"] else formato_hora(row["Salida"])
                        horas_txt = entrada_txt if not salida_txt else f"{entrada_txt} - {salida_txt}"
                        if bool(row.get("Extra_Sabado", False)):
                            c.value = f"{d.day}\nSABADO EXTRA\n{horas_txt}"
                            c.font = _get_font(color=EXCEL_SABADO_EXTRA_FONT, bold=True)
                            c.fill = _get_fill(EXCEL_SABADO_EXTRA_FILL)
                        elif bool(row.get("No_Laborado", False)):
                            motivo = row.get("Motivo_No_Laborado", "")
                            detalle = motivo if motivo else horas_txt
                            c.value = f"{d.day}\nNO LABORADO\n{detalle}"
                            c.font = _get_font(color="475569", bold=True)
                            c.fill = _get_fill("E5E7EB")
                        else:
                            c.value = f"{d.day}\nNO LABORAL\n{horas_txt}"
                            c.font = _get_font(color="475569", bold=True)
                            c.fill = _get_fill("E5E7EB")
                    else:
                        entrada_txt = formato_entrada_reporte(row["Entrada"])
                        salida_txt = formato_salida_reporte(row["Entrada"], row["Salida"])
                        marcas = []
                        if entrada_txt == "SIN ENTRADA":
                            marcas.append("SIN ENTRADA")
                        if salida_txt == "SIN SALIDA":
                            marcas.append("SIN SALIDA")
                        if bool(row.get("Tiene_Retardo", False)):
                            marcas.append("RETARDO")
                        if bool(row.get("Tiene_Salida_Temprana", False)):
                            marcas.append("SALIDA TEMP.")
                        estado_txt = " / ".join(marcas) if marcas else "ASISTENCIA"
                        c.value = f"{d.day}\n{estado_txt}\n{entrada_txt} - {salida_txt}"
                        if "SIN ENTRADA" in estado_txt or "SIN SALIDA" in estado_txt:
                            c.font = _get_font(color="C62828", bold=True)
                            c.fill = _get_fill("FFEBEE")
                        elif "RETARDO" in estado_txt:
                            c.font = _get_font(color="B45309", bold=True)
                            c.fill = _get_fill("FEF3C7")
                        elif "SALIDA TEMP." in estado_txt:
                            c.font = _get_font(color="C2410C", bold=True)
                            c.fill = _get_fill("FFEDD5")
                ws.row_dimensions[current_row].height = 55
                current_row += 1
                primera_semana_emp = False
            if current_row > start_row_emp:
                end_row_emp = current_row - 1
                if end_row_emp > start_row_emp:
                    ws.merge_cells(start_row=start_row_emp, start_column=start_col, end_row=end_row_emp, end_column=start_col)
                    ws.merge_cells(start_row=start_row_emp, start_column=start_col + 1, end_row=end_row_emp, end_column=start_col + 1)
                ws.cell(row=start_row_emp, column=start_col).alignment = _get_alignment(horizontal="left", vertical="center", wrap_text=True)
                ws.cell(row=start_row_emp, column=start_col + 1).alignment = _get_alignment(horizontal="center", vertical="center", wrap_text=True)
        current_row = agregar_leyenda_colores_excel(ws, current_row + 1, start_col, thin_border, incluir_especial=True)
        current_row = agregar_resumen_horarios_especiales_excel(ws, current_row, start_col, total_cols, horarios, empleados, thin_border)
        crear_dashboard_mensual(wb, reporte, unique_dates, horarios, f_inicio, f_fin, thin_border, col_emp)

    ws.freeze_panes = "A7"
    for col in range(start_col, start_col + total_cols):
        column_letter = get_column_letter(col)
        offset = col - start_col
        if modo == "Diario":
            anchos = [32, 18, 12, 13, 13, 12, 16]
            ws.column_dimensions[column_letter].width = anchos[offset] if offset < len(anchos) else 14
        elif modo == "Mensual":
            anchos = [30, 18, 12, 16, 16, 16, 16, 16, 16, 16]
            ws.column_dimensions[column_letter].width = anchos[offset] if offset < len(anchos) else 14
        elif col == start_col:
            ws.column_dimensions[column_letter].width = 32
        else:
            ws.column_dimensions[column_letter].width = 14

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_view.showGridLines = False
    if modo == "Mensual":
        wb.active = 0

    guardado_ok = False
    while not guardado_ok:
        try:
            wb.save(salida)
            guardado_ok = True
        except PermissionError:
            import tkinter.messagebox as mb
            if not mb.askretrycancel("Archivo Abierto", "Excel abierto. Ciérrelo y reintente."):
                return None

    return os.path.abspath(salida)

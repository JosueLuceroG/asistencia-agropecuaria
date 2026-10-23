import os
import sys
import json
import shutil
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, Toplevel
from tkinter import ttk 
from tkinter.ttk import Progressbar, Style
from datetime import datetime, timedelta, date
from zoneinfo import ZoneInfo

# Librerías para ZKTeco y Calendarios
from ui.stable_date_entry import StableDateEntry as DateEntry, standard_calendar_style
from services.attendance_service import append_attendance, record_sync_error, record_sync_success
from ui.backup_actions import create_backup_dialog, restore_backup_dialog
from ui.device_health_window import VentanaEstadoChecadores
from services import auth_service
from services.instance_lock import InstanceLock, InstanceLockedError
from ui.login_window import LoginWindow, PasswordChangeWindow, UserManagementWindow
from device_connectors import hikvision as hikvision_connector

# --- CONFIGURACIÓN DEL SISTEMA Y PERSISTENCIA ---
VERSION = "v 4.2"

def obtener_directorio_ejecucion():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

APP_DIR_LOCAL = os.path.join(os.environ.get("APPDATA", os.path.expanduser("~")), "AsistenciaAgropecuaria")
APP_DIR = os.environ.get("ASISTENCIA_CONFIG_DIR") or obtener_directorio_ejecucion()
PATH_CONFIG_ANTERIOR = os.path.join(obtener_directorio_ejecucion(), "config.json")
PATH_CONFIG_LOCAL = os.path.join(APP_DIR_LOCAL, "config.json")
PATH_CONFIG = os.path.join(APP_DIR, "config_global.json")
PASS_EXCEL = os.environ.get("ASISTENCIA_EXCEL_PASSWORD", "demo")
DEFAULT_PORT = 4370
HIKVISION_DEFAULT_PORT = 80
try:
    TZ_LOCAL = datetime.now().astimezone().tzinfo
except Exception:
    TZ_LOCAL = datetime.now().astimezone().tzinfo
pd = None
Workbook = Font = Alignment = PatternFill = Border = Side = get_column_letter = None
ZK = None

def marca_checador(checador):
    marca = str(checador.get("marca", checador.get("tipo", "zkteco"))).strip().lower()
    if marca in ["hikvision", "hik", "hik-connect"]:
        return "hikvision"
    return "zkteco"

def puerto_default_marca(marca):
    return HIKVISION_DEFAULT_PORT if marca == "hikvision" else DEFAULT_PORT

def cargar_pandas():
    global pd
    if pd is None:
        import pandas as _pd
        pd = _pd
    return pd

def cargar_openpyxl():
    global Workbook, Font, Alignment, PatternFill, Border, Side, get_column_letter
    if Workbook is None:
        from openpyxl import Workbook as _Workbook
        from openpyxl.styles import Font as _Font, Alignment as _Alignment, PatternFill as _PatternFill, Border as _Border, Side as _Side
        from openpyxl.utils import get_column_letter as _get_column_letter
        Workbook = _Workbook
        Font = _Font
        Alignment = _Alignment
        PatternFill = _PatternFill
        Border = _Border
        Side = _Side
        get_column_letter = _get_column_letter

def cargar_zk():
    global ZK
    if ZK is None:
        from zk import ZK as _ZK
        ZK = _ZK
    return ZK

def crear_zk(*args, **kwargs):
    return cargar_zk()(*args, **kwargs)

def conectar_zk_robusto(ip, port=DEFAULT_PORT, password=None, progreso=None, modo_preferido=None):
    intentos_udp = [
        ("UDP rapido", True, 4),
        ("TCP rapido", False, 4),
        ("UDP respaldo", True, 7),
        ("TCP respaldo", False, 7),
    ]
    intentos_tcp = [
        ("TCP rapido", False, 4),
        ("UDP rapido", True, 4),
        ("TCP respaldo", False, 7),
        ("UDP respaldo", True, 7),
    ]
    intentos = intentos_tcp if str(modo_preferido).lower() == "tcp" else intentos_udp
    errores = []
    puerto = puerto_seguro(port)
    for idx, (nombre_intento, usar_udp, timeout) in enumerate(intentos, start=1):
        conn = None
        try:
            if progreso:
                progreso(f"Conectando ({nombre_intento})...", 8 + idx * 6)
            kwargs = {
                "port": puerto,
                "timeout": timeout,
                "force_udp": usar_udp,
                "ommit_ping": True
            }
            if password is not None:
                kwargs["password"] = password
            zk = crear_zk(ip, **kwargs)
            conn = zk.connect()
            return conn, nombre_intento
        except Exception as e:
            errores.append(f"{nombre_intento}: {e}")
            if conn:
                try:
                    conn.disconnect()
                except Exception:
                    pass
    detalle = "\n".join(errores[-4:])
    raise Exception(f"No se pudo conectar al checador {ip}:{puerto} despues de varios intentos.\n{detalle}")

def guardar_preferencia_conexion(nombre_checador, modo_conexion):
    if not nombre_checador:
        return
    modo = "tcp" if "TCP" in str(modo_conexion).upper() else "udp"
    try:
        config = cargar_configuracion()
        actualizado = False
        for checador in config.get("checadores", []):
            if checador.get("nombre") == nombre_checador:
                if checador.get("modo_conexion") != modo:
                    checador["modo_conexion"] = modo
                    actualizado = True
                break
        if actualizado:
            guardar_configuracion(config)
    except Exception:
        pass

def hikvision_base_urls(info):
    ip = info.get("ip")
    puerto = puerto_seguro(info.get("port"), HIKVISION_DEFAULT_PORT)
    candidatos = []
    if puerto == 443:
        candidatos.append(f"https://{ip}:443")
    elif puerto in [80, 8000]:
        candidatos.append(f"http://{ip}:80")
    else:
        esquema = "https" if str(info.get("https", "")).lower() in ["1", "true", "si", "yes"] else "http"
        candidatos.append(f"{esquema}://{ip}:{puerto}")
    for url in [f"http://{ip}:80", f"https://{ip}:443"]:
        if url not in candidatos:
            candidatos.append(url)
    return candidatos

def hikvision_base_url(info):
    return hikvision_base_urls(info)[0]

def _dict_to_hikvision_xml(root_tag, d):
    """
    Convierte un dict en XML para Hikvision ISAPI.
    Ej: {'AcsEventCond': {'searchID': '123', 'major': 5}} ->
        <AcsEventCond><searchID>123</searchID><major>5</major></AcsEventCond>
    """
    def _xml_escape(s):
        s2 = str(s)
        # Manual XML escaping avoiding entity auto-conversion
        amp = chr(38)  # &
        lt = chr(60)   # <
        gt = chr(62)   # >
        quot = chr(34) # "
        apos = chr(39) # '
        s2 = s2.replace(amp, amp + "amp;")
        s2 = s2.replace(lt, amp + "lt;")
        s2 = s2.replace(gt, amp + "gt;")
        s2 = s2.replace(quot, amp + "quot;")
        s2 = s2.replace(apos, amp + "apos;")
        return s2
    parts = ['<?xml version="1.0" encoding="UTF-8"?>']
    xmlns = ' xmlns="http://www.isapi.org/ver20/XMLSchema"'
    def _build(elem_name, value, indent=0):
        prefix = "  " * indent
        if isinstance(value, dict):
            parts.append(prefix + "<" + elem_name + (xmlns if indent == 0 else "") + ">")
            for k, v in value.items():
                _build(k, v, indent + 1)
            parts.append(prefix + "</" + elem_name + ">")
        elif isinstance(value, list):
            for item in value:
                if isinstance(item, dict):
                    parts.append(prefix + "<" + elem_name + ">")
                    for k2, v2 in item.items():
                        _build(k2, v2, indent + 1)
                    parts.append(prefix + "</" + elem_name + ">")
                else:
                    parts.append(prefix + "<" + elem_name + ">" + _xml_escape(str(item)) + "</" + elem_name + ">")
        else:
            parts.append(prefix + "<" + elem_name + ">" + _xml_escape(str(value)) + "</" + elem_name + ">")
    _build(root_tag, d)
    return "\n".join(parts)


def hikvision_request_json(info, path, method="GET", payload=None, timeout=12):
    import urllib.request
    import urllib.error
    import ssl
    usuario = str(info.get("usuario", "admin")).strip() or "admin"
    password = str(info.get("password", ""))
    intentos_payload = [(None, None, {"Accept": "application/json, application/xml, */*"})]
    if payload is not None:
        json_data = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        intentos_payload = [("JSON", json_data, {"Content-Type": "application/json", "Accept": "application/json, application/xml, */*"})]
        root_tag = None
        if "AcsEventCond" in payload:
            root_tag = "AcsEventCond"
        elif "UserInfoSearchCond" in payload:
            root_tag = "UserInfoSearchCond"
        elif "UserInfo" in payload:
            root_tag = "UserInfo"
        if root_tag:
            xml_str = _dict_to_hikvision_xml(root_tag, payload[root_tag])
            intentos_payload.append(("XML", xml_str.encode("utf-8"), {"Content-Type": "application/xml", "Accept": "application/xml, application/json, */*"}))

    errores = []
    for base_url in hikvision_base_urls(info):
        passman = urllib.request.HTTPPasswordMgrWithDefaultRealm()
        passman.add_password(None, base_url, usuario, password)
        handlers = [
            urllib.request.HTTPDigestAuthHandler(passman),
            urllib.request.HTTPBasicAuthHandler(passman)
        ]
        if base_url.startswith("https://"):
            handlers.append(urllib.request.HTTPSHandler(context=ssl._create_unverified_context()))
        opener = urllib.request.build_opener(*handlers)
        for tipo_payload, data, headers in intentos_payload:
            url = f"{base_url}{path}"
            request = urllib.request.Request(url, data=data, headers=headers, method=method)
            try:
                with opener.open(request, timeout=timeout) as response:
                    raw = response.read().decode("utf-8", errors="replace").strip()
                if not raw:
                    return {}
                try:
                    return json.loads(raw)
                except Exception:
                    try:
                        return _hikvision_xml_a_dict(raw)
                    except Exception:
                        return {"raw": raw}
            except urllib.error.HTTPError as e:
                detalle = e.read().decode("utf-8", errors="replace")
                if e.code == 401:
                    raise Exception("Hikvision requiere usuario y contrasena correctos para ISAPI.")
                errores.append(f"{base_url} {tipo_payload or 'GET'}: HTTP {e.code} {detalle[:180]}")
            except Exception as e:
                errores.append(f"{base_url} {tipo_payload or 'GET'}: {e}")
                break
    raise Exception("No se pudo consultar ISAPI Hikvision.\n" + "\n".join(errores[-4:]))


def _hikvision_xml_a_dict(xml_str):
    """
    Convierte XML ISAPI de Hikvision a dict con estructura similar al JSON.
    Soporta namespaces y la estructura tipica de AcsEvent.
    Preserva el nombre del tag raiz para compatibilidad.
    """
    import xml.etree.ElementTree as ET
    root = ET.fromstring(xml_str)
    # Quitar namespace si existe (ej: {http://www.hikvision.com/ver10/XMLSchema})
    def _tag(elem):
        tag = elem.tag
        if '}' in tag:
            return tag.split('}', 1)[1]
        return tag

    root_tag = _tag(root)
    resultado = {}
    info_list_items = []

    for child in root:
        tag = _tag(child)
        text = (child.text or "").strip()
        sub_children = list(child)

        if tag == "InfoList" or tag == "infoList":
            # Procesar lista de eventos
            for evento in sub_children:
                ev_dict = {}
                for campo in evento:
                    ct = _tag(campo)
                    cv = (campo.text or "").strip()
                    ev_dict[ct] = cv
                info_list_items.append(ev_dict)
        elif sub_children:
            # Tiene hijos -> crear sub-dict
            sub_dict = {}
            for sub in sub_children:
                st = _tag(sub)
                sv = (sub.text or "").strip()
                if st in sub_dict:
                    if not isinstance(sub_dict[st], list):
                        sub_dict[st] = [sub_dict[st]]
                    sub_dict[st].append(sv)
                else:
                    sub_dict[st] = sv
            if tag in resultado:
                if not isinstance(resultado[tag], list):
                    resultado[tag] = [resultado[tag]]
                resultado[tag].append(sub_dict)
            else:
                resultado[tag] = sub_dict
        else:
            if tag in resultado:
                if not isinstance(resultado[tag], list):
                    resultado[tag] = [resultado[tag]]
                resultado[tag].append(text)
            else:
                resultado[tag] = text

    if info_list_items:
        resultado["InfoList"] = info_list_items

    # Preservar la estructura {RootTag: {...}} para compatibilidad con _hikvision_buscar_eventos
    if root_tag and root_tag not in resultado:
        return {root_tag: resultado}
    return resultado

def hikvision_search_id():
    return datetime.now().strftime("%Y%m%d%H%M%S%f")

def hikvision_obtener_hora_dispositivo(info):
    data = hikvision_request_json(info, "/ISAPI/System/time?format=json", timeout=8)
    root = data.get("Time", data)
    valor = root.get("localTime") or root.get("time") or root.get("dateTime")
    if not valor:
        return None
    return datetime.fromisoformat(str(valor).replace("Z", "+00:00"))

def hikvision_ajuste_tiempo(info):
    try:
        hora_dispositivo = hikvision_obtener_hora_dispositivo(info)
        if hora_dispositivo is None:
            return timedelta(0)
        if hora_dispositivo.tzinfo is not None:
            hora_dispositivo = hora_dispositivo.astimezone(TZ_LOCAL).replace(tzinfo=None)
        hora_dispositivo = hora_dispositivo.replace(microsecond=0)
        ahora_local = datetime.now(TZ_LOCAL).replace(tzinfo=None, microsecond=0)
        ajuste = ahora_local - hora_dispositivo
        if abs(ajuste.total_seconds()) > 300:
            return ajuste
    except Exception:
        pass
    return timedelta(0)

def hikvision_fecha_iso(fecha, final=False, ajuste_tiempo=None):
    hora = datetime.max.time().replace(microsecond=0) if final else datetime.min.time()
    dt = datetime.combine(fecha, hora)
    if ajuste_tiempo:
        dt = dt - ajuste_tiempo
    dt = dt.replace(tzinfo=TZ_LOCAL)
    return dt.isoformat(timespec="seconds")

def parse_fecha_hikvision(valor, ajuste_tiempo=None):
    texto = str(valor or "").strip()
    if not texto:
        return None
    texto = texto.replace("Z", "+00:00")
    try:
        dt = datetime.fromisoformat(texto)
    except Exception:
        for formato in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S"):
            try:
                dt = datetime.strptime(texto[:19], formato)
                break
            except Exception:
                dt = None
        if dt is None:
            return None
    if dt.tzinfo is not None:
        dt = dt.astimezone(TZ_LOCAL).replace(tzinfo=None)
    if ajuste_tiempo:
        dt = dt + ajuste_tiempo
    return dt

def hikvision_probar_conexion(info):
    hikvision_request_json(info, "/ISAPI/System/deviceInfo?format=json", timeout=8)
    return "ISAPI"

def hikvision_obtener_personas(info):
    personas = set()
    posicion = 0
    max_resultados = 30
    search_id = hikvision_search_id()
    for _ in range(80):
        payload = {
            "UserInfoSearchCond": {
                "searchID": search_id,
                "searchResultPosition": posicion,
                "maxResults": max_resultados
            }
        }
        data = hikvision_request_json(info, "/ISAPI/AccessControl/UserInfo/Search?format=json", method="POST", payload=payload)
        root = data.get("UserInfoSearch", data.get("UserInfoSearchCond", data))
        lista = root.get("UserInfo", root.get("UserInfoList", []))
        if isinstance(lista, dict):
            lista = [lista]
        for item in lista or []:
            nombre = item.get("name") or item.get("employeeName") or item.get("employeeNo") or item.get("employeeNoString")
            if nombre and "admin" not in limpiar_nombre(nombre).lower():
                personas.add(limpiar_nombre(nombre))
        total = int(root.get("totalMatches", root.get("numOfMatches", 0)) or 0)
        cantidad = len(lista or [])
        posicion += cantidad
        if cantidad < max_resultados or (total and posicion >= total):
            break
    return sorted(personas)

def hikvision_obtener_eventos(info, fecha_inicio, fecha_fin, progreso=None, por_dia=True, cancelado=None):
    """Compatibilidad del ejecutable: la descarga real vive en el conector.

    ``por_dia`` se conserva para las llamadas heredadas. El conector central
    utiliza bloques mensuales para evitar miles de solicitudes por rango.
    """
    return hikvision_connector.hikvision_obtener_eventos(
        info,
        fecha_inicio,
        fecha_fin,
        progreso=progreso,
        cancelado=cancelado,
    )
def ruta_recurso(relative_path):
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_path, relative_path)

def configurar_app_windows():
    try:
        import ctypes
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("Asistencia.Agropecuaria.v4")
    except Exception:
        pass

def aplicar_icono(ventana):
    icono = ruta_recurso("icono.ico")
    if not os.path.exists(icono):
        return
    try:
        ventana.iconbitmap(icono)
    except Exception:
        pass
    try:
        from PIL import Image, ImageTk
        with Image.open(icono) as imagen:
            foto = ImageTk.PhotoImage(imagen.copy())
        ventana.wm_iconphoto(True, foto)
        ventana._app_icon_photo = foto
    except Exception:
        try:
            ventana.iconbitmap(icono)
        except Exception:
            pass

def configuracion_default():
    return {
        "checadores": [],
        "personas_sistema": [],
        "personas_eliminadas": [],
        "personas_biometria": [],
        "dias_no_laborados": [],
        "horarios": {
            "entrada": "08:00",
            "salida": "17:00",
            "tolerancia_min": 10,
            "dias_laborales": [0, 1, 2, 3, 4, 5],
            "especiales": []
        }
    }

def puerto_seguro(valor, respaldo=DEFAULT_PORT):
    try:
        puerto = int(str(valor).strip())
    except Exception:
        return respaldo
    if 1 <= puerto <= 65535:
        return puerto
    return respaldo

def entero_no_negativo(valor, respaldo=0):
    try:
        numero = int(str(valor).strip())
    except Exception:
        return respaldo
    return numero if numero >= 0 else respaldo

def opciones_horas(intervalo=5):
    return [f"{hora:02d}:{minuto:02d}" for hora in range(24) for minuto in range(0, 60, intervalo)]

def normalizar_dias_horarios(valor, base_horario):
    resultado = {}
    if not isinstance(valor, dict):
        return resultado
    for dia in [5, 6]:
        item = valor.get(str(dia), valor.get(dia, {}))
        if not isinstance(item, dict) or not item.get("activo", False):
            continue
        resultado[str(dia)] = {
            "activo": True,
            "entrada": item.get("entrada", base_horario.get("entrada", "08:00")),
            "salida": item.get("salida", base_horario.get("salida", "17:00")),
            "tolerancia_min": entero_no_negativo(item.get("tolerancia_min", base_horario.get("tolerancia_min", 10)), base_horario.get("tolerancia_min", 10))
        }
    return resultado

def fecha_iso_segura(valor):
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    texto = str(valor).strip()
    if not texto:
        return ""
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto[:10], formato).date().isoformat()
        except Exception:
            pass
    return ""

def normalizar_dia_no_laborado(item):
    if isinstance(item, dict):
        fecha_txt = fecha_iso_segura(item.get("fecha", ""))
        motivo = str(item.get("motivo", "")).strip()
    else:
        fecha_txt = fecha_iso_segura(item)
        motivo = ""
    if not fecha_txt:
        return None
    return {"fecha": fecha_txt, "motivo": motivo}

def normalizar_biometria_configuracion(config):
    """Keep only non-sensitive biometric workflow states in configuration."""
    estados_huella = {"pendiente", "no_solicitada", "requiere_captura_en_dispositivo", "captura_reportada_sin_verificar"}
    estados_rostro = {
        "pendiente", "no_solicitado", "foto_lista", "enviado_y_confirmado",
        "enviado_sin_verificar", "prueba_fisica_reportada", "error",
    }
    estados_modo_rostro = {"no_aplica", "sin_confirmar", "facial_confirmado", "no_facial", "error"}
    estados_acceso = {"no_aplica", "sin_confirmar", "permiso_confirmado", "sin_permiso", "error"}
    registros = []
    vistos = set()
    for item in config.get("personas_biometria", []):
        if not isinstance(item, dict):
            continue
        persona_id = str(item.get("persona_id", "")).strip()
        checador = str(item.get("checador", "")).strip()
        if not persona_id or not checador:
            continue
        clave = (persona_id.casefold(), checador.casefold())
        if clave in vistos:
            continue
        vistos.add(clave)
        huella = str(item.get("huella", "requiere_captura_en_dispositivo")).strip()
        rostro = str(item.get("rostro", "pendiente")).strip()
        registros.append({
            "persona_id": persona_id,
            "nombre": limpiar_nombre(item.get("nombre", "")),
            "checador": checador,
            "estado": "confirmada" if str(item.get("estado", "")).strip() == "confirmada" else "pendiente_obligatoria",
            "tipo_confirmado": str(item.get("tipo_confirmado", "")).strip() if str(item.get("tipo_confirmado", "")).strip() in {"huella", "rostro"} else "",
            "huella": huella if huella in estados_huella else "requiere_captura_en_dispositivo",
            "rostro": rostro if rostro in estados_rostro else "pendiente",
            "modo_rostro": (
                str(item.get("modo_rostro", "sin_confirmar")).strip()
                if str(item.get("modo_rostro", "sin_confirmar")).strip() in estados_modo_rostro
                else "sin_confirmar"
            ),
            "acceso": (
                str(item.get("acceso", "sin_confirmar")).strip()
                if str(item.get("acceso", "sin_confirmar")).strip() in estados_acceso
                else "sin_confirmar"
            ),
            "actualizado_en": str(item.get("actualizado_en", "")).strip(),
            "detalle": str(item.get("detalle", "")).strip()[:500],
        })
    config["personas_biometria"] = sorted(registros, key=lambda item: (item["checador"].casefold(), item["persona_id"].casefold()))
    # Legacy iVMS bookkeeping never represented an applied permission. Drop it
    # on every save so it cannot be mistaken for a real device state.
    config.pop("grupos_acceso", None)
    config.pop("accesos_personas", None)
    return config

def respaldar_archivo_config_corrupto(path):
    if not os.path.exists(path):
        return None
    marca_tiempo = datetime.now().strftime("%Y%m%d_%H%M%S")
    respaldo = f"{path}.corrupto.{marca_tiempo}.bak"
    try:
        shutil.copy2(path, respaldo)
        return respaldo
    except Exception:
        return None

def normalizar_horario(horario, base=None):
    default = configuracion_default()["horarios"]
    base = base or default
    if not isinstance(horario, dict):
        horario = {}
    limpio = {
        "entrada": horario.get("entrada", base.get("entrada", default["entrada"])),
        "salida": horario.get("salida", base.get("salida", default["salida"])),
        "tolerancia_min": horario.get("tolerancia_min", base.get("tolerancia_min", default["tolerancia_min"])),
        "dias_laborales": horario.get("dias_laborales", base.get("dias_laborales", default["dias_laborales"])),
        "dias_horarios": {},
        "especiales": []
    }
    limpio["tolerancia_min"] = entero_no_negativo(limpio["tolerancia_min"], default["tolerancia_min"])
    if not isinstance(limpio["dias_laborales"], list):
        limpio["dias_laborales"] = list(default["dias_laborales"])
    limpio["dias_horarios"] = normalizar_dias_horarios(horario.get("dias_horarios", base.get("dias_horarios", {})), limpio)
    for especial in horario.get("especiales", []):
        if not isinstance(especial, dict):
            continue
        especial_limpio = {
            "nombre": especial.get("nombre", "Horario especial"),
            "entrada": especial.get("entrada", limpio["entrada"]),
            "salida": especial.get("salida", limpio["salida"]),
            "tolerancia_min": entero_no_negativo(especial.get("tolerancia_min", limpio["tolerancia_min"]), limpio["tolerancia_min"]),
            "dias_laborales": especial.get("dias_laborales", list(limpio["dias_laborales"])),
            "dias_horarios": {},
            "empleados": [limpiar_nombre(emp) for emp in especial.get("empleados", []) if str(emp).strip()]
        }
        especial_limpio["dias_horarios"] = normalizar_dias_horarios(especial.get("dias_horarios", limpio.get("dias_horarios", {})), especial_limpio)
        limpio["especiales"].append({
            **especial_limpio
        })
    return limpio


DIAS_RETENCION_PERSONAS_ELIMINADAS = 60


def _fecha_eliminacion_segura(valor):
    if isinstance(valor, datetime):
        return valor.replace(tzinfo=None)
    if isinstance(valor, date):
        return datetime.combine(valor, datetime.min.time())
    texto = str(valor or "").strip()
    if not texto:
        return None
    try:
        fecha = datetime.fromisoformat(texto.replace("Z", "+00:00"))
        return fecha.replace(tzinfo=None)
    except (TypeError, ValueError):
        pass
    for formato in ("%Y-%m-%d %H:%M", "%Y-%m-%d", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    return None


def eliminar_personas_definitivamente_config(config, nombres):
    """Elimina personas de referencias locales y conserva el historial de asistencia."""
    if not isinstance(config, dict):
        return []
    nombres_limpios = sorted({limpiar_nombre(nombre) for nombre in nombres if str(nombre).strip()})
    claves = {nombre.casefold() for nombre in nombres_limpios}
    if not claves:
        return []

    def conservar_nombre(nombre):
        return limpiar_nombre(nombre).casefold() not in claves

    config["personas_sistema"] = [
        persona for persona in config.get("personas_sistema", []) if conservar_nombre(persona)
    ]
    config["personas_eliminadas"] = [
        item for item in config.get("personas_eliminadas", [])
        if conservar_nombre(item.get("nombre", "") if isinstance(item, dict) else item)
    ]
    config["personas_biometria"] = [
        item for item in config.get("personas_biometria", [])
        if not isinstance(item, dict) or conservar_nombre(item.get("nombre", ""))
    ]

    horarios = []
    if isinstance(config.get("horarios"), dict):
        horarios.append(config["horarios"])
    for checador in config.get("checadores", []):
        if not isinstance(checador, dict):
            continue
        for campo in ("empleados", "empleados_descargados"):
            checador[campo] = [
                persona for persona in checador.get(campo, []) if conservar_nombre(persona)
            ]
        if isinstance(checador.get("horarios"), dict):
            horarios.append(checador["horarios"])

    for horario in horarios:
        for especial in horario.get("especiales", []):
            if isinstance(especial, dict):
                especial["empleados"] = [
                    persona for persona in especial.get("empleados", []) if conservar_nombre(persona)
                ]
    return nombres_limpios


def purgar_personas_eliminadas_vencidas(config, ahora=None, dias=DIAS_RETENCION_PERSONAS_ELIMINADAS):
    """Elimina definitivamente del sistema local a quien cumpla el periodo de retención."""
    ahora = (ahora or datetime.now()).replace(tzinfo=None)
    vencidas = []
    for item in config.get("personas_eliminadas", []) if isinstance(config, dict) else []:
        if not isinstance(item, dict):
            continue
        fecha = _fecha_eliminacion_segura(item.get("fecha", ""))
        nombre = limpiar_nombre(item.get("nombre", ""))
        if nombre and fecha is not None and ahora - fecha >= timedelta(days=dias):
            vencidas.append(nombre)
    return eliminar_personas_definitivamente_config(config, vencidas)

def normalizar_configuracion(config):
    default = configuracion_default()
    if not isinstance(config, dict):
        config = {}
    if not isinstance(config.get("checadores"), list):
        config["checadores"] = default["checadores"]
    horarios = config.get("horarios")
    if not isinstance(horarios, dict):
        horarios = {}
    horarios = normalizar_horario(horarios)
    config["horarios"] = horarios
    checadores = []
    for checador in config["checadores"]:
        if not isinstance(checador, dict):
            continue
        checador.setdefault("nombre", "Sin Nombre")
        checador.setdefault("ip", "")
        checador["marca"] = marca_checador(checador)
        checador["port"] = puerto_seguro(checador.get("port"), puerto_default_marca(checador["marca"]))
        if "modelo" in checador:
            checador["modelo"] = str(checador.get("modelo", "")).strip().upper()
        if checador["marca"] == "hikvision":
            checador["usuario"] = str(checador.get("usuario", "admin")).strip() or "admin"
            checador["password"] = str(checador.get("password", ""))
        checador["empleados"] = sorted({limpiar_nombre(emp) for emp in checador.get("empleados", []) if str(emp).strip()})
        checador["empleados_descargados"] = sorted({limpiar_nombre(emp) for emp in checador.get("empleados_descargados", []) if str(emp).strip()})
        if not checador["empleados_descargados"] and checador["empleados"]:
            checador["empleados_descargados"] = list(checador["empleados"])
        checador["horarios"] = normalizar_horario(checador.get("horarios"), horarios)
        checadores.append(checador)
    if not checadores:
        for checador in default["checadores"]:
            nuevo = dict(checador)
            nuevo["marca"] = marca_checador(nuevo)
            nuevo["port"] = puerto_seguro(nuevo.get("port"), puerto_default_marca(nuevo["marca"]))
            nuevo["empleados"] = []
            nuevo["horarios"] = normalizar_horario(None, horarios)
            checadores.append(nuevo)
    config["checadores"] = checadores
    personas_sistema = {limpiar_nombre(emp) for emp in config.get("personas_sistema", []) if str(emp).strip()}
    for checador in checadores:
        personas_sistema.update(limpiar_nombre(emp) for emp in checador.get("empleados", []) if str(emp).strip())
        for especial in checador.get("horarios", {}).get("especiales", []):
            personas_sistema.update(limpiar_nombre(emp) for emp in especial.get("empleados", []) if str(emp).strip())
    config["personas_sistema"] = sorted(personas_sistema)
    eliminadas = []
    for item in config.get("personas_eliminadas", []):
        if isinstance(item, dict):
            nombre = limpiar_nombre(item.get("nombre", ""))
            if nombre:
                eliminadas.append({
                    "nombre": nombre,
                    "fecha": item.get("fecha", ""),
                    "checador": item.get("checador", "")
                })
        elif str(item).strip():
            eliminadas.append({"nombre": limpiar_nombre(item), "fecha": "", "checador": ""})
    vistos = set()
    limpias = []
    for item in eliminadas:
        key = item["nombre"].lower()
        if key in vistos:
            continue
        vistos.add(key)
        limpias.append(item)
    config["personas_eliminadas"] = limpias
    eliminadas_nombres = {item["nombre"].lower() for item in limpias}
    config["personas_sistema"] = [emp for emp in config["personas_sistema"] if emp.lower() not in eliminadas_nombres]
    dias_no_laborados = {}
    for item in config.get("dias_no_laborados", []):
        limpio = normalizar_dia_no_laborado(item)
        if limpio:
            dias_no_laborados[limpio["fecha"]] = limpio
    config["dias_no_laborados"] = [dias_no_laborados[fecha] for fecha in sorted(dias_no_laborados)]
    config = normalizar_biometria_configuracion(config)
    purgar_personas_eliminadas_vencidas(config)
    return config

def cargar_configuracion():
    os.makedirs(APP_DIR, exist_ok=True)
    if not os.path.exists(PATH_CONFIG):
        for origen in (PATH_CONFIG_ANTERIOR, PATH_CONFIG_LOCAL):
            if not origen or not os.path.exists(origen):
                continue
            try:
                if os.path.abspath(origen) != os.path.abspath(PATH_CONFIG):
                    shutil.copy2(origen, PATH_CONFIG)
                    break
            except Exception:
                pass
    if not os.path.exists(PATH_CONFIG):
        default = normalizar_configuracion(configuracion_default())
        guardar_configuracion(default)
        return default
    try:
        with open(PATH_CONFIG, "r", encoding="utf-8") as f:
            config = json.load(f)
    except Exception:
        respaldar_archivo_config_corrupto(PATH_CONFIG)
        config = configuracion_default()
    config = normalizar_configuracion(config)
    guardar_configuracion(config)
    return config

def guardar_configuracion(config):
    os.makedirs(APP_DIR, exist_ok=True)
    temp_path = f"{PATH_CONFIG}.tmp"
    with open(temp_path, "w", encoding="utf-8") as f:
        json.dump(normalizar_configuracion(config), f, indent=4)
    os.replace(temp_path, PATH_CONFIG)

meses = {1: "enero", 2: "febrero", 3: "marzo", 4: "abril", 5: "mayo", 6: "junio", 
         7: "julio", 8: "agosto", 9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre"}
meses_abrev = {1: "Ene", 2: "Feb", 3: "Mar", 4: "Abr", 5: "May", 6: "Jun", 7: "Jul", 8: "Ago", 9: "Sep", 10: "Oct", 11: "Nov", 12: "Dic"}
dias_sem = {0: "Lunes", 1: "Martes", 2: "Miércoles", 3: "Jueves", 4: "Viernes", 5: "Sábado", 6: "Domingo"}

# --- PALETA DE COLORES ---
BG_APP = "#0F172A"         
BG_CARD = "#1E293B"        
BG_CARD_ALT = "#253347"    
BORDER_CARD = "#334155"    
FG_TEXT = "#F8FAFC"        
FG_DIM = "#94A3B8"         
ACCENT = "#3B82F6"         
SUCCESS = "#10B981"        
EXCEL_COL = "#16A34A"      
DANGER = "#EF4444"         
EXCEL_SABADO_EXTRA_FILL = "CCFBF1"
EXCEL_SABADO_EXTRA_FONT = "0F766E"
HORA_CORTE_SALIDA_MIN = 14 * 60

# --- VARIABLES GLOBALES ---
datos_memoria_crudos = None  
checador_memoria = ""
rango_memoria_inicio = None
rango_memoria_fin = None
ultimo_reporte = ""
titulo_actual = ""
period_frame_widgets = []
usuario_actual = None
bloqueo_instancia = None


def tiene_permiso(permiso):
    return auth_service.can(usuario_actual, permiso)


def requerir_permiso(permiso):
    if tiene_permiso(permiso):
        return True
    nombre = auth_service.PERMISSIONS.get(permiso, "esta acción")
    messagebox.showwarning("Acceso restringido", f"Tu usuario no tiene permiso para: {nombre}.")
    return False


def ejecutar_accion_protegida(permiso, accion):
    if requerir_permiso(permiso):
        return accion()
    return None


def actualizar_etiqueta_usuario():
    if "label_usuario_actual" not in globals() or usuario_actual is None:
        return
    label_usuario_actual.config(text=f"Sesión: {usuario_actual['display_name']} ({usuario_actual['username']})")

# --- UTILIDADES ---
def limpiar_nombre(nombre):
    return str(nombre).replace(".", " ").title().strip()

def formato_hora(hora):
    try:
        pd = cargar_pandas()
        if pd.isna(hora): return "---"
        return hora.strftime("%I:%M %p").lstrip("0")
    except: return str(hora)

def formato_entrada_reporte(entrada):
    return "SIN ENTRADA" if hora_vacia(entrada) else formato_hora(entrada)

def formato_salida_reporte(entrada, salida):
    if hora_vacia(salida):
        return "SIN SALIDA"
    if not hora_vacia(entrada) and entrada == salida:
        return "SIN SALIDA"
    return formato_hora(salida)

def parse_hora_config(valor, respaldo):
    try:
        return datetime.strptime(str(valor).strip(), "%H:%M").time()
    except Exception:
        return datetime.strptime(respaldo, "%H:%M").time()

def minutos_de_hora(hora):
    return hora.hour * 60 + hora.minute

def segundos_de_hora(hora):
    return hora.hour * 3600 + hora.minute * 60 + hora.second

def hora_vacia(hora):
    if hora is None:
        return True
    try:
        pd = cargar_pandas()
        return bool(pd.isna(hora))
    except Exception:
        return False

def rango_fechas(fecha_inicio, fecha_fin):
    dias = []
    actual = fecha_inicio
    while actual <= fecha_fin:
        dias.append(actual)
        actual += timedelta(days=1)
    return dias

def obtener_mapa_dias_no_laborados(config=None):
    config = config or cargar_configuracion()
    mapa = {}
    for item in config.get("dias_no_laborados", []):
        limpio = normalizar_dia_no_laborado(item)
        if not limpio:
            continue
        try:
            fecha = datetime.strptime(limpio["fecha"], "%Y-%m-%d").date()
        except Exception:
            continue
        mapa[fecha] = limpio.get("motivo", "")
    return mapa

def es_dia_no_laborado(fecha, dias_no_laborados=None):
    if fecha is None:
        return False
    if dias_no_laborados is None:
        dias_no_laborados = obtener_mapa_dias_no_laborados()
    return fecha in dias_no_laborados

def motivo_dia_no_laborado(fecha, dias_no_laborados=None):
    if fecha is None:
        return ""
    if dias_no_laborados is None:
        dias_no_laborados = obtener_mapa_dias_no_laborados()
    return dias_no_laborados.get(fecha, "")

def opciones_meses_reporte():
    hoy = date.today()
    return [f"{meses[m].capitalize()} {hoy.year}" for m in range(1, hoy.month + 1)]

def rango_mes_actual(indice_mes):
    hoy = date.today()
    mes = max(1, min(indice_mes, hoy.month))
    inicio = date(hoy.year, mes, 1)
    if mes == 12:
        fin_mes = date(hoy.year, 12, 31)
    else:
        fin_mes = date(hoy.year, mes + 1, 1) - timedelta(days=1)
    fin = min(fin_mes, hoy) if mes == hoy.month else fin_mes
    return inicio, fin

def obtener_horarios_config(nombre_checador=None):
    config = cargar_configuracion()
    if nombre_checador:
        info = next((c for c in config["checadores"] if c["nombre"] == nombre_checador), None)
        if info:
            return normalizar_horario(info.get("horarios"), config.get("horarios"))
    return normalizar_horario(config.get("horarios"))

def obtener_horario_empleado(nombre_empleado, horarios):
    nombre_normalizado = limpiar_nombre(nombre_empleado).lower()
    horarios = normalizar_horario(horarios)
    for especial in horarios.get("especiales", []):
        empleados = {limpiar_nombre(emp).lower() for emp in especial.get("empleados", [])}
        if nombre_normalizado in empleados:
            return normalizar_horario(especial, horarios)
    return horarios

def es_dia_laboral(fecha, horarios=None, dias_no_laborados=None):
    if dias_no_laborados is not None and es_dia_no_laborado(fecha, dias_no_laborados):
        return False
    horarios = horarios or obtener_horarios_config()
    return fecha.weekday() in horarios.get("dias_laborales", [0, 1, 2, 3, 4, 5])

def obtener_horario_para_fecha(horarios, fecha):
    horarios = normalizar_horario(horarios)
    if fecha is None:
        return horarios
    especial_dia = horarios.get("dias_horarios", {}).get(str(fecha.weekday()))
    if especial_dia and especial_dia.get("activo"):
        horario_dia = dict(horarios)
        horario_dia["entrada"] = especial_dia.get("entrada", horarios.get("entrada", "08:00"))
        horario_dia["salida"] = especial_dia.get("salida", horarios.get("salida", "17:00"))
        horario_dia["tolerancia_min"] = especial_dia.get("tolerancia_min", horarios.get("tolerancia_min", 10))
        return horario_dia
    return horarios

def formato_duracion(minutos):
    if minutos is None or minutos < 0:
        return "---"
    horas = minutos // 60
    mins = minutos % 60
    return f"{horas}:{mins:02d}"

def calcular_detalles_asistencia(entrada, salida, horarios, falta=False, fecha=None, dias_no_laborados=None):
    if fecha is not None and es_dia_no_laborado(fecha, dias_no_laborados):
        if hora_vacia(entrada) or hora_vacia(salida):
            horas_trabajadas = "---"
        else:
            ent_min = minutos_de_hora(entrada)
            sal_min = minutos_de_hora(salida)
            if salida == entrada:
                horas_trabajadas = "---"
            else:
                if sal_min < ent_min:
                    sal_min += 24 * 60
                horas_trabajadas = formato_duracion(sal_min - ent_min)
        return {
            "Retardo": "No cuenta",
            "Retardo_SN": "NO",
            "Tiempo_Retardo": "0 min",
            "Salida_Temprana": "No cuenta",
            "Horas_Trabajadas": horas_trabajadas,
            "Tiene_Retardo": False,
            "Tiene_Salida_Temprana": False,
            "Dia_Laboral": False,
            "Dia_No_Laboral": True,
            "Extra_Sabado": False,
            "Domingo_No_Laboral": False,
            "Sin_Entrada": False,
            "No_Laborado": True,
            "Motivo_No_Laborado": motivo_dia_no_laborado(fecha, dias_no_laborados)
        }

    if falta:
        return {
            "Retardo": "FALTA",
            "Retardo_SN": "FALTA",
            "Tiempo_Retardo": "FALTA",
            "Salida_Temprana": "FALTA",
            "Horas_Trabajadas": "0:00",
            "Tiene_Retardo": False,
            "Tiene_Salida_Temprana": False,
            "Dia_Laboral": True,
            "Dia_No_Laboral": False,
            "Extra_Sabado": False,
            "Domingo_No_Laboral": False,
            "Sin_Entrada": False,
            "No_Laborado": False,
            "Motivo_No_Laborado": ""
        }

    dia_laboral = True if fecha is None else es_dia_laboral(fecha, horarios, dias_no_laborados)
    if fecha is not None and not dia_laboral:
        if hora_vacia(entrada) or hora_vacia(salida) or salida == entrada:
            horas_trabajadas = "---"
        else:
            ent_min = minutos_de_hora(entrada)
            sal_min = minutos_de_hora(salida)
            if sal_min < ent_min:
                sal_min += 24 * 60
            horas_trabajadas = formato_duracion(sal_min - ent_min)
        return {
            "Retardo": "No cuenta",
            "Retardo_SN": "NO",
            "Tiempo_Retardo": "0 min",
            "Salida_Temprana": "No cuenta",
            "Horas_Trabajadas": horas_trabajadas,
            "Tiene_Retardo": False,
            "Tiene_Salida_Temprana": False,
            "Dia_Laboral": False,
            "Dia_No_Laboral": True,
            "Extra_Sabado": fecha.weekday() == 5,
            "Domingo_No_Laboral": fecha.weekday() == 6,
            "Sin_Entrada": False,
            "No_Laborado": False,
            "Motivo_No_Laborado": ""
        }

    horarios = obtener_horario_para_fecha(horarios, fecha)
    hora_entrada_cfg = parse_hora_config(horarios.get("entrada"), "08:00")
    hora_salida_cfg = parse_hora_config(horarios.get("salida"), "17:00")
    tolerancia = entero_no_negativo(horarios.get("tolerancia_min", 10), 10)

    if hora_vacia(entrada) and not hora_vacia(salida):
        return {
            "Retardo": "SIN ENTRADA",
            "Retardo_SN": "SIN ENTRADA",
            "Tiempo_Retardo": "SIN ENTRADA",
            "Salida_Temprana": "No",
            "Horas_Trabajadas": "---",
            "Tiene_Retardo": False,
            "Tiene_Salida_Temprana": False,
            "Dia_Laboral": True,
            "Dia_No_Laboral": False,
            "Extra_Sabado": False,
            "Domingo_No_Laboral": False,
            "Sin_Entrada": True,
            "No_Laborado": False,
            "Motivo_No_Laborado": ""
        }

    ent_min = minutos_de_hora(entrada)
    sal_min = minutos_de_hora(salida)
    entrada_limite = minutos_de_hora(hora_entrada_cfg) + tolerancia
    salida_esperada = minutos_de_hora(hora_salida_cfg)

    retardo_min = ent_min - entrada_limite
    tiene_retardo = retardo_min > 0
    retardo = f"Si (+{retardo_min} min)" if tiene_retardo else "No"
    retardo_sn = "SI" if tiene_retardo else "NO"
    tiempo_retardo = f"{retardo_min} min" if tiene_retardo else "0 min"

    tiene_salida_temprana = salida != entrada and sal_min < salida_esperada
    salida_temprana = f"Si (-{salida_esperada - sal_min} min)" if tiene_salida_temprana else "No"

    if salida == entrada:
        horas_trabajadas = "---"
    else:
        if sal_min < ent_min:
            sal_min += 24 * 60
        horas_trabajadas = formato_duracion(sal_min - ent_min)

    return {
        "Retardo": retardo,
        "Retardo_SN": retardo_sn,
        "Tiempo_Retardo": tiempo_retardo,
        "Salida_Temprana": salida_temprana,
        "Horas_Trabajadas": horas_trabajadas,
        "Tiene_Retardo": tiene_retardo,
        "Tiene_Salida_Temprana": tiene_salida_temprana,
        "Dia_Laboral": True,
        "Dia_No_Laboral": False,
        "Extra_Sabado": False,
        "Domingo_No_Laboral": False,
        "Sin_Entrada": False,
        "No_Laborado": False,
        "Motivo_No_Laborado": ""
    }

def preparar_base_asistencia(df_crudo):
    pd = cargar_pandas()
    if df_crudo is None or df_crudo.empty:
        return pd.DataFrame(columns=["nombre", "fecha_hora", "Fecha_D", "Hora_D"])
    df = df_crudo.dropna(subset=["fecha_hora", "nombre"]).copy()
    df = df[~df["nombre"].str.contains("admin", case=False, na=False)]
    df["nombre"] = df["nombre"].apply(limpiar_nombre)
    df["Fecha_D"] = df["fecha_hora"].dt.date
    df["Hora_D"] = df["fecha_hora"].dt.time
    return df

def obtener_entrada_salida_por_corte(horas):
    horas = sorted([h for h in horas if not hora_vacia(h)], key=segundos_de_hora)
    if not horas:
        return None, None
    entradas = [h for h in horas if minutos_de_hora(h) < HORA_CORTE_SALIDA_MIN]
    salidas = [h for h in horas if minutos_de_hora(h) >= HORA_CORTE_SALIDA_MIN]
    entrada = entradas[0] if entradas else None
    if salidas:
        salida = salidas[-1]
    else:
        salida = entrada if entrada else horas[-1]
    return entrada, salida

def obtener_entrada_salida_primera_ultima(horas):
    horas = sorted([h for h in horas if not hora_vacia(h)], key=segundos_de_hora)
    if not horas:
        return None, None
    return horas[0], horas[-1]

def usa_primera_ultima_checador(checador_info):
    return bool(checador_info and marca_checador(checador_info) == "hikvision")

def agrupar_entrada_salida(df_rango, usar_primera_ultima=False):
    pd = cargar_pandas()
    if df_rango.empty:
        return pd.DataFrame(columns=["nombre", "Fecha_D", "Entrada", "Salida"])
    registros = []
    for (nombre, fecha), grupo in df_rango.groupby(["nombre", "Fecha_D"]):
        horas = grupo["Hora_D"].tolist()
        if usar_primera_ultima:
            entrada, salida = obtener_entrada_salida_primera_ultima(horas)
        else:
            entrada, salida = obtener_entrada_salida_por_corte(horas)
        registros.append({"nombre": nombre, "Fecha_D": fecha, "Entrada": entrada, "Salida": salida})
    return pd.DataFrame(registros)

def construir_reporte_asistencia(df_crudo, fecha_inicio, fecha_fin, horarios=None, incluir_faltas=True, empleados_asignados=None, usar_primera_ultima=False):
    pd = cargar_pandas()
    horarios = horarios or obtener_horarios_config()
    dias_no_laborados = obtener_mapa_dias_no_laborados()
    base = preparar_base_asistencia(df_crudo)
    empleados_config = {limpiar_nombre(emp) for emp in (empleados_asignados or []) if str(emp).strip()}
    empleados_especiales = set()
    for especial in normalizar_horario(horarios).get("especiales", []):
        empleados_especiales.update(limpiar_nombre(emp) for emp in especial.get("empleados", []) if str(emp).strip())
    empleados_permitidos = empleados_config | empleados_especiales
    if empleados_permitidos and not base.empty:
        base = base[base["nombre"].isin(empleados_permitidos)].copy()
    empleados_base = set(base["nombre"].unique()) if not base.empty else set()
    empleados = sorted(empleados_base | empleados_permitidos)
    fechas_periodo = rango_fechas(fecha_inicio, fecha_fin)
    df_rango = base[(base["Fecha_D"] >= fecha_inicio) & (base["Fecha_D"] <= fecha_fin)].copy()

    if df_rango.empty:
        reporte = pd.DataFrame(columns=["nombre", "Fecha_D", "Entrada", "Salida", "Falta"])
    else:
        reporte = agrupar_entrada_salida(df_rango, usar_primera_ultima=usar_primera_ultima)
        reporte["Falta"] = False

    if incluir_faltas and empleados:
        existentes = set(zip(reporte["nombre"], reporte["Fecha_D"])) if not reporte.empty else set()
        faltas = []
        for emp in empleados:
            horario_emp = obtener_horario_empleado(emp, horarios)
            for fecha in fechas_periodo:
                if es_dia_laboral(fecha, horario_emp, dias_no_laborados) and (emp, fecha) not in existentes:
                    faltas.append({"nombre": emp, "Fecha_D": fecha, "Entrada": None, "Salida": None, "Falta": True})
        if faltas:
            reporte = pd.concat([reporte, pd.DataFrame(faltas)], ignore_index=True)

    if reporte.empty:
        return reporte, base, fechas_periodo

    detalles = []
    for _, row in reporte.iterrows():
        horario_emp = obtener_horario_empleado(row["nombre"], horarios)
        detalles.append(calcular_detalles_asistencia(row["Entrada"], row["Salida"], horario_emp, bool(row.get("Falta", False)), row["Fecha_D"], dias_no_laborados))
    detalles_df = pd.DataFrame(detalles)
    reporte = pd.concat([reporte.reset_index(drop=True), detalles_df], axis=1)
    reporte = reporte.sort_values(by=["nombre", "Fecha_D"]).reset_index(drop=True)
    return reporte, base, fechas_periodo

def set_boton_estado(btn, estado, color_activo):
    if estado == "normal":
        btn.config(state="normal", bg=color_activo, fg=FG_TEXT, cursor="hand2")
    else:
        btn.config(state="disabled", bg="#334155", fg="#64748B", cursor="arrow")

def ajustar_ventana(ventana, ancho, alto, min_ancho=None, min_alto=None):
    ventana.update_idletasks()
    margen_ancho = 70
    margen_alto = 95
    screen_w = ventana.winfo_screenwidth()
    screen_h = ventana.winfo_screenheight()
    ancho_final = min(ancho, max(360, screen_w - margen_ancho))
    alto_final = min(alto, max(360, screen_h - margen_alto))
    x = max(0, (screen_w - ancho_final) // 2)
    y = max(0, (screen_h - alto_final) // 2 - 15)
    ventana.geometry(f"{ancho_final}x{alto_final}+{x}+{y}")
    ventana.minsize(min(min_ancho or ancho_final, ancho_final), min(min_alto or alto_final, alto_final))

# --- VENTANAS MEJORADAS (UX/UI) ---

class VentanaConfiguracion(Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        aplicar_icono(self)
        self.title("Administrador de Dispositivos")
        self.configure(bg=BG_APP)
        self.resizable(True, True)
        self.config_local = cargar_configuracion()
        self.init_ui()
        ajustar_ventana(self, 780, 650, 700, 560)

    def init_ui(self):
        # Header de la ventana
        header = tk.Frame(self, bg=BG_APP, pady=8)
        header.pack(fill="x", padx=20)
        tk.Label(header, text="Configuración de Dispositivos", bg=BG_APP, fg=FG_TEXT, font=("Segoe UI", 15, "bold")).pack(side="left")
        tk.Label(header, text="Gestiona tus biométricos", bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 9)).pack(side="left", padx=10, pady=(5,0))

        main_frame = tk.Frame(self, bg=BG_APP)
        main_frame.pack(fill="both", expand=True, padx=20, pady=(0, 16))

        # PANEL IZQUIERDO: Lista de dispositivos (Estilo Menú Lateral)
        col_izq = tk.Frame(main_frame, bg=BG_CARD, highlightthickness=1, highlightbackground=BORDER_CARD)
        col_izq.pack(side="left", fill="y", ipadx=8, ipady=8)
        
        tk.Label(col_izq, text="Equipos Registrados", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 10, "bold")).pack(anchor="w", padx=10, pady=(8, 8))
        
        # Lista sin bordes para que se funda con el fondo
        self.listbox = tk.Listbox(col_izq, bg=BG_CARD, fg=FG_TEXT, bd=0, highlightthickness=0, selectbackground=ACCENT, font=("Segoe UI", 11), activestyle="none")
        self.listbox.pack(fill="both", expand=True, padx=10)
        self.listbox.bind("<<ListboxSelect>>", self.cargar_datos)

        tk.Button(col_izq, text="+ Agregar Nuevo", command=self.nuevo_registro, bg=SUCCESS, fg="white", font=("Segoe UI", 10, "bold"), bd=0, pady=8, cursor="hand2").pack(fill="x", padx=10, pady=8)

        # PANEL DERECHO: Formulario de edición (Estilo Tarjeta)
        col_der = tk.Frame(main_frame, bg=BG_APP)
        col_der.pack(side="right", fill="both", expand=True, padx=(18, 0))

        card_form = tk.Frame(col_der, bg=BG_CARD, padx=22, pady=18, highlightthickness=1, highlightbackground=BORDER_CARD)
        card_form.pack(fill="both", expand=True)

        tk.Label(card_form, text="Detalles de Conexión", bg=BG_CARD, fg=FG_TEXT, font=("Segoe UI", 13, "bold")).pack(anchor="w", pady=(0, 14))

        self.var_nombre = tk.StringVar()
        self.var_ip = tk.StringVar()
        self.var_puerto = tk.StringVar()
        self.var_marca = tk.StringVar(value="ZKTeco")
        self.var_usuario = tk.StringVar()
        self.var_password = tk.StringVar()

        marca_frame = tk.Frame(card_form, bg=BG_CARD)
        marca_frame.pack(fill="x", pady=(0, 8))
        tk.Label(marca_frame, text="Marca:", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 10)).pack(anchor="w")
        self.cb_marca = ttk.Combobox(marca_frame, textvariable=self.var_marca, values=["ZKTeco", "Hikvision"], state="readonly", font=("Segoe UI", 10))
        self.cb_marca.pack(fill="x", ipady=4, pady=(4, 0))
        self.cb_marca.bind("<<ComboboxSelected>>", self.actualizar_formulario_marca)

        # Generador de inputs estilizados
        for lbl_text, var in [("Nombre Asignado:", self.var_nombre), ("Dirección IPv4:", self.var_ip), ("Puerto de Red:", self.var_puerto)]:
            frame_input = tk.Frame(card_form, bg=BG_CARD)
            frame_input.pack(fill="x", pady=(0, 8))
            tk.Label(frame_input, text=lbl_text, bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 10)).pack(anchor="w")
            ent = tk.Entry(frame_input, textvariable=var, bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 11), bd=0, insertbackground=FG_TEXT, disabledbackground=BG_APP)
            ent.pack(fill="x", ipady=6, pady=(4,0))
            if lbl_text == "Nombre Asignado:": self.ent_nombre = ent
            elif lbl_text == "Dirección IPv4:": self.ent_ip = ent
            else: self.ent_puerto = ent

        self.frame_usuario = tk.Frame(card_form, bg=BG_CARD)
        self.frame_usuario.pack(fill="x", pady=(0, 8))
        tk.Label(self.frame_usuario, text="Usuario Hikvision:", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 10)).pack(anchor="w")
        self.ent_usuario = tk.Entry(self.frame_usuario, textvariable=self.var_usuario, bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 11), bd=0, insertbackground=FG_TEXT, disabledbackground=BG_APP)
        self.ent_usuario.pack(fill="x", ipady=6, pady=(4, 0))

        self.frame_password = tk.Frame(card_form, bg=BG_CARD)
        self.frame_password.pack(fill="x", pady=(0, 8))
        tk.Label(self.frame_password, text="Password Hikvision:", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 10)).pack(anchor="w")
        self.ent_password = tk.Entry(self.frame_password, textvariable=self.var_password, show="*", bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 11), bd=0, insertbackground=FG_TEXT, disabledbackground=BG_APP)
        self.ent_password.pack(fill="x", ipady=6, pady=(4, 0))

        # Separador
        self.sep_config = tk.Frame(card_form, bg=BORDER_CARD, height=1)
        self.sep_config.pack(fill="x", pady=10)

        # Botones de Acción
        f_btns = tk.Frame(card_form, bg=BG_CARD)
        f_btns.pack(fill="x")
        
        self.btn_modificar = tk.Button(f_btns, text="Modificar Datos", command=self.toggle_edicion, bg=ACCENT, fg="white", font=("Segoe UI", 10, "bold"), bd=0, pady=7, cursor="hand2", width=15)
        self.btn_modificar.pack(side="left", padx=(0, 10))

        tk.Button(f_btns, text="Probar Conexión", command=self.probar_conexion, bg=BG_CARD_ALT, fg=SUCCESS, font=("Segoe UI", 10, "bold"), bd=1, pady=7, cursor="hand2").pack(side="left", fill="x", expand=True)
        
        tk.Button(card_form, text="Eliminar Checador", command=self.eliminar_equipo, bg="#3B1F2A", fg=DANGER, font=("Segoe UI", 10, "bold"), bd=1, relief="flat", activebackground="#4A2633", activeforeground="#FCA5A5", cursor="hand2", pady=7).pack(fill="x", pady=(12, 0))

        self.actualizar_lista()
        self.actualizar_formulario_marca()

    def actualizar_lista(self):
        self.listbox.delete(0, tk.END)
        for c in self.config_local["checadores"]:
            marca_txt = "Hikvision" if marca_checador(c) == "hikvision" else "ZKTeco"
            self.listbox.insert(tk.END, f"  {c['nombre']}  [{marca_txt}]")

    def actualizar_formulario_marca(self, event=None):
        marca = "hikvision" if self.var_marca.get() == "Hikvision" else "zkteco"
        puerto_actual = puerto_seguro(self.var_puerto.get(), None)
        if not self.var_puerto.get().strip() or puerto_actual in [DEFAULT_PORT, HIKVISION_DEFAULT_PORT]:
            self.var_puerto.set(str(puerto_default_marca(marca)))
        if marca == "hikvision":
            self.frame_usuario.pack(fill="x", pady=(0, 8), before=self.sep_config)
            self.frame_password.pack(fill="x", pady=(0, 8), before=self.sep_config)
            if not self.var_usuario.get().strip():
                self.var_usuario.set("admin")
        else:
            self.frame_usuario.pack_forget()
            self.frame_password.pack_forget()
            self.var_usuario.set("")
            self.var_password.set("")

    def cargar_datos(self, event):
        idx = self.listbox.curselection()
        if not idx: return
        data = self.config_local["checadores"][idx[0]]
        self.estado_campos("normal")
        marca = marca_checador(data)
        self.var_marca.set("Hikvision" if marca == "hikvision" else "ZKTeco")
        self.var_nombre.set(data["nombre"]); self.var_ip.set(data["ip"]); self.var_puerto.set(data.get("port", puerto_default_marca(marca)))
        self.var_usuario.set(data.get("usuario", "admin" if marca == "hikvision" else ""))
        self.var_password.set(data.get("password", ""))
        self.actualizar_formulario_marca()
        self.estado_campos("disabled")
        self.btn_modificar.config(text="Modificar Datos", bg=ACCENT)

    def estado_campos(self, estado):
        self.cb_marca.config(state="readonly" if estado == "normal" else "disabled")
        self.ent_nombre.config(state=estado)
        self.ent_ip.config(state=estado)
        self.ent_puerto.config(state=estado)
        self.ent_usuario.config(state=estado)
        self.ent_password.config(state=estado)

    def nuevo_registro(self):
        self.listbox.selection_clear(0, tk.END)
        self.estado_campos("normal")
        self.var_marca.set("ZKTeco")
        self.var_nombre.set(""); self.var_ip.set(""); self.var_puerto.set(str(DEFAULT_PORT))
        self.var_usuario.set(""); self.var_password.set("")
        self.actualizar_formulario_marca()
        self.ent_nombre.focus()
        self.btn_modificar.config(text="Guardar Nuevo", bg=SUCCESS)

    def toggle_edicion(self):
        if self.btn_modificar.cget("text") in ["Modificar Datos"]:
            if not self.var_nombre.get(): return
            self.estado_campos("normal")
            self.ent_nombre.focus()
            self.btn_modificar.config(text="Guardar Cambios", bg=SUCCESS)
        else:
            n, i, p = self.var_nombre.get().strip(), self.var_ip.get().strip(), self.var_puerto.get().strip()
            if not n or not i: return messagebox.showwarning("Aviso", "El nombre y la IP son campos obligatorios.")

            puerto = puerto_seguro(p, None)
            if puerto is None:
                return messagebox.showwarning("Puerto invalido", "El puerto debe ser un numero entre 1 y 65535.")

            marca = "hikvision" if self.var_marca.get() == "Hikvision" else "zkteco"
            nuevo_dato = {"nombre": n, "ip": i, "port": puerto, "marca": marca}
            if marca == "hikvision":
                usuario = self.var_usuario.get().strip() or "admin"
                nuevo_dato["usuario"] = usuario
                nuevo_dato["password"] = self.var_password.get()
            idx = self.listbox.curselection()
            if idx:
                dato_actual = self.config_local["checadores"][idx[0]]
                horario_actual = dato_actual.get("horarios", self.config_local.get("horarios"))
                empleados_actuales = dato_actual.get("empleados", [])
                descargados_actuales = dato_actual.get("empleados_descargados", [])
                modo_actual = dato_actual.get("modo_conexion")
                if "modelo" in dato_actual:
                    nuevo_dato["modelo"] = dato_actual["modelo"]
                if modo_actual:
                    nuevo_dato["modo_conexion"] = modo_actual
                nuevo_dato["horarios"] = normalizar_horario(horario_actual, self.config_local.get("horarios"))
                nuevo_dato["empleados"] = empleados_actuales
                nuevo_dato["empleados_descargados"] = descargados_actuales
                self.config_local["checadores"][idx[0]] = nuevo_dato
            else:
                nuevo_dato["horarios"] = normalizar_horario(self.config_local.get("horarios"))
                nuevo_dato["empleados"] = []
                nuevo_dato["empleados_descargados"] = []
                self.config_local["checadores"].append(nuevo_dato)

            guardar_configuracion(self.config_local)
            self.actualizar_lista()
            actualizar_combos_globales()
            self.estado_campos("disabled")
            self.btn_modificar.config(text="Modificar Datos", bg=ACCENT)
            messagebox.showinfo("Guardado", "La configuración del dispositivo se ha actualizado correctamente.")

    def probar_conexion(self):
        ip, puerto = self.var_ip.get(), self.var_puerto.get()
        if not ip: return
        puerto = puerto_seguro(puerto, None)
        if puerto is None:
            return messagebox.showwarning("Puerto invalido", "El puerto debe ser un numero entre 1 y 65535.")
        conn = None
        try:
            conn = zk.connect()
            messagebox.showinfo("Conexión Exitosa", f"Se estableció comunicación con el equipo en la IP:\n{ip}")
        except Exception as e:
            messagebox.showerror("Error de Red", f"No se pudo conectar a {ip}.\nVerifique que el equipo esté encendido y en la misma red.\n\nDetalle: {e}")
        finally:
            if conn: conn.disconnect()

    def probar_conexion(self):
        ip, puerto = self.var_ip.get(), self.var_puerto.get()
        if not ip:
            return
        puerto = puerto_seguro(puerto, None)
        if puerto is None:
            return messagebox.showwarning("Puerto invalido", "El puerto debe ser un numero entre 1 y 65535.")
        conn = None
        try:
            marca = "hikvision" if self.var_marca.get() == "Hikvision" else "zkteco"
            if marca == "hikvision":
                info = {
                    "ip": ip,
                    "port": puerto,
                    "marca": "hikvision",
                    "usuario": self.var_usuario.get().strip() or "admin",
                    "password": self.var_password.get()
                }
                modo = hikvision_probar_conexion(info)
                messagebox.showinfo("Conexion exitosa", f"Hikvision respondio correctamente en:\n{ip}\n\nModo: {modo}")
            else:
                conn, modo = conectar_zk_robusto(ip, puerto)
                messagebox.showinfo("Conexion exitosa", f"Se establecio comunicacion con el equipo en la IP:\n{ip}\n\nModo: {modo}")
        except Exception as e:
            messagebox.showerror("Error de red", f"No se pudo conectar a {ip}.\nVerifique que el equipo este encendido y en la misma red.\n\nDetalle: {e}")
        finally:
            if conn:
                try:
                    conn.disconnect()
                except Exception:
                    pass

    def eliminar_equipo(self):
        idx = self.listbox.curselection()
        if not idx:
            return messagebox.showwarning("Selección requerida", "Por favor, selecciona un equipo de la lista de la izquierda primero.")
        
        equipo = self.config_local["checadores"][idx[0]]["nombre"]
        
        confirmar = messagebox.askyesno("Eliminar Dispositivo", f"¿Estás seguro que deseas eliminar permanentemente el checador '{equipo}' de la configuración?\n\nNota: Esto NO borrará tus archivos de Excel previamente generados.")
        
        if confirmar:
            self.config_local["checadores"].pop(idx[0])
            guardar_configuracion(self.config_local)
            self.actualizar_lista()
            actualizar_combos_globales()
            self.nuevo_registro()


class VentanaDiasNoLaborados(Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        aplicar_icono(self)
        self.title("Dias no laborados")
        self.configure(bg=BG_APP)
        self.resizable(True, True)
        self.config_local = cargar_configuracion()
        self.dias_actuales = []
        self.idx_actual = None
        self.var_motivo = tk.StringVar()
        self.cal = None
        self.init_ui()
        ajustar_ventana(self, 720, 640, 640, 560)

    def init_ui(self):
        header = tk.Frame(self, bg=BG_APP)
        header.pack(fill="x", padx=24, pady=(18, 10))
        tk.Label(header, text="Dias no laborados", bg=BG_APP, fg=FG_TEXT, font=("Segoe UI", 16, "bold")).pack(anchor="w")
        tk.Label(header, text="Estas fechas son globales y no generan falta en ningun checador.", bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 9)).pack(anchor="w", pady=(2, 0))

        form = tk.Frame(self, bg=BG_CARD, padx=20, pady=16, highlightthickness=1, highlightbackground=BORDER_CARD)
        form.pack(fill="x", padx=24, pady=(0, 12))

        tk.Label(form, text="Selecciona la fecha:", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 10)).pack(anchor="w")
        cal_style = {
            'background': BG_CARD, 'foreground': FG_TEXT, 'headersbackground': BG_APP, 'headersforeground': FG_DIM,
            'normalbackground': BG_CARD, 'normalforeground': FG_TEXT, 'weekendbackground': BG_CARD, 'weekendforeground': FG_DIM,
            'othermonthbackground': BG_APP, 'othermonthwebackground': BG_APP, 'othermonthforeground': "#334155",
            'othermonthweforeground': "#334155", 'selectbackground': ACCENT, 'selectforeground': FG_TEXT,
            'bottombackground': BG_APP, 'bordercolor': BORDER_CARD, 'borderwidth': 0, 'font': ("Segoe UI", 10)
        }
        self.cal = DateEntry(form, locale='es_MX', date_pattern='dd/MM/yyyy', **cal_style)
        self.cal.pack(fill="x", ipady=5, pady=(4, 12))

        tk.Label(form, text="Motivo * (obligatorio):", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 10)).pack(anchor="w")
        self.entry_motivo = tk.Entry(form, textvariable=self.var_motivo, bg=BG_CARD_ALT, fg=FG_TEXT, insertbackground=FG_TEXT, bd=0, font=("Segoe UI", 11))
        self.entry_motivo.pack(fill="x", ipady=7, pady=(4, 12))

        botones_form = tk.Frame(form, bg=BG_CARD)
        botones_form.pack(fill="x")
        tk.Button(botones_form, text="Agregar / actualizar fecha", command=self.agregar_actualizar, bg=SUCCESS, fg="white", font=("Segoe UI", 9, "bold"), bd=0, padx=14, pady=9, cursor="hand2").pack(side="left")
        tk.Button(botones_form, text="Limpiar", command=self.nuevo, bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 9), bd=0, padx=14, pady=9, cursor="hand2").pack(side="left", padx=(8, 0))

        lista_card = tk.Frame(self, bg=BG_CARD, padx=16, pady=14, highlightthickness=1, highlightbackground=BORDER_CARD)
        lista_card.pack(fill="both", expand=True, padx=24, pady=(0, 12))
        tk.Label(lista_card, text="Fechas guardadas", bg=BG_CARD, fg=FG_TEXT, font=("Segoe UI", 12, "bold")).pack(anchor="w", pady=(0, 8))
        self.lst_dias = tk.Listbox(lista_card, bg=BG_CARD_ALT, fg=FG_TEXT, bd=0, highlightthickness=0, selectbackground=ACCENT, font=("Segoe UI", 10), activestyle="none")
        self.lst_dias.pack(fill="both", expand=True)
        self.lst_dias.bind("<<ListboxSelect>>", self.cargar_seleccion)

        acciones = tk.Frame(self, bg=BG_APP)
        acciones.pack(fill="x", padx=24, pady=(0, 18))
        tk.Button(acciones, text="Eliminar seleccionado", command=self.eliminar_seleccionado, bg="#3B1F2A", fg=DANGER, font=("Segoe UI", 9, "bold"), bd=0, padx=14, pady=9, cursor="hand2").pack(side="left")
        tk.Button(acciones, text="Guardar cambios", command=self.guardar, bg=SUCCESS, fg="white", font=("Segoe UI", 10, "bold"), bd=0, padx=18, pady=9, cursor="hand2").pack(side="right")

        self.cargar_lista()

    def obtener_fecha_form(self):
        try:
            return self.cal.get_date()
        except Exception:
            return None

    def poner_fecha_form(self, fecha):
        try:
            self.cal.set_date(fecha)
        except Exception:
            pass

    def cargar_lista(self):
        self.dias_actuales = list(self.config_local.get("dias_no_laborados", []))
        self.lst_dias.delete(0, tk.END)
        for item in self.dias_actuales:
            fecha_txt = item.get("fecha", "")
            motivo = item.get("motivo", "")
            try:
                fecha = datetime.strptime(fecha_txt, "%Y-%m-%d").date()
                fecha_mostrar = fecha.strftime("%d/%m/%Y")
                dia_txt = dias_sem[fecha.weekday()]
            except Exception:
                fecha_mostrar = fecha_txt
                dia_txt = ""
            detalle = f"{fecha_mostrar} | {dia_txt}"
            if motivo:
                detalle = f"{detalle} | {motivo}"
            self.lst_dias.insert(tk.END, detalle)

    def nuevo(self):
        self.idx_actual = None
        self.var_motivo.set("")
        self.lst_dias.selection_clear(0, tk.END)

    def cargar_seleccion(self, event=None):
        idx = self.lst_dias.curselection()
        if not idx:
            return
        self.idx_actual = idx[0]
        item = self.dias_actuales[self.idx_actual]
        try:
            self.poner_fecha_form(datetime.strptime(item.get("fecha", ""), "%Y-%m-%d").date())
        except Exception:
            pass
        self.var_motivo.set(item.get("motivo", ""))

    def agregar_actualizar(self):
        fecha = self.obtener_fecha_form()
        if not fecha:
            return messagebox.showwarning("Fecha requerida", "Selecciona una fecha valida.")
        motivo = self.var_motivo.get().strip()
        if not motivo:
            messagebox.showwarning("Motivo requerido", "Escribe un motivo para el dia no laborado.\nEs obligatorio para poder guardar.")
            self.entry_motivo.focus()
            return
        fecha_txt = fecha.isoformat()
        nuevo = {"fecha": fecha_txt, "motivo": motivo}
        fechas = [item.get("fecha") for item in self.dias_actuales]
        if self.idx_actual is not None and 0 <= self.idx_actual < len(self.dias_actuales):
            self.dias_actuales[self.idx_actual] = nuevo
        elif fecha_txt in fechas:
            if not messagebox.askyesno("Fecha duplicada", f"La fecha {fecha.strftime('%d/%m/%Y')} ya esta registrada.\n¿Actualizar el motivo?"):
                return
            self.dias_actuales[fechas.index(fecha_txt)] = nuevo
        else:
            self.dias_actuales.append(nuevo)
        self.dias_actuales = sorted(self.dias_actuales, key=lambda item: item.get("fecha", ""))
        self.config_local["dias_no_laborados"] = self.dias_actuales
        self.cargar_lista()
        self.nuevo()
        messagebox.showinfo("Agregado", f"Fecha {fecha.strftime('%d/%m/%Y')} agregada correctamente.\n\nRecuerda presionar 'Guardar cambios' para persistir.")

    def eliminar_seleccionado(self):
        idx = self.lst_dias.curselection()
        if not idx:
            return messagebox.showwarning("Seleccion requerida", "Selecciona una fecha para eliminar.")
        self.dias_actuales.pop(idx[0])
        self.config_local["dias_no_laborados"] = self.dias_actuales
        self.cargar_lista()
        self.nuevo()

    def guardar(self):
        self.config_local["dias_no_laborados"] = self.dias_actuales
        guardar_configuracion(self.config_local)
        self.config_local = cargar_configuracion()
        self.cargar_lista()
        messagebox.showinfo("Guardado", "Dias no laborados guardados correctamente.")


class VentanaPersonasChecador(Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        aplicar_icono(self)
        self.title("Asignar Personas a Checadores")
        self.configure(bg=BG_APP)
        self.resizable(True, True)
        self.config_local = cargar_configuracion()
        self.empleados_reloj = []
        self.empleados_registrados = []
        self.empleados_eliminados = []
        self.disponibles_actuales = []
        self.registrados_actuales = []
        self.eliminados_actuales = []
        self.descargando = False
        self.verificando_eliminacion = False
        self.ventana_eliminados = None
        self.lst_eliminados_popup = None
        self.lbl_eliminados_popup = None
        self._checador_cargado = ""
        self._cargando_personas = False
        self.estado_inicial_registrados = set()
        self.movimientos_aprobados = {}
        self.cambios_pendientes = False

        self.var_checador = tk.StringVar()
        self.var_buscar_disponibles = tk.StringVar()
        self.var_buscar_registrados = tk.StringVar()
        self.init_ui()
        self.protocol("WM_DELETE_WINDOW", self.cerrar_ventana_personas)
        ajustar_ventana(self, 1080, 700, 940, 600)

    def init_ui(self):
        header = tk.Frame(self, bg=BG_APP)
        header.pack(fill="x", padx=24, pady=(18, 12))
        tk.Label(header, text="Asignar Personas a Checadores", bg=BG_APP, fg=FG_TEXT, font=("Segoe UI", 16, "bold")).pack(side="left")

        nombres = [c["nombre"] for c in self.config_local["checadores"]]
        self.cb_checador = ttk.Combobox(header, textvariable=self.var_checador, values=nombres, state="readonly", width=24, font=("Segoe UI", 10))
        self.cb_checador.pack(side="right")
        self.cb_checador.bind("<<ComboboxSelected>>", self.cargar_checador)
        self.cb_checador.set("Selecciona un checador")
        tk.Button(header, text="Eliminadas", command=self.abrir_eliminados, bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 8, "bold"), bd=0, padx=10, pady=6, cursor="hand2").pack(side="right", padx=(0, 8))

        ayuda = tk.Label(
            self,
            text=("Selecciona el checador destino. Puedes agregar o quitar varias personas. "
                  "Si alguien pertenece a otro checador, podrás decidir individualmente si deseas moverlo."),
            bg=BG_APP,
            fg=FG_DIM,
            font=("Segoe UI", 9),
            justify="left",
            wraplength=980,
        )
        ayuda.pack(fill="x", padx=24, pady=(0, 10))

        self.lbl_cambios_personas = tk.Label(
            self, text="Sin cambios pendientes", bg=BG_CARD, fg=SUCCESS,
            font=("Segoe UI", 9, "bold"), anchor="w", padx=12, pady=8,
            highlightthickness=1, highlightbackground=BORDER_CARD,
        )
        self.lbl_cambios_personas.pack(fill="x", padx=24, pady=(0, 10))

        cont = tk.Frame(self, bg=BG_APP)
        cont.pack(fill="both", expand=True, padx=24, pady=(0, 12))
        cont.grid_columnconfigure(0, weight=1, uniform="personas")
        cont.grid_columnconfigure(1, weight=0, minsize=150)
        cont.grid_columnconfigure(2, weight=1, uniform="personas")
        cont.grid_rowconfigure(0, weight=1)

        card_disp = tk.Frame(cont, bg=BG_CARD, padx=14, pady=14, highlightthickness=1, highlightbackground=BORDER_CARD)
        card_disp.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        tk.Label(card_disp, text="Personas disponibles", bg=BG_CARD, fg=FG_TEXT, font=("Segoe UI", 12, "bold")).pack(anchor="w")
        tk.Label(card_disp, text="Incluye personas libres y de otros checadores. Su asignación actual aparece junto al nombre.", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 8), wraplength=390, justify="left").pack(anchor="w", pady=(2, 8))
        tk.Entry(
            card_disp, textvariable=self.var_buscar_disponibles, bg=BG_CARD_ALT, fg=FG_TEXT,
            insertbackground=FG_TEXT, bd=0, font=("Segoe UI", 9),
        ).pack(fill="x", ipady=6, pady=(0, 7))
        acciones_disp = tk.Frame(card_disp, bg=BG_CARD)
        acciones_disp.pack(fill="x", pady=(0, 7))
        tk.Button(acciones_disp, text="Seleccionar todas", command=lambda: self.seleccionar_todas(self.lst_disponibles), bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 8, "bold"), bd=0, pady=5, cursor="hand2").pack(side="left", fill="x", expand=True, padx=(0, 3))
        tk.Button(acciones_disp, text="Limpiar selección", command=lambda: self.limpiar_seleccion(self.lst_disponibles), bg=BG_CARD_ALT, fg=FG_DIM, font=("Segoe UI", 8, "bold"), bd=0, pady=5, cursor="hand2").pack(side="left", fill="x", expand=True, padx=(3, 0))
        self.lst_disponibles = tk.Listbox(card_disp, selectmode=tk.EXTENDED, exportselection=False, bg=BG_CARD_ALT, fg=FG_TEXT, bd=0, highlightthickness=0, selectbackground=ACCENT, font=("Segoe UI", 10), activestyle="none")
        self.lst_disponibles.pack(fill="both", expand=True)
        self.lst_disponibles.bind("<Double-Button-1>", lambda _event: self.pasar_a_registrados())
        self.lst_disponibles.bind("<<ListboxSelect>>", self.actualizar_contadores_seleccion)
        self.lbl_disponibles = tk.Label(card_disp, text="Selecciona un checador para ver disponibles.", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 8))
        self.lbl_disponibles.pack(anchor="w", pady=(8, 0))

        centro = tk.Frame(cont, bg=BG_APP)
        centro.grid(row=0, column=1, sticky="ns", padx=6)
        centro.grid_rowconfigure(0, weight=1)
        centro.grid_rowconfigure(3, weight=1)
        tk.Button(centro, text="Añadir\nseleccionadas  →", command=self.pasar_a_registrados, bg=ACCENT, fg="white", font=("Segoe UI", 9, "bold"), bd=0, width=17, pady=9, cursor="hand2").grid(row=1, column=0, pady=(0, 10))
        tk.Button(centro, text="←  Quitar\nseleccionadas", command=self.quitar_de_registrados, bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 9, "bold"), bd=0, width=17, pady=9, cursor="hand2").grid(row=2, column=0)

        card_reg = tk.Frame(cont, bg=BG_CARD, padx=14, pady=14, highlightthickness=1, highlightbackground=BORDER_CARD)
        card_reg.grid(row=0, column=2, sticky="nsew", padx=(10, 0))
        tk.Label(card_reg, text="Personas en este checador", bg=BG_CARD, fg=FG_TEXT, font=("Segoe UI", 12, "bold")).pack(anchor="w")
        tk.Label(card_reg, text="Estas personas aparecerán en reportes y horarios de este checador después de guardar.", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 8), wraplength=390, justify="left").pack(anchor="w", pady=(2, 8))
        tk.Entry(
            card_reg, textvariable=self.var_buscar_registrados, bg=BG_CARD_ALT, fg=FG_TEXT,
            insertbackground=FG_TEXT, bd=0, font=("Segoe UI", 9),
        ).pack(fill="x", ipady=6, pady=(0, 7))
        acciones_reg = tk.Frame(card_reg, bg=BG_CARD)
        acciones_reg.pack(fill="x", pady=(0, 7))
        tk.Button(acciones_reg, text="Seleccionar todas", command=lambda: self.seleccionar_todas(self.lst_registrados), bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 8, "bold"), bd=0, pady=5, cursor="hand2").pack(side="left", fill="x", expand=True, padx=(0, 3))
        tk.Button(acciones_reg, text="Limpiar selección", command=lambda: self.limpiar_seleccion(self.lst_registrados), bg=BG_CARD_ALT, fg=FG_DIM, font=("Segoe UI", 8, "bold"), bd=0, pady=5, cursor="hand2").pack(side="left", fill="x", expand=True, padx=(3, 0))
        self.lst_registrados = tk.Listbox(card_reg, selectmode=tk.EXTENDED, exportselection=False, bg=BG_CARD_ALT, fg=FG_TEXT, bd=0, highlightthickness=0, selectbackground=ACCENT, font=("Segoe UI", 10), activestyle="none")
        self.lst_registrados.pack(fill="both", expand=True)
        self.lst_registrados.bind("<Double-Button-1>", lambda _event: self.quitar_de_registrados())
        self.lst_registrados.bind("<<ListboxSelect>>", self.actualizar_contadores_seleccion)
        self.lbl_registrados = tk.Label(card_reg, text="0 personas registradas.", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 8))
        self.lbl_registrados.pack(anchor="w", pady=(8, 0))

        acciones = tk.Frame(self, bg=BG_APP)
        acciones.pack(fill="x", padx=24, pady=(0, 20))
        self.btn_cargar = tk.Button(acciones, text="Cargar personas desde este reloj", command=self.cargar_personas_reloj, bg=BG_CARD_ALT, fg=SUCCESS, font=("Segoe UI", 10, "bold"), bd=0, padx=14, pady=9, cursor="hand2")
        self.btn_cargar.pack(side="left")
        self.btn_eliminar_personas = tk.Button(acciones, text="Eliminar seleccionadas del sistema…", command=self.solicitar_eliminar_personas, bg=BG_CARD_ALT, fg=DANGER, font=("Segoe UI", 9, "bold"), bd=0, padx=12, pady=9, cursor="hand2")
        self.btn_eliminar_personas.pack(side="left", padx=(8, 0))
        self.lbl_estado_personas = tk.Label(acciones, text="Selecciona un checador.", bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 9))
        self.lbl_estado_personas.pack(side="left", fill="x", expand=True, padx=12)
        self.btn_guardar_personas = tk.Button(acciones, text="Revisar y guardar cambios", command=self.guardar_personas, bg=SUCCESS, fg="white", font=("Segoe UI", 10, "bold"), bd=0, padx=18, pady=9, cursor="hand2", state="disabled")
        self.btn_guardar_personas.pack(side="right")

        self.var_buscar_disponibles.trace_add("write", lambda *_args: self.actualizar_listas())
        self.var_buscar_registrados.trace_add("write", lambda *_args: self.actualizar_listas())

        self.limpiar_seleccion_inicial()

    def obtener_checador_actual(self):
        nombre = self.var_checador.get()
        return next((c for c in self.config_local["checadores"] if c["nombre"] == nombre), None)

    def limpiar_seleccion_inicial(self):
        self._cargando_personas = True
        self.var_checador.set("")
        self.cb_checador.set("Selecciona un checador")
        self.empleados_reloj = []
        self.empleados_registrados = []
        self.empleados_eliminados = []
        self.disponibles_actuales = []
        self.registrados_actuales = []
        self.eliminados_actuales = []
        self.lst_disponibles.delete(0, tk.END)
        self.lst_registrados.delete(0, tk.END)
        self.actualizar_lista_eliminados_popup()
        self.lbl_disponibles.config(text="Selecciona un checador para ver disponibles del sistema.")
        self.lbl_registrados.config(text="0 personas registradas.")
        self.lbl_estado_personas.config(text="Selecciona un checador.", fg=FG_DIM)
        self._checador_cargado = ""
        self.estado_inicial_registrados = set()
        self.movimientos_aprobados = {}
        self._cargando_personas = False
        self.actualizar_estado_cambios()

    def obtener_asignaciones_personas(self):
        asignaciones = {}
        for checador in self.config_local["checadores"]:
            nombre_checador = checador.get("nombre", "Sin nombre")
            for empleado in checador.get("empleados", []):
                empleado = limpiar_nombre(empleado)
                if empleado:
                    asignaciones.setdefault(empleado, []).append(nombre_checador)
        return asignaciones

    def obtener_banco_personas(self):
        eliminadas = {
            limpiar_nombre(item.get("nombre", ""))
            for item in self.config_local.get("personas_eliminadas", [])
            if isinstance(item, dict)
        }
        personas = {
            limpiar_nombre(emp) for emp in self.config_local.get("personas_sistema", [])
            if str(emp).strip()
        }
        for checador in self.config_local.get("checadores", []):
            personas.update(
                limpiar_nombre(emp) for emp in checador.get("empleados", [])
                if str(emp).strip()
            )
            personas.update(
                limpiar_nombre(emp) for emp in checador.get("empleados_descargados", [])
                if str(emp).strip()
            )
        return sorted(emp for emp in personas if emp and emp not in eliminadas)

    def cargar_checador(self, event=None):
        nombre_nuevo = self.var_checador.get()
        if (
            self._checador_cargado and nombre_nuevo != self._checador_cargado
            and self.cambios_pendientes
            and not messagebox.askyesno(
                "Cambios sin guardar",
                "Hay cambios pendientes. ¿Cambiar de checador y descartarlos?",
                parent=self,
            )
        ):
            self._cargando_personas = True
            self.var_checador.set(self._checador_cargado)
            self._cargando_personas = False
            return
        checador = self.obtener_checador_actual()
        if not checador:
            return
        self._cargando_personas = True
        self.config_local = cargar_configuracion()
        checador = self.obtener_checador_actual()
        if not checador:
            self._cargando_personas = False
            return
        descargados_checador = checador.get("empleados_descargados", [])
        if not descargados_checador:
            descargados_checador = list(checador.get("empleados", []))
        self.empleados_reloj = list(descargados_checador)
        self.empleados_registrados = sorted({limpiar_nombre(emp) for emp in checador.get("empleados", []) if str(emp).strip()})
        self.empleados_eliminados = list(self.config_local.get("personas_eliminadas", []))
        self.estado_inicial_registrados = set(self.empleados_registrados)
        self.movimientos_aprobados = {}
        self._checador_cargado = checador["nombre"]
        self.lbl_estado_personas.config(text=f"Editando asignaciones de {checador['nombre']}.", fg=FG_DIM)
        self._cargando_personas = False
        self.actualizar_listas()
        self.actualizar_estado_cambios()

    def actualizar_listas(self):
        if not hasattr(self, "lst_disponibles"):
            return
        asignaciones = self.obtener_asignaciones_personas()
        registrados = sorted({limpiar_nombre(emp) for emp in self.empleados_registrados if str(emp).strip()})
        banco = self.obtener_banco_personas()
        filtro_disponibles = limpiar_nombre(self.var_buscar_disponibles.get()).lower()
        filtro_registrados = limpiar_nombre(self.var_buscar_registrados.get()).lower()
        disponibles = [
            emp for emp in banco
            if emp not in registrados and (not filtro_disponibles or filtro_disponibles in emp.lower())
        ]
        registrados_visibles = [
            emp for emp in registrados
            if not filtro_registrados or filtro_registrados in emp.lower()
        ]

        self.disponibles_actuales = disponibles
        self.registrados_actuales = registrados_visibles
        self.eliminados_actuales = list(self.config_local.get("personas_eliminadas", []))

        self.lst_disponibles.delete(0, tk.END)
        for empleado in disponibles:
            origenes = [nombre for nombre in asignaciones.get(empleado, []) if nombre != self.var_checador.get()]
            detalle = f"{empleado}  —  {', '.join(origenes)}" if origenes else empleado
            self.lst_disponibles.insert(tk.END, detalle)

        self.lst_registrados.delete(0, tk.END)
        for empleado in registrados_visibles:
            self.lst_registrados.insert(tk.END, empleado)

        self.actualizar_lista_eliminados_popup()

        seleccion_disp = len(self.lst_disponibles.curselection())
        seleccion_reg = len(self.lst_registrados.curselection())
        total_disponibles = len([emp for emp in banco if emp not in registrados])
        self.lbl_disponibles.config(text=f"{len(disponibles)} visibles de {total_disponibles} | {seleccion_disp} seleccionadas")
        self.lbl_registrados.config(text=f"{len(registrados_visibles)} visibles de {len(registrados)} | {seleccion_reg} seleccionadas")

    def seleccionar_todas(self, lista):
        if lista.size():
            lista.selection_set(0, tk.END)
        self.actualizar_contadores_seleccion()

    def limpiar_seleccion(self, lista):
        lista.selection_clear(0, tk.END)
        self.actualizar_contadores_seleccion()

    def actualizar_contadores_seleccion(self, event=None):
        if not self._checador_cargado:
            return
        seleccion_disp = len(self.lst_disponibles.curselection())
        seleccion_reg = len(self.lst_registrados.curselection())
        total_registrados = len(self.empleados_registrados)
        total_disponibles = len([emp for emp in self.obtener_banco_personas() if emp not in set(self.empleados_registrados)])
        self.lbl_disponibles.config(text=f"{len(self.disponibles_actuales)} visibles de {total_disponibles} | {seleccion_disp} seleccionadas")
        self.lbl_registrados.config(text=f"{len(self.registrados_actuales)} visibles de {total_registrados} | {seleccion_reg} seleccionadas")

    def actualizar_estado_cambios(self):
        actuales = {limpiar_nombre(emp) for emp in self.empleados_registrados if str(emp).strip()}
        agregadas = actuales - self.estado_inicial_registrados
        quitadas = self.estado_inicial_registrados - actuales
        self.cambios_pendientes = bool(agregadas or quitadas)
        if not hasattr(self, "lbl_cambios_personas"):
            return
        if self.cambios_pendientes:
            partes = []
            if agregadas:
                partes.append(f"{len(agregadas)} por agregar")
            if quitadas:
                partes.append(f"{len(quitadas)} por quitar")
            movimientos = len([emp for emp in agregadas if emp in self.movimientos_aprobados])
            if movimientos:
                partes.append(f"{movimientos} se moverán desde otro checador")
            self.lbl_cambios_personas.config(text="Cambios pendientes: " + " | ".join(partes), fg="#F59E0B")
            self.btn_guardar_personas.config(state="normal", bg=SUCCESS)
        else:
            self.lbl_cambios_personas.config(text="Sin cambios pendientes", fg=SUCCESS)
            self.btn_guardar_personas.config(state="disabled", bg="#334155")

    def cerrar_ventana_personas(self):
        if self.cambios_pendientes and not messagebox.askyesno(
            "Cambios sin guardar",
            "Hay cambios pendientes. ¿Cerrar y descartarlos?",
            parent=self,
        ):
            return
        self.destroy()

    def formato_eliminado(self, item):
        nombre = item.get("nombre", "")
        fecha = item.get("fecha", "")
        checador = item.get("checador", "")
        detalle = f"{nombre} | {checador}" if checador else nombre
        if fecha:
            detalle = f"{detalle} | {fecha[:10]}"
            fecha_eliminacion = _fecha_eliminacion_segura(fecha)
            if fecha_eliminacion:
                vence = fecha_eliminacion + timedelta(days=DIAS_RETENCION_PERSONAS_ELIMINADAS)
                detalle = f"{detalle} | purga {vence:%d/%m/%Y}"
        return detalle

    def actualizar_lista_eliminados_popup(self):
        lista = self.lst_eliminados_popup
        if not lista or not lista.winfo_exists():
            return
        lista.delete(0, tk.END)
        for item in self.eliminados_actuales:
            lista.insert(tk.END, self.formato_eliminado(item))
        if self.lbl_eliminados_popup and self.lbl_eliminados_popup.winfo_exists():
            total = len(self.eliminados_actuales)
            self.lbl_eliminados_popup.config(text=f"{total} persona{'s' if total != 1 else ''} en eliminación temporal")

    def abrir_eliminados(self):
        self.config_local = cargar_configuracion()
        self.eliminados_actuales = list(self.config_local.get("personas_eliminadas", []))
        if self.ventana_eliminados and self.ventana_eliminados.winfo_exists():
            self.actualizar_lista_eliminados_popup()
            self.ventana_eliminados.lift()
            return
        win = Toplevel(self)
        aplicar_icono(win)
        win.title("Personas eliminadas")
        win.configure(bg=BG_APP)
        win.resizable(True, True)
        self.ventana_eliminados = win
        tk.Label(win, text="Personas eliminadas", bg=BG_APP, fg=FG_TEXT, font=("Segoe UI", 14, "bold")).pack(anchor="w", padx=18, pady=(16, 4))
        tk.Label(
            win,
            text=("Puedes recuperar una selección o todas las personas. Después de 60 días se eliminan "
                  "automáticamente del sistema local. La eliminación definitiva no borra el historial ni el usuario físico del reloj."),
            bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 9), justify="left", wraplength=690,
        ).pack(anchor="w", padx=18, pady=(0, 6))
        self.lbl_eliminados_popup = tk.Label(win, text="", bg=BG_APP, fg=FG_TEXT, font=("Segoe UI", 9, "bold"))
        self.lbl_eliminados_popup.pack(anchor="w", padx=18, pady=(0, 8))
        self.lst_eliminados_popup = tk.Listbox(win, selectmode=tk.EXTENDED, exportselection=False, bg=BG_CARD_ALT, fg=FG_TEXT, bd=0, highlightthickness=0, selectbackground=ACCENT, font=("Segoe UI", 10), activestyle="none")
        self.lst_eliminados_popup.pack(fill="both", expand=True, padx=18, pady=(0, 12))
        acciones = tk.Frame(win, bg=BG_APP)
        acciones.pack(fill="x", padx=18, pady=(0, 16))
        tk.Button(acciones, text="Eliminar definitivamente seleccionadas", command=self.eliminar_definitivamente_seleccionadas, bg=DANGER, fg="white", font=("Segoe UI", 9, "bold"), bd=0, padx=12, pady=8, cursor="hand2").pack(side="left")
        tk.Button(acciones, text="Recuperar todas", command=self.recuperar_todas_personas, bg=ACCENT, fg="white", font=("Segoe UI", 9, "bold"), bd=0, padx=12, pady=8, cursor="hand2").pack(side="right")
        tk.Button(acciones, text="Recuperar seleccionadas", command=self.recuperar_persona, bg=SUCCESS, fg="white", font=("Segoe UI", 9, "bold"), bd=0, padx=12, pady=8, cursor="hand2").pack(side="right", padx=(0, 8))
        tk.Button(acciones, text="Cerrar", command=win.destroy, bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 9), bd=0, padx=12, pady=8, cursor="hand2").pack(side="right", padx=(0, 8))
        self.actualizar_lista_eliminados_popup()
        ajustar_ventana(win, 780, 430, 700, 360)

    def cargar_personas_reloj(self):
        checador = self.obtener_checador_actual()
        if self.descargando:
            return
        if not checador:
            return messagebox.showwarning("Seleccion requerida", "Primero selecciona un checador.")
        self.descargando = True
        self.btn_cargar.config(state="disabled", text="Cargando...", bg="#334155")
        self.lbl_estado_personas.config(text="Conectando al checador...", fg="#F59E0B")
        threading.Thread(target=self.worker_cargar_personas, args=(dict(checador),), daemon=True).start()

    def worker_cargar_personas(self, checador):
        conn = None
        try:
            if marca_checador(checador) == "hikvision":
                empleados = hikvision_obtener_personas(checador)
                if not empleados:
                    eventos = hikvision_obtener_eventos(checador, date.today() - timedelta(days=90), date.today())
                    empleados = sorted({limpiar_nombre(r["nombre"]) for r in eventos if r.get("nombre")})
                self.after(0, lambda: self.finalizar_carga_personas(empleados, None))
                return
            conn, _modo = conectar_zk_robusto(checador["ip"], checador.get("port"), modo_preferido=checador.get("modo_conexion"))
            users = conn.get_users()
            empleados = sorted({limpiar_nombre(u.name) for u in users if u.name and "admin" not in limpiar_nombre(u.name).lower()})
            self.after(0, lambda: self.finalizar_carga_personas(empleados, None))
        except Exception as e:
            self.after(0, lambda error=str(e): self.finalizar_carga_personas([], error))
        finally:
            if conn:
                try:
                    conn.disconnect()
                except Exception:
                    pass

    def finalizar_carga_personas(self, empleados, error):
        self.descargando = False
        self.btn_cargar.config(state="normal", text="Actualizar personas desde este reloj", bg=BG_CARD_ALT)
        if error:
            self.empleados_reloj = []
            self.actualizar_listas()
            self.lbl_estado_personas.config(text="Error al cargar personas.", fg=DANGER)
            return messagebox.showerror("Error", f"No se pudieron cargar personas del checador:\n{error}")
        checador = self.obtener_checador_actual()
        if not checador:
            return
        nombre_checador = checador.get("nombre", "")
        empleados_norm = sorted({limpiar_nombre(emp) for emp in empleados if str(emp).strip()})
        for c in self.config_local["checadores"]:
            if c.get("nombre") == nombre_checador:
                c["empleados_descargados"] = empleados_norm
                break
        guardar_configuracion(self.config_local)
        self.config_local = cargar_configuracion()
        for c in self.config_local["checadores"]:
            if c.get("nombre") == nombre_checador:
                self.empleados_reloj = list(c.get("empleados_descargados", empleados_norm))
                break
        else:
            self.empleados_reloj = list(empleados_norm)
        self.actualizar_listas()
        self.lbl_estado_personas.config(text=f"{len(empleados)} personas descargadas de '{nombre_checador}'. Usa las flechas y guarda.", fg=SUCCESS)

    def pasar_a_registrados(self):
        seleccion = list(self.lst_disponibles.curselection())
        if not seleccion:
            return messagebox.showwarning("Selección requerida", "Selecciona una o más personas para añadir.", parent=self)
        empleados = [self.disponibles_actuales[idx] for idx in seleccion]
        asignaciones = self.obtener_asignaciones_personas()
        nombre_actual = self.var_checador.get()
        conflictos = {
            empleado: [nombre for nombre in asignaciones.get(empleado, []) if nombre != nombre_actual]
            for empleado in empleados
        }
        conflictos = {empleado: origenes for empleado, origenes in conflictos.items() if origenes}
        aprobados = set(conflictos)
        if conflictos:
            aprobados = self.revisar_conflictos_movimiento(conflictos)
        for empleado in empleados:
            if empleado in conflictos and empleado not in aprobados:
                continue
            if empleado not in self.empleados_registrados:
                self.empleados_registrados.append(empleado)
            if empleado in conflictos:
                self.movimientos_aprobados[empleado] = list(conflictos[empleado])
        self.actualizar_listas()
        self.actualizar_estado_cambios()

    def revisar_conflictos_movimiento(self, conflictos):
        resultado = set()
        dialogo = Toplevel(self)
        aplicar_icono(dialogo)
        dialogo.title("Revisar personas asignadas a otros checadores")
        dialogo.configure(bg=BG_APP)
        dialogo.transient(self)
        dialogo.grab_set()

        tk.Label(dialogo, text="Revisar movimientos", bg=BG_APP, fg=FG_TEXT, font=("Segoe UI", 14, "bold")).pack(anchor="w", padx=20, pady=(18, 4))
        tk.Label(
            dialogo,
            text=("Estas personas ya pertenecen a otro checador. Selecciona únicamente las que deseas mover; "
                  "las demás conservarán su asignación actual."),
            bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 9), justify="left", wraplength=620,
        ).pack(anchor="w", padx=20, pady=(0, 10))

        empleados = list(conflictos)
        lista = tk.Listbox(
            dialogo, selectmode=tk.EXTENDED, bg=BG_CARD_ALT, fg=FG_TEXT, bd=0,
            highlightthickness=0, selectbackground=ACCENT, font=("Segoe UI", 10), activestyle="none",
        )
        lista.pack(fill="both", expand=True, padx=20, pady=(0, 10))
        for empleado in empleados:
            lista.insert(tk.END, f"{empleado}  —  actualmente en {', '.join(conflictos[empleado])}")

        seleccion = tk.Frame(dialogo, bg=BG_APP)
        seleccion.pack(fill="x", padx=20, pady=(0, 10))
        tk.Button(seleccion, text="Seleccionar todas", command=lambda: lista.selection_set(0, tk.END), bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 8, "bold"), bd=0, padx=10, pady=6, cursor="hand2").pack(side="left")
        tk.Button(seleccion, text="Limpiar selección", command=lambda: lista.selection_clear(0, tk.END), bg=BG_CARD_ALT, fg=FG_DIM, font=("Segoe UI", 8, "bold"), bd=0, padx=10, pady=6, cursor="hand2").pack(side="left", padx=(6, 0))

        def aceptar():
            resultado.update(empleados[idx] for idx in lista.curselection())
            dialogo.destroy()

        botones = tk.Frame(dialogo, bg=BG_APP)
        botones.pack(fill="x", padx=20, pady=(0, 18))
        tk.Button(botones, text="Cancelar", command=dialogo.destroy, bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 9, "bold"), bd=0, padx=14, pady=8, cursor="hand2").pack(side="right")
        tk.Button(botones, text="Mover seleccionadas", command=aceptar, bg=ACCENT, fg="white", font=("Segoe UI", 9, "bold"), bd=0, padx=14, pady=8, cursor="hand2").pack(side="right", padx=(0, 8))
        dialogo.protocol("WM_DELETE_WINDOW", dialogo.destroy)
        ajustar_ventana(dialogo, 680, 440, 560, 340)
        self.wait_window(dialogo)
        return resultado

    def quitar_de_registrados(self):
        seleccion = sorted(self.lst_registrados.curselection(), reverse=True)
        if not seleccion:
            return messagebox.showwarning("Selección requerida", "Selecciona una o más personas para quitar de este checador.", parent=self)
        for idx in seleccion:
            empleado = self.registrados_actuales[idx]
            self.empleados_registrados = [emp for emp in self.empleados_registrados if emp != empleado]
            self.movimientos_aprobados.pop(empleado, None)
        self.actualizar_listas()
        self.actualizar_estado_cambios()

    def solicitar_eliminar_personas(self):
        if self.verificando_eliminacion:
            return
        if self.cambios_pendientes:
            return messagebox.showwarning(
                "Cambios pendientes",
                "Guarda o descarta primero los cambios de asignación antes de eliminar personas del sistema.",
                parent=self,
            )
        empleados = {
            self.disponibles_actuales[idx] for idx in self.lst_disponibles.curselection()
        }
        empleados.update(
            self.registrados_actuales[idx] for idx in self.lst_registrados.curselection()
        )
        empleados = sorted(empleados)
        if not empleados:
            return messagebox.showwarning(
                "Selección requerida",
                "Selecciona una o más personas en cualquiera de las dos listas.",
                parent=self,
            )
        if not messagebox.askyesno(
            "Verificar eliminación",
            f"Se verificarán los registros recientes de {len(empleados)} persona(s).\n\n"
            "Eliminar significa moverlas a 'Personas eliminadas' y retirarlas de todos los checadores y horarios especiales. "
            "No se borrarán físicamente del dispositivo.\n\n¿Continuar?",
            parent=self,
        ):
            return
        self.verificando_eliminacion = True
        self.btn_eliminar_personas.config(state="disabled", text="Verificando...")
        self.lbl_estado_personas.config(text=f"Verificando 0 de {len(empleados)} personas...", fg="#F59E0B")
        configuracion = cargar_configuracion()
        nombre_actual = self.var_checador.get()
        threading.Thread(target=self.worker_validar_eliminaciones, args=(configuracion, empleados, nombre_actual), daemon=True).start()

    def worker_validar_eliminaciones(self, configuracion, empleados, nombre_actual):
        aptas = []
        bloqueadas = []
        errores = []
        for posicion, empleado in enumerate(empleados, start=1):
            checadores = [
                dict(checador) for checador in configuracion.get("checadores", [])
                if empleado in {limpiar_nombre(emp) for emp in checador.get("empleados", [])}
            ]
            if not checadores:
                actual = next((c for c in configuracion.get("checadores", []) if c.get("nombre") == nombre_actual), None)
                checadores = [dict(actual)] if actual else []
            tiene_registro = False
            error_persona = []
            for checador in checadores:
                try:
                    if self.persona_tiene_registro_ultima_semana(checador, empleado):
                        tiene_registro = True
                        break
                except Exception as exc:
                    error_persona.append(f"{checador.get('nombre', 'Checador')}: {exc}")
            if tiene_registro:
                bloqueadas.append(empleado)
            elif error_persona:
                errores.append((empleado, error_persona))
            else:
                aptas.append(empleado)
            self.after(0, lambda posicion=posicion, total=len(empleados): self.lbl_estado_personas.config(text=f"Verificando {posicion} de {total} personas...", fg="#F59E0B"))
        self.after(0, lambda: self.finalizar_validar_eliminaciones(aptas, bloqueadas, errores))

    def persona_tiene_registro_ultima_semana(self, checador, empleado):
        cutoff = datetime.now() - timedelta(days=7)
        empleado_norm = limpiar_nombre(empleado).lower()

        if datos_memoria_crudos is not None and not datos_memoria_crudos.empty and checador_memoria == checador.get("nombre", ""):
            base = preparar_base_asistencia(datos_memoria_crudos)
            if not base.empty:
                reciente = base[(base["nombre"].str.lower() == empleado_norm) & (base["fecha_hora"] >= cutoff)]
                return not reciente.empty

        conn = None
        try:
            if marca_checador(checador) == "hikvision":
                eventos = hikvision_obtener_eventos(checador, cutoff.date(), date.today())
                return any(limpiar_nombre(r.get("nombre", "")).lower() == empleado_norm and r.get("fecha_hora") >= cutoff for r in eventos)
            conn, _modo = conectar_zk_robusto(checador["ip"], checador.get("port"), modo_preferido=checador.get("modo_conexion"))
            users = conn.get_users()
            user_dict = {u.user_id: limpiar_nombre(u.name) for u in users}
            for att in conn.get_attendance():
                nombre = limpiar_nombre(user_dict.get(att.user_id, ""))
                if nombre.lower() == empleado_norm and att.timestamp >= cutoff:
                    return True
            return False
        finally:
            if conn:
                try:
                    conn.disconnect()
                except Exception:
                    pass

    def finalizar_validar_eliminaciones(self, aptas, bloqueadas, errores):
        self.verificando_eliminacion = False
        self.btn_eliminar_personas.config(state="normal", text="Eliminar seleccionadas del sistema…")
        lineas = [f"Aptas para eliminar: {len(aptas)}"]
        if bloqueadas:
            lineas.append(f"Bloqueadas por registros en los últimos 7 días: {len(bloqueadas)}")
            lineas.extend(f"  • {empleado}" for empleado in bloqueadas[:8])
            if len(bloqueadas) > 8:
                lineas.append(f"  • y {len(bloqueadas) - 8} más")
        if errores:
            lineas.append(f"No verificadas por error: {len(errores)}")
            lineas.extend(f"  • {empleado}: {detalles[0]}" for empleado, detalles in errores[:5])
            if len(errores) > 5:
                lineas.append(f"  • y {len(errores) - 5} más")
        if not aptas:
            self.lbl_estado_personas.config(text="No hay personas aptas para eliminar.", fg=DANGER)
            return messagebox.showwarning("Resultado de la verificación", "\n".join(lineas), parent=self)
        lineas.append("")
        lineas.append("¿Mover las personas aptas a 'Personas eliminadas'?")
        if not messagebox.askyesno("Confirmar eliminación", "\n".join(lineas), parent=self):
            self.lbl_estado_personas.config(text="Eliminación cancelada.", fg=FG_DIM)
            return
        self.eliminar_personas_sistema(aptas)

    def eliminar_personas_sistema(self, empleados):
        empleados = sorted({limpiar_nombre(empleado) for empleado in empleados if limpiar_nombre(empleado)})
        empleados_set = set(empleados)
        self.config_local = cargar_configuracion()
        self.config_local["personas_sistema"] = sorted({
            limpiar_nombre(emp) for emp in self.config_local.get("personas_sistema", [])
            if limpiar_nombre(emp) not in empleados_set
        })
        eliminadas = [
            item for item in self.config_local.get("personas_eliminadas", [])
            if limpiar_nombre(item.get("nombre", "")) not in empleados_set
        ]
        asignaciones = self.obtener_asignaciones_personas()
        fecha = datetime.now().strftime("%Y-%m-%d %H:%M")
        for empleado in empleados:
            eliminadas.append({"nombre": empleado, "fecha": fecha, "checador": ", ".join(asignaciones.get(empleado, []))})
        self.config_local["personas_eliminadas"] = eliminadas
        for checador in self.config_local["checadores"]:
            checador["empleados"] = sorted({
                limpiar_nombre(emp) for emp in checador.get("empleados", [])
                if limpiar_nombre(emp) not in empleados_set
            })
            horarios = normalizar_horario(checador.get("horarios"), self.config_local.get("horarios"))
            for especial in horarios.get("especiales", []):
                especial["empleados"] = sorted({
                    limpiar_nombre(emp) for emp in especial.get("empleados", [])
                    if limpiar_nombre(emp) not in empleados_set
                })
            checador["horarios"] = horarios
        guardar_configuracion(self.config_local)
        self.config_local = cargar_configuracion()
        self.empleados_registrados = [emp for emp in self.empleados_registrados if emp not in empleados_set]
        self.estado_inicial_registrados -= empleados_set
        for empleado in empleados:
            self.movimientos_aprobados.pop(empleado, None)
        self.cargar_checador()
        self.lbl_estado_personas.config(text=f"{len(empleados)} persona(s) se movieron a eliminadas.", fg=SUCCESS)
        if hasattr(self.master, "recargar_personas_registradas"):
            self.master.recargar_personas_registradas()

    def obtener_eliminados_seleccionados(self):
        if not self.lst_eliminados_popup or not self.lst_eliminados_popup.winfo_exists():
            return []
        return [
            self.eliminados_actuales[indice]
            for indice in self.lst_eliminados_popup.curselection()
            if 0 <= indice < len(self.eliminados_actuales)
        ]

    def refrescar_despues_gestion_eliminados(self, mensaje):
        self.config_local = cargar_configuracion()
        self.cargar_checador()
        self.eliminados_actuales = list(self.config_local.get("personas_eliminadas", []))
        self.actualizar_lista_eliminados_popup()
        self.lbl_estado_personas.config(text=mensaje, fg=SUCCESS)
        if hasattr(self.master, "recargar_personas_registradas"):
            self.master.recargar_personas_registradas()

    def recuperar_personas_eliminadas(self, items):
        empleados = sorted({
            limpiar_nombre(item.get("nombre", "")) for item in items
            if isinstance(item, dict) and limpiar_nombre(item.get("nombre", ""))
        })
        if not empleados:
            return 0
        empleados_clave = {empleado.casefold() for empleado in empleados}
        self.config_local = cargar_configuracion()
        personas = {limpiar_nombre(emp) for emp in self.config_local.get("personas_sistema", []) if str(emp).strip()}
        personas.update(empleados)
        self.config_local["personas_sistema"] = sorted(personas)
        self.config_local["personas_eliminadas"] = [
            item for item in self.config_local.get("personas_eliminadas", [])
            if limpiar_nombre(item.get("nombre", "")).casefold() not in empleados_clave
        ]
        guardar_configuracion(self.config_local)
        self.refrescar_despues_gestion_eliminados(f"{len(empleados)} persona(s) recuperada(s) al banco general.")
        return len(empleados)

    def recuperar_persona(self):
        seleccionadas = self.obtener_eliminados_seleccionados()
        if not seleccionadas:
            return messagebox.showwarning("Seleccion requerida", "Selecciona una persona eliminada para recuperarla.")
        self.recuperar_personas_eliminadas(seleccionadas)

    def recuperar_todas_personas(self):
        items = list(self.eliminados_actuales)
        if not items:
            return messagebox.showinfo("Sin personas", "No hay personas eliminadas para recuperar.")
        if not messagebox.askyesno(
            "Recuperar todas",
            f"¿Recuperar las {len(items)} personas al banco general?\n\nNo se reasignarán automáticamente a ningún checador.",
            parent=self.ventana_eliminados or self,
        ):
            return
        self.recuperar_personas_eliminadas(items)

    def eliminar_definitivamente_seleccionadas(self):
        seleccionadas = self.obtener_eliminados_seleccionados()
        if not seleccionadas:
            return messagebox.showwarning("Selección requerida", "Selecciona una o más personas para eliminarlas definitivamente.")
        nombres = [limpiar_nombre(item.get("nombre", "")) for item in seleccionadas]
        vista = "\n".join(f"• {nombre}" for nombre in nombres[:8])
        if len(nombres) > 8:
            vista += f"\n• ... y {len(nombres) - 8} más"
        if not messagebox.askyesno(
            "Eliminar definitivamente",
            (f"Esta acción no se puede deshacer y eliminará {len(nombres)} persona(s) del sistema local:\n\n"
             f"{vista}\n\nEl historial de asistencia y los usuarios físicos de los relojes no se borrarán."),
            icon="warning",
            parent=self.ventana_eliminados or self,
        ):
            return
        self.config_local = cargar_configuracion()
        eliminadas = eliminar_personas_definitivamente_config(self.config_local, nombres)
        guardar_configuracion(self.config_local)
        self.refrescar_despues_gestion_eliminados(f"{len(eliminadas)} persona(s) eliminada(s) definitivamente.")
        messagebox.showinfo(
            "Eliminación definitiva completada",
            f"Se eliminaron definitivamente {len(eliminadas)} persona(s) del sistema local.",
            parent=self.ventana_eliminados or self,
        )

    def guardar_personas(self):
        checador = self.obtener_checador_actual()
        if not checador:
            return messagebox.showwarning("Seleccion requerida", "Primero selecciona un checador.")
        nombre_actual = checador.get("nombre")
        seleccionados_set = {limpiar_nombre(emp) for emp in self.empleados_registrados if str(emp).strip()}
        agregadas = seleccionados_set - self.estado_inicial_registrados
        quitadas = self.estado_inicial_registrados - seleccionados_set
        movimientos = {empleado for empleado in agregadas if empleado in self.movimientos_aprobados}
        if not agregadas and not quitadas:
            return messagebox.showinfo("Sin cambios", "No hay cambios pendientes para guardar.", parent=self)

        resumen = [f"Checador destino: {nombre_actual}", ""]
        resumen.append(f"Personas por agregar: {len(agregadas)}")
        resumen.append(f"Personas por quitar: {len(quitadas)}")
        resumen.append(f"Personas que se moverán desde otro checador: {len(movimientos)}")
        if movimientos:
            resumen.append("")
            resumen.append("Movimientos confirmados:")
            for empleado in sorted(movimientos)[:10]:
                resumen.append(f"  • {empleado} desde {', '.join(self.movimientos_aprobados[empleado])}")
            if len(movimientos) > 10:
                resumen.append(f"  • y {len(movimientos) - 10} más")
        resumen.append("")
        resumen.append("¿Aplicar estos cambios?")
        if not messagebox.askyesno("Revisar cambios", "\n".join(resumen), parent=self):
            return

        seleccionados = sorted(seleccionados_set)
        self.config_local = cargar_configuracion()
        personas = {limpiar_nombre(emp) for emp in self.config_local.get("personas_sistema", []) if str(emp).strip()}
        personas.update(seleccionados)
        self.config_local["personas_sistema"] = sorted(personas)
        for item in self.config_local["checadores"]:
            if item.get("nombre") == nombre_actual:
                item["empleados"] = seleccionados
                retirar_de_especiales = quitadas
            else:
                item["empleados"] = sorted({
                    limpiar_nombre(emp) for emp in item.get("empleados", [])
                    if str(emp).strip() and limpiar_nombre(emp) not in movimientos
                })
                retirar_de_especiales = movimientos
            if retirar_de_especiales:
                horarios = normalizar_horario(item.get("horarios"), self.config_local.get("horarios"))
                for especial in horarios.get("especiales", []):
                    especial["empleados"] = sorted({
                        limpiar_nombre(emp) for emp in especial.get("empleados", [])
                        if limpiar_nombre(emp) not in retirar_de_especiales
                    })
                item["horarios"] = horarios
        guardar_configuracion(self.config_local)
        self.config_local = cargar_configuracion()
        self._cargando_personas = True
        self.var_checador.set(nombre_actual)
        self.cb_checador.set(nombre_actual)
        self._cargando_personas = False
        self.cargar_checador()
        if hasattr(self.master, "recargar_personas_registradas"):
            self.master.recargar_personas_registradas()
        messagebox.showinfo(
            "Cambios guardados",
            f"Se actualizaron las personas de {nombre_actual}.\n\n"
            f"Agregadas: {len(agregadas)}\nQuitadas: {len(quitadas)}\nMovidas: {len(movimientos)}",
            parent=self,
        )

class VentanaHorarios(Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        aplicar_icono(self)
        self.title("Horarios por Checador")
        self.configure(bg=BG_APP)
        self.resizable(True, True)
        self.config_local = cargar_configuracion()
        self.idx_especial_actual = None
        self.empleados_checador = []
        self.empleados_registrados = []
        self.empleados_especial = []
        self._cargando_horario = False
        self._checador_cargado = ""
        self.cambios_pendientes = False

        self.var_checador = tk.StringVar()
        self.var_entrada = tk.StringVar()
        self.var_salida = tk.StringVar()
        self.var_tolerancia = tk.StringVar()
        self.var_esp_nombre = tk.StringVar()
        self.var_esp_entrada = tk.StringVar()
        self.var_esp_salida = tk.StringVar()
        self.var_esp_tolerancia = tk.StringVar()
        self.var_empleado = tk.StringVar()
        self.var_empleado_reg = tk.StringVar()
        self.var_filtro_empleado = tk.StringVar()
        self.dias_vars = {}
        self.esp_dias_vars = {}
        self.fin_semana_vars = {}
        self.esp_fin_semana_vars = {}

        self.init_ui()
        self.configurar_seguimiento_cambios()
        self.protocol("WM_DELETE_WINDOW", self.cerrar_horarios)
        ajustar_ventana(self, 1060, 760, 920, 660)

    def init_ui(self):
        header = tk.Frame(self, bg=BG_APP)
        header.pack(fill="x", padx=24, pady=(18, 12))
        tk.Label(header, text="Horarios por Checador", bg=BG_APP, fg=FG_TEXT, font=("Segoe UI", 16, "bold")).pack(side="left")

        nombres = [c["nombre"] for c in self.config_local["checadores"]]
        self.cb_checador = ttk.Combobox(header, textvariable=self.var_checador, values=nombres, state="readonly", width=24, font=("Segoe UI", 10))
        self.cb_checador.pack(side="right")
        self.cb_checador.bind("<<ComboboxSelected>>", self.cargar_checador)
        self.cb_checador.set("Selecciona un checador")

        resumen = tk.Frame(self, bg=BG_CARD, padx=18, pady=11, highlightthickness=1, highlightbackground=BORDER_CARD)
        resumen.pack(fill="x", padx=24, pady=(0, 12))
        self.lbl_resumen_horario = tk.Label(
            resumen, text="Selecciona un checador para revisar su cobertura semanal.",
            bg=BG_CARD, fg=FG_TEXT, font=("Segoe UI", 10, "bold"),
            justify="left", anchor="w", wraplength=790,
        )
        self.lbl_resumen_horario.pack(side="left", fill="x", expand=True)
        self.lbl_cambios_horario = tk.Label(
            resumen, text="Sin cambios pendientes", bg=BG_CARD, fg=SUCCESS,
            font=("Segoe UI", 9, "bold"), padx=10,
        )
        self.lbl_cambios_horario.pack(side="right")

        cuerpo = tk.Frame(self, bg=BG_APP)
        cuerpo.pack(fill="both", expand=True, padx=24, pady=(0, 20))
        self.canvas_horarios = tk.Canvas(cuerpo, bg=BG_APP, highlightthickness=0)
        scroll_horarios = tk.Scrollbar(cuerpo, orient="vertical", command=self.canvas_horarios.yview, bg=BG_APP, troughcolor=BG_CARD)
        self.canvas_horarios.configure(yscrollcommand=scroll_horarios.set)
        scroll_horarios.pack(side="right", fill="y")
        self.canvas_horarios.pack(side="left", fill="both", expand=True)

        cont = tk.Frame(self.canvas_horarios, bg=BG_APP)
        self.scroll_window_horarios = self.canvas_horarios.create_window((0, 0), window=cont, anchor="nw")
        cont.bind("<Configure>", lambda event: self.canvas_horarios.configure(scrollregion=self.canvas_horarios.bbox("all")))
        self.canvas_horarios.bind("<Configure>", lambda event: self.canvas_horarios.itemconfig(self.scroll_window_horarios, width=event.width))
        self.canvas_horarios.bind("<Enter>", lambda event: self.canvas_horarios.bind_all("<MouseWheel>", self.scroll_horarios))
        self.canvas_horarios.bind("<Leave>", lambda event: self.canvas_horarios.unbind_all("<MouseWheel>"))

        cont.grid_columnconfigure(0, weight=1)
        cont.grid_columnconfigure(1, weight=1)
        cont.grid_rowconfigure(0, weight=1)

        card_general = tk.Frame(cont, bg=BG_CARD, padx=20, pady=18, highlightthickness=1, highlightbackground=BORDER_CARD)
        card_general.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        tk.Label(card_general, text="Horario General del Checador", bg=BG_CARD, fg=FG_TEXT, font=("Segoe UI", 13, "bold")).pack(anchor="w", pady=(0, 12))

        for texto, var, tipo in [("Entrada (HH:MM):", self.var_entrada, "hora"), ("Salida (HH:MM):", self.var_salida, "hora"), ("Tolerancia retardo (min):", self.var_tolerancia, "numero")]:
            fila = tk.Frame(card_general, bg=BG_CARD)
            fila.pack(fill="x", pady=(0, 10))
            tk.Label(fila, text=texto, bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 10)).pack(anchor="w")
            if tipo == "hora":
                ttk.Combobox(fila, textvariable=var, values=opciones_horas(), state="readonly", font=("Segoe UI", 11)).pack(fill="x", ipady=4, pady=(4, 0))
            else:
                tk.Spinbox(fila, from_=0, to=240, increment=1, textvariable=var, bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 11), bd=0, insertbackground=FG_TEXT).pack(fill="x", ipady=5, pady=(4, 0))

        tk.Label(card_general, text="Ejemplo claro: si la entrada es 08:30 y la tolerancia es 10 min, puede checar hasta 08:40 sin retardo. A las 08:41 ya cuenta 1 min de retardo.", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 8), wraplength=380, justify="left").pack(anchor="w", pady=(0, 8))
        tk.Label(card_general, text="Días Laborales:", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 10, "bold")).pack(anchor="w", pady=(4, 8))
        dias_frame = tk.Frame(card_general, bg=BG_CARD)
        dias_frame.pack(fill="x")
        for idx, nombre in dias_sem.items():
            var = tk.BooleanVar()
            self.dias_vars[idx] = var
            tk.Checkbutton(dias_frame, text=nombre[:3], variable=var, bg=BG_CARD, fg=FG_TEXT, selectcolor=BG_CARD_ALT, activebackground=BG_CARD, activeforeground=FG_TEXT, font=("Segoe UI", 9), bd=0).pack(side="left", padx=(0, 8))

        self.crear_controles_fin_semana(card_general, self.fin_semana_vars, "Horario distinto para guardias de fin de semana")
        self.lbl_validacion_general = tk.Label(
            card_general, text="", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 8, "bold"),
            justify="left", anchor="w", wraplength=390,
        )
        self.lbl_validacion_general.pack(fill="x", pady=(9, 0))
        tk.Button(card_general, text="Guardar Horario del Checador", command=self.guardar_horario_general, bg=SUCCESS, fg="white", font=("Segoe UI", 10, "bold"), bd=0, pady=9, cursor="hand2").pack(fill="x", pady=(16, 0))

        tk.Frame(card_general, bg=BORDER_CARD, height=1).pack(fill="x", pady=14)
        tk.Label(card_general, text="Personas registradas en este checador", bg=BG_CARD, fg=FG_TEXT, font=("Segoe UI", 11, "bold")).pack(anchor="w")
        tk.Label(card_general, text="Se administran desde Empleados > Asignar Personas. Esta lista alimenta los horarios especiales y reportes.", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 8), wraplength=380, justify="left").pack(anchor="w", pady=(2, 8))
        self.lst_registrados = tk.Listbox(card_general, height=5, bg=BG_CARD_ALT, fg=FG_TEXT, bd=0, highlightthickness=0, selectbackground=ACCENT, font=("Segoe UI", 9), activestyle="none")
        self.lst_registrados.pack(fill="both", expand=True)
        reg_actions = tk.Frame(card_general, bg=BG_CARD)
        reg_actions.pack(fill="x", pady=(6, 0))
        tk.Button(reg_actions, text="Actualizar lista", command=self.recargar_personas_registradas, bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 8, "bold"), bd=0, pady=5, cursor="hand2").pack(side="left", fill="x", expand=True, padx=(0, 5))
        tk.Button(reg_actions, text="Asignar personas", command=lambda: VentanaPersonasChecador(self), bg=ACCENT, fg="white", font=("Segoe UI", 8, "bold"), bd=0, pady=5, cursor="hand2").pack(side="right", fill="x", expand=True, padx=(5, 0))

        card_especial = tk.Frame(cont, bg=BG_CARD, padx=20, pady=18, highlightthickness=1, highlightbackground=BORDER_CARD)
        card_especial.grid(row=0, column=1, sticky="nsew", padx=(10, 0))
        tk.Label(card_especial, text="Horarios especiales por persona", bg=BG_CARD, fg=FG_TEXT, font=("Segoe UI", 13, "bold")).pack(anchor="w")
        tk.Label(card_especial, text="1. Selecciona un horario guardado o presiona Nuevo.", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 9)).pack(anchor="w", pady=(2, 8))

        top_especial = tk.Frame(card_especial, bg=BG_CARD)
        top_especial.pack(fill="x", pady=(0, 10))
        self.lst_especiales = tk.Listbox(top_especial, height=4, bg=BG_CARD_ALT, fg=FG_TEXT, bd=0, highlightthickness=0, selectbackground=ACCENT, font=("Segoe UI", 10), activestyle="none")
        self.lst_especiales.pack(side="left", fill="both", expand=True)
        self.lst_especiales.bind("<<ListboxSelect>>", self.cargar_especial)
        tk.Button(top_especial, text="Nuevo", command=lambda: self.nuevo_especial(confirmar=True), bg=SUCCESS, fg="white", font=("Segoe UI", 9, "bold"), bd=0, width=8, cursor="hand2").pack(side="right", fill="y", padx=(8, 0))

        tk.Label(card_especial, text="2. Define nombre, entrada, salida y minutos de tolerancia.", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(2, 6))
        for texto, var, tipo in [("Nombre del horario:", self.var_esp_nombre, "texto"), ("Entrada (HH:MM):", self.var_esp_entrada, "hora"), ("Salida (HH:MM):", self.var_esp_salida, "hora"), ("Tolerancia retardo (min):", self.var_esp_tolerancia, "numero")]:
            fila = tk.Frame(card_especial, bg=BG_CARD)
            fila.pack(fill="x", pady=(0, 8))
            tk.Label(fila, text=texto, bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 9)).pack(anchor="w")
            if tipo == "hora":
                ttk.Combobox(fila, textvariable=var, values=opciones_horas(), state="readonly", font=("Segoe UI", 10)).pack(fill="x", ipady=3, pady=(3, 0))
            elif tipo == "numero":
                tk.Spinbox(fila, from_=0, to=240, increment=1, textvariable=var, bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 10), bd=0, insertbackground=FG_TEXT).pack(fill="x", ipady=4, pady=(3, 0))
            else:
                tk.Entry(fila, textvariable=var, bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 10), bd=0, insertbackground=FG_TEXT).pack(fill="x", ipady=5, pady=(3, 0))

        tk.Button(
            card_especial, text="Copiar horario general del checador",
            command=self.copiar_horario_general_a_especial, bg=BG_CARD_ALT, fg=ACCENT,
            font=("Segoe UI", 8, "bold"), bd=0, pady=5, cursor="hand2",
        ).pack(fill="x", pady=(0, 7))

        tk.Label(card_especial, text="Ejemplo claro: si este horario entra 08:30 y tiene 10 min de tolerancia, 08:40 pasa normal; 08:41 queda con 1 min de retardo.", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 8), wraplength=380, justify="left").pack(anchor="w", pady=(0, 8))
        esp_dias_frame = tk.Frame(card_especial, bg=BG_CARD)
        esp_dias_frame.pack(fill="x", pady=(2, 8))
        for idx, nombre in dias_sem.items():
            var = tk.BooleanVar()
            self.esp_dias_vars[idx] = var
            tk.Checkbutton(esp_dias_frame, text=nombre[:3], variable=var, bg=BG_CARD, fg=FG_TEXT, selectcolor=BG_CARD_ALT, activebackground=BG_CARD, activeforeground=FG_TEXT, font=("Segoe UI", 8), bd=0).pack(side="left", padx=(0, 5))

        self.crear_controles_fin_semana(card_especial, self.esp_fin_semana_vars, "Horario distinto sabado/domingo para este horario")
        tk.Label(card_especial, text="3. Agrega personas registradas en este checador.", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(4, 6))
        filtro_frame = tk.Frame(card_especial, bg=BG_CARD)
        filtro_frame.pack(fill="x", pady=(0, 6))
        tk.Entry(
            filtro_frame, textvariable=self.var_filtro_empleado, bg=BG_CARD_ALT, fg=FG_TEXT,
            insertbackground=FG_TEXT, bd=0, font=("Segoe UI", 9),
        ).pack(side="left", fill="x", expand=True, ipady=5)
        emp_frame = tk.Frame(card_especial, bg=BG_CARD)
        emp_frame.pack(fill="x", pady=(0, 8))
        self.cb_empleado = ttk.Combobox(emp_frame, textvariable=self.var_empleado, values=[], font=("Segoe UI", 9))
        self.cb_empleado.pack(side="left", fill="x", expand=True)
        tk.Button(emp_frame, text="Actualizar", command=self.actualizar_empleados_combo, bg=BG_CARD_ALT, fg=SUCCESS, font=("Segoe UI", 8, "bold"), bd=0, padx=8, pady=4, cursor="hand2").pack(side="left", padx=(6, 0))
        tk.Button(emp_frame, text="Agregar al Horario", command=self.agregar_empleado_especial, bg=ACCENT, fg="white", font=("Segoe UI", 8, "bold"), bd=0, padx=8, pady=4, cursor="hand2").pack(side="left", padx=(6, 0))

        self.lst_empleados = tk.Listbox(card_especial, height=4, bg=BG_CARD_ALT, fg=FG_TEXT, bd=0, highlightthickness=0, selectbackground=ACCENT, font=("Segoe UI", 9), activestyle="none")
        self.lst_empleados.pack(fill="both", expand=True)
        self.lbl_conteo_empleados_especial = tk.Label(
            card_especial, text="0 personas asignadas", bg=BG_CARD, fg=FG_DIM,
            font=("Segoe UI", 8, "bold"), anchor="w",
        )
        self.lbl_conteo_empleados_especial.pack(fill="x", pady=(5, 0))
        quitar_frame = tk.Frame(card_especial, bg=BG_CARD)
        quitar_frame.pack(fill="x", pady=(6, 0))
        tk.Button(quitar_frame, text="Quitar seleccionada", command=self.quitar_empleado_especial, bg=BG_CARD_ALT, fg=DANGER, font=("Segoe UI", 9, "bold"), bd=0, pady=5, cursor="hand2").pack(side="left", fill="x", expand=True, padx=(0, 4))
        tk.Button(quitar_frame, text="Quitar todas", command=self.quitar_todos_empleados_especial, bg=BG_CARD_ALT, fg=DANGER, font=("Segoe UI", 9, "bold"), bd=0, pady=5, cursor="hand2").pack(side="left", fill="x", expand=True, padx=(4, 0))

        self.lbl_validacion_especial = tk.Label(
            card_especial, text="", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 8, "bold"),
            justify="left", anchor="w", wraplength=390,
        )
        self.lbl_validacion_especial.pack(fill="x", pady=(8, 0))

        acciones = tk.Frame(card_especial, bg=BG_CARD)
        acciones.pack(fill="x", pady=(10, 0))
        tk.Button(acciones, text="Guardar Horario Especial", command=self.guardar_especial, bg=SUCCESS, fg="white", font=("Segoe UI", 9, "bold"), bd=0, pady=8, cursor="hand2").pack(side="left", fill="x", expand=True, padx=(0, 5))
        tk.Button(acciones, text="Eliminar Especial", command=self.eliminar_especial, bg="#3B1F2A", fg=DANGER, font=("Segoe UI", 9, "bold"), bd=0, pady=8, cursor="hand2").pack(side="right", fill="x", expand=True, padx=(5, 0))

        self.limpiar_horarios()

    def configurar_seguimiento_cambios(self):
        generales = [self.var_entrada, self.var_salida, self.var_tolerancia]
        especiales = [self.var_esp_nombre, self.var_esp_entrada, self.var_esp_salida, self.var_esp_tolerancia]
        for var in generales:
            var.trace_add("write", self.marcar_cambio_general)
        for var in especiales:
            var.trace_add("write", self.marcar_cambio_especial)
        self.var_filtro_empleado.trace_add("write", lambda *_args: self.actualizar_empleados_combo())
        for var in self.dias_vars.values():
            var.trace_add("write", self.marcar_cambio_general)
        for var in self.esp_dias_vars.values():
            var.trace_add("write", self.marcar_cambio_especial)
        for datos in self.fin_semana_vars.values():
            for var in datos.values():
                var.trace_add("write", self.marcar_cambio_general)
        for datos in self.esp_fin_semana_vars.values():
            for var in datos.values():
                var.trace_add("write", self.marcar_cambio_especial)

    def marcar_cambio_general(self, *_args):
        if self._cargando_horario:
            return
        self.cambios_pendientes = True
        self.cambios_general = True
        self.actualizar_estado_horarios()

    def marcar_cambio_especial(self, *_args):
        if self._cargando_horario:
            return
        self.cambios_pendientes = True
        self.cambios_especial = True
        self.actualizar_estado_horarios()

    def limpiar_estado_cambios(self, general=True, especial=True):
        if general:
            self.cambios_general = False
        if especial:
            self.cambios_especial = False
        self.cambios_pendientes = bool(
            getattr(self, "cambios_general", False) or getattr(self, "cambios_especial", False)
        )
        self.actualizar_estado_horarios()

    def cerrar_horarios(self):
        if self.cambios_pendientes and not messagebox.askyesno(
            "Cambios sin guardar",
            "Hay cambios pendientes en los horarios. ¿Cerrar y descartarlos?",
            parent=self,
        ):
            return
        self.destroy()

    @staticmethod
    def _resumen_dias(dias):
        dias = sorted(set(dias))
        if dias == [0, 1, 2, 3, 4]:
            return "lunes a viernes"
        if dias == [0, 1, 2, 3, 4, 5]:
            return "lunes a sábado"
        if dias == list(range(7)):
            return "todos los días"
        return ", ".join(dias_sem[dia].lower() for dia in dias) if dias else "sin días laborales"

    def actualizar_estado_horarios(self):
        if not hasattr(self, "lbl_resumen_horario"):
            return
        checador = self.obtener_checador_actual()
        if not checador:
            resumen = "Selecciona un checador para revisar su cobertura semanal."
        else:
            dias = [idx for idx, var in self.dias_vars.items() if var.get()]
            especiales = checador.get("horarios", {}).get("especiales", [])
            asignados = {
                limpiar_nombre(emp)
                for especial in especiales
                for emp in especial.get("empleados", [])
                if str(emp).strip()
            }
            resumen = (
                f"{checador['nombre']}: {self._resumen_dias(dias)} · "
                f"{self.var_entrada.get() or '--:--'}–{self.var_salida.get() or '--:--'} · "
                f"tolerancia {self.var_tolerancia.get() or '0'} min · "
                f"{len(especiales)} especiales · {len(asignados)} personas con excepción"
            )
        self.lbl_resumen_horario.config(text=resumen)
        pendiente = bool(getattr(self, "cambios_general", False) or getattr(self, "cambios_especial", False))
        self.lbl_cambios_horario.config(
            text="Cambios pendientes" if pendiente else "Sin cambios pendientes",
            fg="#F59E0B" if pendiente else SUCCESS,
        )
        self.actualizar_validaciones_visibles()

    def actualizar_validaciones_visibles(self):
        if not hasattr(self, "lbl_validacion_general"):
            return
        dias = [idx for idx, var in self.dias_vars.items() if var.get()]
        entrada = self.var_entrada.get().strip()
        salida = self.var_salida.get().strip()
        if not self.obtener_checador_actual():
            texto_general, color_general = "Selecciona un checador para comenzar.", FG_DIM
        elif not dias:
            texto_general, color_general = "⚠ Selecciona al menos un día laboral.", "#F59E0B"
        elif entrada and entrada == salida:
            texto_general, color_general = "⚠ Entrada y salida no pueden ser iguales.", DANGER
        elif entrada and salida and salida < entrada:
            texto_general, color_general = "ℹ La salida es menor que la entrada; se interpretará como turno nocturno.", "#F59E0B"
        else:
            texto_general, color_general = "✓ Horario base listo para guardar.", SUCCESS
        self.lbl_validacion_general.config(text=texto_general, fg=color_general)

        if not self.var_esp_nombre.get().strip():
            texto_especial, color_especial = "Escribe un nombre y asigna personas para crear una excepción.", FG_DIM
        elif not any(var.get() for var in self.esp_dias_vars.values()):
            texto_especial, color_especial = "⚠ Selecciona al menos un día para este horario.", "#F59E0B"
        elif self.var_esp_entrada.get() == self.var_esp_salida.get():
            texto_especial, color_especial = "⚠ Entrada y salida no pueden ser iguales.", DANGER
        elif not self.empleados_especial:
            texto_especial, color_especial = "⚠ El horario todavía no tiene personas asignadas.", "#F59E0B"
        else:
            conflictos = self.obtener_conflictos_especiales()
            if conflictos:
                texto_especial = "⚠ Ya tienen otro horario especial: " + ", ".join(conflictos[:3])
                if len(conflictos) > 3:
                    texto_especial += f" y {len(conflictos) - 3} más"
                color_especial = DANGER
            else:
                texto_especial, color_especial = "✓ Excepción lista para guardar.", SUCCESS
        self.lbl_validacion_especial.config(text=texto_especial, fg=color_especial)

    def limpiar_horarios(self):
        self._cargando_horario = True
        self.var_checador.set("")
        self.cb_checador.set("Selecciona un checador")
        for var in [self.var_entrada, self.var_salida, self.var_tolerancia, self.var_esp_nombre, self.var_esp_entrada, self.var_esp_salida, self.var_esp_tolerancia, self.var_empleado, self.var_empleado_reg]:
            var.set("")
        for var in self.dias_vars.values():
            var.set(False)
        for var in self.esp_dias_vars.values():
            var.set(False)
        self.cargar_dias_horarios_vars(self.fin_semana_vars, {}, {"entrada": "09:00", "salida": "13:00", "tolerancia_min": 10})
        self.cargar_dias_horarios_vars(self.esp_fin_semana_vars, {}, {"entrada": "09:00", "salida": "13:00", "tolerancia_min": 10})
        self.empleados_registrados = []
        self.empleados_especial = []
        self.idx_especial_actual = None
        self.lst_registrados.delete(0, tk.END)
        self.lst_especiales.delete(0, tk.END)
        self.lst_empleados.delete(0, tk.END)
        self.actualizar_empleados_combo()
        self._cargando_horario = False
        self._checador_cargado = ""
        self.cambios_general = False
        self.cambios_especial = False
        self.cambios_pendientes = False
        if hasattr(self, "lbl_resumen_horario"):
            self.actualizar_estado_horarios()

    def crear_controles_fin_semana(self, parent, destino, titulo):
        tk.Label(parent, text=titulo, bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(10, 4))
        tk.Label(parent, text="Si el dia no esta marcado como laboral, no genera falta; si hay checada, usa este horario para evaluar retardo.", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 8), wraplength=380, justify="left").pack(anchor="w", pady=(0, 4))
        for dia in [5, 6]:
            fila = tk.Frame(parent, bg=BG_CARD)
            fila.pack(fill="x", pady=(2, 2))
            activo = tk.BooleanVar()
            entrada = tk.StringVar(value="09:00" if dia == 5 else "09:00")
            salida = tk.StringVar(value="13:00" if dia == 5 else "13:00")
            tolerancia = tk.StringVar(value="10")
            destino[dia] = {"activo": activo, "entrada": entrada, "salida": salida, "tolerancia": tolerancia}
            tk.Checkbutton(fila, text=dias_sem[dia], variable=activo, bg=BG_CARD, fg=FG_TEXT, selectcolor=BG_CARD_ALT, activebackground=BG_CARD, activeforeground=FG_TEXT, font=("Segoe UI", 8), bd=0).pack(side="left", padx=(0, 6))
            ttk.Combobox(fila, textvariable=entrada, values=opciones_horas(), state="readonly", width=6, font=("Segoe UI", 8)).pack(side="left", padx=2)
            ttk.Combobox(fila, textvariable=salida, values=opciones_horas(), state="readonly", width=6, font=("Segoe UI", 8)).pack(side="left", padx=2)
            tk.Spinbox(fila, from_=0, to=240, increment=1, textvariable=tolerancia, width=4, bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 8), bd=0).pack(side="left", padx=2)

    def cargar_dias_horarios_vars(self, vars_dict, dias_horarios, base_horario):
        dias_horarios = normalizar_dias_horarios(dias_horarios, base_horario)
        for dia, vars_dia in vars_dict.items():
            datos = dias_horarios.get(str(dia), {})
            vars_dia["activo"].set(bool(datos.get("activo", False)))
            vars_dia["entrada"].set(datos.get("entrada", "09:00"))
            vars_dia["salida"].set(datos.get("salida", "13:00"))
            vars_dia["tolerancia"].set(str(datos.get("tolerancia_min", base_horario.get("tolerancia_min", 10))))

    def obtener_dias_horarios_vars(self, vars_dict):
        resultado = {}
        for dia, vars_dia in vars_dict.items():
            if not vars_dia["activo"].get():
                continue
            tolerancia = entero_no_negativo(vars_dia["tolerancia"].get(), 10)
            resultado[str(dia)] = {
                "activo": True,
                "entrada": vars_dia["entrada"].get().strip(),
                "salida": vars_dia["salida"].get().strip(),
                "tolerancia_min": tolerancia
            }
        return resultado

    def scroll_horarios(self, event):
        if hasattr(self, "canvas_horarios"):
            self.canvas_horarios.yview_scroll(int(-1 * (event.delta / 120)), "units")

    def obtener_checador_actual(self):
        nombre = self.var_checador.get()
        return next((c for c in self.config_local["checadores"] if c["nombre"] == nombre), None)

    def cargar_checador(self, event=None):
        nombre_nuevo = self.var_checador.get()
        if (
            self._checador_cargado and nombre_nuevo != self._checador_cargado
            and self.cambios_pendientes
            and not messagebox.askyesno(
                "Cambios sin guardar",
                "Hay cambios pendientes. ¿Cambiar de checador y descartarlos?",
                parent=self,
            )
        ):
            self.var_checador.set(self._checador_cargado)
            return
        checador = self.obtener_checador_actual()
        if not checador:
            return
        self._cargando_horario = True
        horario = normalizar_horario(checador.get("horarios"), self.config_local.get("horarios"))
        checador["horarios"] = horario
        self.var_entrada.set(horario["entrada"])
        self.var_salida.set(horario["salida"])
        self.var_tolerancia.set(str(horario["tolerancia_min"]))
        dias = set(horario.get("dias_laborales", []))
        for idx, var in self.dias_vars.items():
            var.set(idx in dias)
        self.cargar_dias_horarios_vars(self.fin_semana_vars, horario.get("dias_horarios", {}), horario)
        self.empleados_registrados = sorted({limpiar_nombre(emp) for emp in checador.get("empleados", []) if str(emp).strip()})
        self.actualizar_lista_empleados_checador()
        self.actualizar_empleados_combo()
        self.actualizar_lista_especiales()
        self.nuevo_especial(confirmar=False)
        self._cargando_horario = False
        self._checador_cargado = checador["nombre"]
        self.limpiar_estado_cambios()

    def recargar_personas_registradas(self):
        self.config_local = cargar_configuracion()
        self.cargar_checador()

    def validar_horario_form(self, entrada, salida, tolerancia, dias):
        try:
            datetime.strptime(entrada, "%H:%M")
            datetime.strptime(salida, "%H:%M")
            tolerancia = int(tolerancia)
        except Exception:
            messagebox.showwarning("Datos inválidos", "Usa formato HH:MM para horas y número entero para tolerancia.")
            return None
        if tolerancia < 0:
            messagebox.showwarning("Datos invalidos", "La tolerancia no puede ser negativa.")
            return None
        if not dias:
            messagebox.showwarning("Días requeridos", "Selecciona al menos un día laboral.")
            return None
        if entrada == salida:
            messagebox.showwarning("Horario inválido", "La entrada y la salida no pueden ser iguales.")
            return None
        return tolerancia

    def guardar_horario_general(self):
        checador = self.obtener_checador_actual()
        if not checador:
            return messagebox.showwarning("Seleccion requerida", "Primero selecciona un checador.")
        if getattr(self, "cambios_especial", False):
            return messagebox.showwarning(
                "Horario especial pendiente",
                "Guarda o descarta primero los cambios del horario especial para no perderlos.",
                parent=self,
            )
        dias = [idx for idx, var in self.dias_vars.items() if var.get()]
        tolerancia = self.validar_horario_form(self.var_entrada.get().strip(), self.var_salida.get().strip(), self.var_tolerancia.get().strip(), dias)
        if tolerancia is None:
            return
        horario_actual = normalizar_horario(checador.get("horarios"), self.config_local.get("horarios"))
        checador["horarios"] = {
            "entrada": self.var_entrada.get().strip(),
            "salida": self.var_salida.get().strip(),
            "tolerancia_min": tolerancia,
            "dias_laborales": dias,
            "dias_horarios": self.obtener_dias_horarios_vars(self.fin_semana_vars),
            "especiales": horario_actual.get("especiales", [])
        }
        guardar_configuracion(self.config_local)
        self.config_local = cargar_configuracion()
        messagebox.showinfo("Guardado", f"Horario general de {checador['nombre']} guardado correctamente.")
        self.cargar_checador()

    def actualizar_lista_especiales(self):
        self.lst_especiales.delete(0, tk.END)
        checador = self.obtener_checador_actual()
        if not checador:
            return
        for especial in checador.get("horarios", {}).get("especiales", []):
            total_emp = len(especial.get("empleados", []))
            self.lst_especiales.insert(tk.END, f"{especial.get('nombre', 'Horario especial')} ({total_emp})")

    def nuevo_especial(self, confirmar=False):
        checador_actual = self.obtener_checador_actual()
        if not checador_actual:
            return
        if confirmar and getattr(self, "cambios_especial", False) and not messagebox.askyesno(
            "Cambios sin guardar",
            "El horario especial actual tiene cambios pendientes. ¿Descartarlos y crear uno nuevo?",
            parent=self,
        ):
            return
        cargando_anterior = self._cargando_horario
        self._cargando_horario = True
        horario = normalizar_horario(checador_actual.get("horarios"))
        self.idx_especial_actual = None
        self.lst_especiales.selection_clear(0, tk.END)
        self.var_esp_nombre.set("")
        self.var_esp_entrada.set(horario["entrada"])
        self.var_esp_salida.set(horario["salida"])
        self.var_esp_tolerancia.set(str(horario["tolerancia_min"]))
        for idx, var in self.esp_dias_vars.items():
            var.set(idx in horario.get("dias_laborales", []))
        self.cargar_dias_horarios_vars(self.esp_fin_semana_vars, {}, horario)
        self.empleados_especial = []
        self.actualizar_lista_empleados_especial()
        self._cargando_horario = cargando_anterior
        if not cargando_anterior:
            self.limpiar_estado_cambios(general=False, especial=True)

    def copiar_horario_general_a_especial(self):
        if not self.obtener_checador_actual():
            return messagebox.showwarning("Selección requerida", "Primero selecciona un checador.", parent=self)
        self.var_esp_entrada.set(self.var_entrada.get())
        self.var_esp_salida.set(self.var_salida.get())
        self.var_esp_tolerancia.set(self.var_tolerancia.get())
        for idx, var in self.esp_dias_vars.items():
            var.set(self.dias_vars[idx].get())
        self.cargar_dias_horarios_vars(
            self.esp_fin_semana_vars,
            self.obtener_dias_horarios_vars(self.fin_semana_vars),
            {
                "entrada": self.var_entrada.get(),
                "salida": self.var_salida.get(),
                "tolerancia_min": entero_no_negativo(self.var_tolerancia.get(), 10),
            },
        )
        self.marcar_cambio_especial()

    def cargar_especial(self, event=None):
        idx = self.lst_especiales.curselection()
        checador = self.obtener_checador_actual()
        if not idx or not checador:
            return
        if getattr(self, "cambios_especial", False) and not messagebox.askyesno(
            "Cambios sin guardar",
            "El horario especial actual tiene cambios pendientes. ¿Descartarlos y abrir el seleccionado?",
            parent=self,
        ):
            self.lst_especiales.selection_clear(0, tk.END)
            if self.idx_especial_actual is not None:
                self.lst_especiales.selection_set(self.idx_especial_actual)
            return
        self._cargando_horario = True
        self.idx_especial_actual = idx[0]
        especial = checador["horarios"]["especiales"][self.idx_especial_actual]
        self.var_esp_nombre.set(especial.get("nombre", ""))
        self.var_esp_entrada.set(especial.get("entrada", checador["horarios"]["entrada"]))
        self.var_esp_salida.set(especial.get("salida", checador["horarios"]["salida"]))
        self.var_esp_tolerancia.set(str(especial.get("tolerancia_min", checador["horarios"]["tolerancia_min"])))
        dias = set(especial.get("dias_laborales", checador["horarios"]["dias_laborales"]))
        for idx_dia, var in self.esp_dias_vars.items():
            var.set(idx_dia in dias)
        self.cargar_dias_horarios_vars(self.esp_fin_semana_vars, especial.get("dias_horarios", {}), especial)
        registrados = set(self.empleados_registrados)
        self.empleados_especial = [emp for emp in especial.get("empleados", []) if emp in registrados]
        self.actualizar_lista_empleados_especial()
        self._cargando_horario = False
        self.limpiar_estado_cambios(general=False, especial=True)

    def obtener_conflictos_especiales(self):
        checador = self.obtener_checador_actual()
        if not checador:
            return []
        asignados_otro = set()
        for idx, especial in enumerate(checador.get("horarios", {}).get("especiales", [])):
            if idx == self.idx_especial_actual:
                continue
            asignados_otro.update(
                limpiar_nombre(emp) for emp in especial.get("empleados", []) if str(emp).strip()
            )
        return sorted(set(self.empleados_especial) & asignados_otro)

    def guardar_especial(self):
        checador = self.obtener_checador_actual()
        if not checador:
            return messagebox.showwarning("Seleccion requerida", "Primero selecciona un checador.")
        nombre = self.var_esp_nombre.get().strip()
        if not nombre:
            return messagebox.showwarning("Nombre requerido", "Escribe un nombre para el horario especial.")
        dias = [idx for idx, var in self.esp_dias_vars.items() if var.get()]
        tolerancia = self.validar_horario_form(self.var_esp_entrada.get().strip(), self.var_esp_salida.get().strip(), self.var_esp_tolerancia.get().strip(), dias)
        if tolerancia is None:
            return
        if not self.empleados_especial:
            return messagebox.showwarning(
                "Personas requeridas",
                "Asigna al menos una persona a este horario especial.",
                parent=self,
            )
        conflictos = self.obtener_conflictos_especiales()
        if conflictos:
            return messagebox.showwarning(
                "Asignación duplicada",
                "Estas personas ya pertenecen a otro horario especial:\n\n" + "\n".join(conflictos),
                parent=self,
            )
        especial = {
            "nombre": nombre,
            "entrada": self.var_esp_entrada.get().strip(),
            "salida": self.var_esp_salida.get().strip(),
            "tolerancia_min": tolerancia,
            "dias_laborales": dias,
            "dias_horarios": self.obtener_dias_horarios_vars(self.esp_fin_semana_vars),
            "empleados": sorted({limpiar_nombre(emp) for emp in self.empleados_especial if str(emp).strip() and limpiar_nombre(emp) in self.empleados_registrados})
        }
        checador["horarios"] = normalizar_horario(checador.get("horarios"), self.config_local.get("horarios"))
        if self.idx_especial_actual is None:
            checador["horarios"]["especiales"].append(especial)
        else:
            checador["horarios"]["especiales"][self.idx_especial_actual] = especial
        guardar_configuracion(self.config_local)
        self.config_local = cargar_configuracion()
        self.actualizar_lista_especiales()
        messagebox.showinfo("Guardado", "Horario especial guardado correctamente.")
        self.nuevo_especial(confirmar=False)

    def eliminar_especial(self):
        checador = self.obtener_checador_actual()
        if not checador or self.idx_especial_actual is None:
            return messagebox.showwarning("Selección requerida", "Selecciona un horario especial para eliminar.")
        nombre = checador["horarios"]["especiales"][self.idx_especial_actual].get("nombre", "Horario especial")
        if not messagebox.askyesno("Eliminar Horario", f"¿Eliminar el horario especial '{nombre}'?"):
            return
        checador["horarios"]["especiales"].pop(self.idx_especial_actual)
        guardar_configuracion(self.config_local)
        self.config_local = cargar_configuracion()
        self.actualizar_lista_especiales()
        self.nuevo_especial(confirmar=False)

    def cargar_empleados_checador(self):
        checador = self.obtener_checador_actual()
        if not checador:
            return
        self.cb_empleado.config(values=["Cargando empleados..."])
        threading.Thread(target=self.worker_cargar_empleados, args=(checador,), daemon=True).start()

    def worker_cargar_empleados(self, checador):
        conn = None
        try:
            if marca_checador(checador) == "hikvision":
                empleados = hikvision_obtener_personas(checador)
                self.after(0, lambda: self.finalizar_carga_empleados(empleados, None))
                return
            conn, _modo = conectar_zk_robusto(checador["ip"], checador.get("port"), modo_preferido=checador.get("modo_conexion"))
            users = conn.get_users()
            empleados = sorted({limpiar_nombre(u.name) for u in users if u.name and "admin" not in limpiar_nombre(u.name).lower()})
            self.after(0, lambda: self.finalizar_carga_empleados(empleados, None))
        except Exception as e:
            self.after(0, lambda error=str(e): self.finalizar_carga_empleados([], error))
        finally:
            if conn:
                try:
                    conn.disconnect()
                except Exception:
                    pass

    def finalizar_carga_empleados(self, empleados, error):
        if error:
            self.cb_empleado_reg.config(values=[])
            return messagebox.showerror("Error", f"No se pudieron cargar empleados del checador:\n{error}")
        self.empleados_checador = empleados
        self.cb_empleado_reg.config(values=empleados)
        if empleados:
            self.var_empleado_reg.set(empleados[0])
        messagebox.showinfo("Empleados cargados", f"Se cargaron {len(empleados)} empleados del checador seleccionado.")

    def agregar_empleado_checador(self):
        empleado = limpiar_nombre(self.var_empleado_reg.get())
        if not empleado:
            return
        if empleado not in self.empleados_registrados:
            self.empleados_registrados.append(empleado)
            self.empleados_registrados.sort()
            self.actualizar_lista_empleados_checador()
            self.actualizar_empleados_combo()
        self.var_empleado_reg.set("")

    def quitar_empleado_checador(self):
        idx = self.lst_registrados.curselection()
        if not idx:
            return
        empleado = self.empleados_registrados.pop(idx[0])
        self.empleados_especial = [emp for emp in self.empleados_especial if emp != empleado]
        self.actualizar_lista_empleados_checador()
        self.actualizar_lista_empleados_especial()
        self.actualizar_empleados_combo()

    def guardar_empleados_checador(self):
        checador = self.obtener_checador_actual()
        if not checador:
            return
        checador["empleados"] = sorted({limpiar_nombre(emp) for emp in self.empleados_registrados if str(emp).strip()})
        guardar_configuracion(self.config_local)
        self.config_local = cargar_configuracion()
        messagebox.showinfo("Guardado", f"Lista de empleados de {checador['nombre']} guardada correctamente.")
        self.cargar_checador()

    def actualizar_lista_empleados_checador(self):
        self.lst_registrados.delete(0, tk.END)
        for empleado in self.empleados_registrados:
            self.lst_registrados.insert(tk.END, empleado)

    def actualizar_empleados_combo(self):
        filtro = limpiar_nombre(self.var_filtro_empleado.get()).lower()
        disponibles = [
            emp for emp in self.empleados_registrados
            if emp not in self.empleados_especial and (not filtro or filtro in emp.lower())
        ]
        self.cb_empleado.config(values=disponibles)
        if self.var_empleado.get() not in disponibles:
            self.var_empleado.set(disponibles[0] if disponibles else "")

    def agregar_empleado_especial(self):
        empleado = limpiar_nombre(self.var_empleado.get())
        if not empleado:
            return
        if empleado not in self.empleados_registrados:
            return messagebox.showwarning("Empleado no registrado", "Primero agrega este empleado a la lista del checador.")
        if empleado not in self.empleados_especial:
            self.empleados_especial.append(empleado)
            self.empleados_especial.sort()
            self.actualizar_lista_empleados_especial()
            self.marcar_cambio_especial()
        self.var_empleado.set("")

    def quitar_empleado_especial(self):
        idx = self.lst_empleados.curselection()
        if not idx:
            return
        self.empleados_especial.pop(idx[0])
        self.actualizar_lista_empleados_especial()
        self.marcar_cambio_especial()

    def quitar_todos_empleados_especial(self):
        if not self.empleados_especial:
            return
        if not messagebox.askyesno(
            "Quitar personas",
            "¿Quitar a todas las personas de este horario especial?",
            parent=self,
        ):
            return
        self.empleados_especial = []
        self.actualizar_lista_empleados_especial()
        self.marcar_cambio_especial()

    def actualizar_lista_empleados_especial(self):
        self.lst_empleados.delete(0, tk.END)
        for empleado in self.empleados_especial:
            self.lst_empleados.insert(tk.END, empleado)
        if hasattr(self, "lbl_conteo_empleados_especial"):
            total = len(self.empleados_especial)
            self.lbl_conteo_empleados_especial.config(
                text=f"{total} persona{'s' if total != 1 else ''} asignada{'s' if total != 1 else ''}"
            )
        if hasattr(self, "cb_empleado"):
            self.actualizar_empleados_combo()
        if hasattr(self, "lbl_validacion_especial"):
            self.actualizar_validaciones_visibles()


class VentanaHistorico(Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        aplicar_icono(self)
        self.title("Visor Avanzado de Empleados")
        self.configure(bg=BG_APP)
        self.resizable(True, True)
        
        self.df_dispositivo = None
        pd = cargar_pandas()
        self.df_historico_actual = pd.DataFrame()
        self.descargando_historico = False
        self.cancelar_descarga_historico = threading.Event()
        self.config_local = cargar_configuracion()
        self.nombres_completos = []
        self.usar_primera_ultima_actual = False
        self.vista_historico = tk.StringVar(value="resumen")
        self.fechas_historico_aplicadas = False
        self.init_ui()
        ajustar_ventana(self, 1000, 640, 940, 580)

    def init_ui(self):
        # HEADER CONTROLES (Estilo Panel de Control)
        header_frame = tk.Frame(self, bg=BG_APP)
        header_frame.pack(fill="x", padx=20, pady=(18, 12))
        
        # Título del panel
        titulo_bar = tk.Frame(header_frame, bg=BG_APP)
        titulo_bar.pack(fill="x", pady=(0, 10))
        tk.Label(titulo_bar, text="Panel de Búsqueda Histórica", bg=BG_APP, fg=FG_TEXT, font=("Segoe UI", 15, "bold")).pack(side="left")
        self.lbl_historico_estado = tk.Label(titulo_bar, text="Sin datos cargados", bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 9, "bold"))
        self.lbl_historico_estado.pack(side="right", pady=(4, 0))

        controles = tk.Frame(header_frame, bg=BG_APP)
        controles.pack(fill="x")
        for col in range(3):
            controles.grid_columnconfigure(col, weight=1, uniform="historial")

        # Fila de controles usando bloques compactos
        # Columna 1: Origen
        tk.Label(controles, text="1. Descargar datos", bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 9, "bold")).grid(row=0, column=0, sticky="w", pady=(0, 6))
        f_origen = tk.Frame(controles, bg=BG_CARD, padx=12, pady=10, highlightthickness=1, highlightbackground=BORDER_CARD)
        f_origen.grid(row=1, column=0, sticky="ew", padx=(0, 8))
        acciones_descarga = tk.Frame(f_origen, bg=BG_CARD)
        acciones_descarga.pack(side="right")
        self.btn_extraer = tk.Button(acciones_descarga, text="Extraer", command=self.extraer_empleados, bg=ACCENT, fg="white", font=("Segoe UI", 10, "bold"), bd=0, width=11, pady=4, cursor="hand2")
        self.btn_extraer.pack(side="left")
        self.btn_cancelar_descarga = tk.Button(acciones_descarga, text="Cancelar", command=self.cancelar_extraccion_historica, state="disabled", bg=DANGER, fg="white", font=("Segoe UI", 10, "bold"), bd=0, width=10, pady=4, cursor="hand2")
        self.btn_cancelar_descarga.pack(side="left", padx=(6, 0))
        selector_checador = tk.Frame(f_origen, bg=BG_CARD)
        selector_checador.pack(side="left", fill="x", expand=True, padx=(0, 8))
        self.cb_checadores = ttk.Combobox(selector_checador, values=[c["nombre"] for c in self.config_local["checadores"]], state="readonly", width=17, font=("Segoe UI", 10))
        self.cb_checadores.pack(fill="x", expand=True, ipady=2)

        # Columna 2: Búsqueda
        tk.Label(controles, text="2. Buscar empleado", bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 9, "bold")).grid(row=0, column=1, sticky="w", pady=(0, 6), padx=8)
        f_buscar = tk.Frame(controles, bg=BG_CARD, padx=12, pady=10, highlightthickness=1, highlightbackground=BORDER_CARD)
        f_buscar.grid(row=1, column=1, sticky="ew", padx=8)
        self.cb_empleados = ttk.Combobox(f_buscar, width=25, font=("Segoe UI", 10))
        self.cb_empleados.pack(side="left", fill="x", expand=True, ipady=2)
        self.cb_empleados.bind("<KeyRelease>", self.filtrar_teclado)
        self.cb_empleados.bind("<<ComboboxSelected>>", self.cargar_todo_empleado)

        # Columna 3: Filtro Fechas
        tk.Label(controles, text="3. Rango de fechas", bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 9, "bold")).grid(row=0, column=2, sticky="w", pady=(0, 6), padx=(8, 0))
        f_fechas = tk.Frame(controles, bg=BG_CARD, padx=12, pady=10, highlightthickness=1, highlightbackground=BORDER_CARD)
        f_fechas.grid(row=1, column=2, sticky="ew", padx=(8, 0))
        calendar_style = standard_calendar_style(
            BG_CARD, FG_TEXT, BG_APP, FG_DIM, BORDER_CARD, SUCCESS,
        )
        self.cal_desde = DateEntry(
            f_fechas, locale='es_MX', date_pattern='dd/MM/yyyy', width=10,
            **calendar_style,
        )
        self.cal_desde.pack(side="left", fill="x", expand=True)
        tk.Label(f_fechas, text="-", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 12)).pack(side="left", padx=5)
        self.cal_hasta = DateEntry(
            f_fechas, locale='es_MX', date_pattern='dd/MM/yyyy', width=10,
            **calendar_style,
        )
        self.cal_hasta.pack(side="left", fill="x", expand=True)
        tk.Button(f_fechas, text="Aplicar", command=self.filtrar_fechas, bg=SUCCESS, fg="white", font=("Segoe UI", 10, "bold"), bd=0, padx=10, pady=4, cursor="hand2").pack(side="left", padx=(8, 0))

        hoy = date.today()
        self.cal_desde.set_date(hoy.replace(day=1))
        self.cal_hasta.set_date(hoy)

        self.progreso_historico = Progressbar(header_frame, mode="determinate", maximum=100, value=0)
        self.progreso_historico.pack(fill="x", pady=(10, 0))

        acciones_hist = tk.Frame(header_frame, bg=BG_APP)
        acciones_hist.pack(fill="x", pady=(10, 0))
        tk.Button(acciones_hist, text="Ver Todos", command=self.ver_todos_empleados, bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 9, "bold"), bd=0, padx=14, pady=7, cursor="hand2").pack(side="left", padx=(0, 8))
        tk.Button(acciones_hist, text="Limpiar Tabla", command=self.limpiar_tabla, bg=BG_CARD_ALT, fg=FG_DIM, font=("Segoe UI", 9, "bold"), bd=0, padx=14, pady=7, cursor="hand2").pack(side="left", padx=(0, 8))
        tk.Label(acciones_hist, text="Vista:", bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 9, "bold")).pack(side="left", padx=(18, 6))
        self.btn_vista_resumen = tk.Button(acciones_hist, text="Resumen diario", command=self.mostrar_resumen, bg=ACCENT, fg="white", font=("Segoe UI", 9, "bold"), bd=0, padx=12, pady=7, cursor="hand2")
        self.btn_vista_resumen.pack(side="left", padx=(0, 6))
        self.btn_vista_movimientos = tk.Button(acciones_hist, text="Movimientos detallados", command=self.mostrar_movimientos, bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 9, "bold"), bd=0, padx=12, pady=7, cursor="hand2")
        self.btn_vista_movimientos.pack(side="left")

        # TABLA DE RESULTADOS
        f_tabla = tk.Frame(self, bg=BG_APP)
        f_tabla.pack(fill="both", expand=True, padx=20, pady=(0, 18))

        # Estilo Premium para Treeview
        style = ttk.Style()
        style.theme_use("clam")
        style.configure("Custom.Treeview", background=BG_CARD, fieldbackground=BG_CARD, foreground=FG_TEXT, borderwidth=0, font=("Segoe UI", 10), rowheight=32)
        style.configure("Custom.Treeview.Heading", background=BG_CARD_ALT, foreground=FG_TEXT, font=("Segoe UI", 11, "bold"), borderwidth=0, padding=5)
        style.map("Custom.Treeview", background=[('selected', ACCENT)], foreground=[('selected', "white")])

        self.tree = ttk.Treeview(f_tabla, style="Custom.Treeview", columns=("Nombre", "Fecha", "Dia", "Entrada", "Salida"), show="headings")
        self.tree.heading("Nombre", text="Nombre del Empleado")
        self.tree.heading("Fecha", text="Fecha")
        self.tree.heading("Dia", text="Día")
        self.tree.heading("Entrada", text="Hora Entrada")
        self.tree.heading("Salida", text="Hora Salida")

        self.tree.column("Nombre", width=280, anchor="w")
        self.tree.column("Fecha", width=115, anchor="center")
        self.tree.column("Dia", width=115, anchor="center")
        self.tree.column("Entrada", width=140, anchor="center")
        self.tree.column("Salida", width=140, anchor="center")

        # Scrollbar estilizado
        scroll = tk.Scrollbar(f_tabla, orient="vertical", command=self.tree.yview, bg=BG_APP, troughcolor=BG_CARD)
        self.tree.configure(yscrollcommand=scroll.set)
        scroll.pack(side="right", fill="y")
        self.tree.pack(side="left", fill="both", expand=True)

        self.tree_movimientos = ttk.Treeview(
            f_tabla,
            style="Custom.Treeview",
            columns=("Nombre", "Fecha", "Dia", "Hora"),
            show="headings"
        )
        self.tree_movimientos.heading("Nombre", text="Nombre del Empleado")
        self.tree_movimientos.heading("Fecha", text="Fecha")
        self.tree_movimientos.heading("Dia", text="Dia")
        self.tree_movimientos.heading("Hora", text="Hora de Checada")
        self.tree_movimientos.column("Nombre", width=340, anchor="w")
        self.tree_movimientos.column("Fecha", width=150, anchor="center")
        self.tree_movimientos.column("Dia", width=150, anchor="center")
        self.tree_movimientos.column("Hora", width=180, anchor="center")
        scroll_movimientos = tk.Scrollbar(f_tabla, orient="vertical", command=self.tree_movimientos.yview, bg=BG_APP, troughcolor=BG_CARD)
        self.tree_movimientos.configure(yscrollcommand=scroll_movimientos.set)
        self.scroll_resumen = scroll
        self.scroll_movimientos = scroll_movimientos

    def actualizar_botones_vista(self):
        es_movimientos = self.vista_historico.get() == "movimientos"
        self.btn_vista_resumen.config(
            bg=BG_CARD_ALT if es_movimientos else ACCENT,
            fg=FG_TEXT if es_movimientos else "white"
        )
        self.btn_vista_movimientos.config(
            bg=ACCENT if es_movimientos else BG_CARD_ALT,
            fg="white" if es_movimientos else FG_TEXT
        )

    def mostrar_resumen(self):
        self.vista_historico.set("resumen")
        self.actualizar_botones_vista()
        self.scroll_movimientos.pack_forget()
        self.tree_movimientos.pack_forget()
        self.scroll_resumen.pack(side="right", fill="y")
        self.tree.pack(side="left", fill="both", expand=True)
        self.procesar_y_mostrar(aplicar_fechas=self.fechas_historico_aplicadas)

    def mostrar_movimientos(self):
        self.vista_historico.set("movimientos")
        self.actualizar_botones_vista()
        self.scroll_resumen.pack_forget()
        self.tree.pack_forget()
        self.scroll_movimientos.pack(side="right", fill="y")
        self.tree_movimientos.pack(side="left", fill="both", expand=True)
        self.procesar_y_mostrar(aplicar_fechas=self.fechas_historico_aplicadas)

    def extraer_empleados(self):
        checador_sel = self.cb_checadores.get()
        if not checador_sel:
            return
        info = next((c for c in self.config_local["checadores"] if c["nombre"] == checador_sel), None)
        if not info or self.descargando_historico:
            return
        try:
            fecha_inicio = self.cal_desde.get_date()
            fecha_fin = self.cal_hasta.get_date()
        except Exception:
            return messagebox.showwarning("Rango de fechas", "Selecciona un rango de fechas válido para la descarga.")
        if fecha_inicio > fecha_fin:
            return messagebox.showwarning("Rango de fechas", "La fecha inicial no puede ser posterior a la fecha final.")

        self.usar_primera_ultima_actual = usa_primera_ultima_checador(info)
        self.fechas_historico_aplicadas = True
        self.descargando_historico = True
        self.cancelar_descarga_historico.clear()
        self.btn_extraer.config(state="disabled", text="Descargando...", bg="#334155")
        self.btn_cancelar_descarga.config(state="normal")
        self.actualizar_progreso_historico(
            0,
            f"Preparando descarga de {info['nombre']} del {fecha_inicio:%d/%m/%Y} al {fecha_fin:%d/%m/%Y}…",
        )
        threading.Thread(
            target=self.worker_extraer_empleados,
            args=(info, fecha_inicio, fecha_fin),
            daemon=True,
        ).start()

    def cancelar_extraccion_historica(self):
        if not self.descargando_historico:
            return
        self.cancelar_descarga_historico.set()
        self.btn_cancelar_descarga.config(state="disabled")
        self.lbl_historico_estado.config(text="Cancelando descarga al terminar la consulta actual…", fg="#F59E0B")

    def actualizar_progreso_historico(self, valor, texto):
        pagina_sin_total = texto.startswith("Descargando página") and " de " not in texto
        if pagina_sin_total:
            if str(self.progreso_historico.cget("mode")) != "indeterminate":
                self.progreso_historico.config(mode="indeterminate")
                self.progreso_historico.start(10)
        elif str(self.progreso_historico.cget("mode")) != "determinate":
            self.progreso_historico.stop()
            self.progreso_historico.config(mode="determinate")
        self.progreso_historico.config(value=max(0, min(100, int(valor))))
        self.lbl_historico_estado.config(text=texto, fg="#F59E0B")

    def worker_extraer_empleados(self, info, fecha_inicio, fecha_fin):
        conn = None
        try:
            if marca_checador(info) == "hikvision":
                def actualizar_estado_hikvision(texto, valor):
                    self.after(
                        0,
                        lambda texto=texto, valor=valor: self.actualizar_progreso_historico(valor, texto),
                    )

                records = hikvision_obtener_eventos(
                    info,
                    fecha_inicio,
                    fecha_fin,
                    progreso=actualizar_estado_hikvision,
                    cancelado=self.cancelar_descarga_historico,
                )
                if self.cancelar_descarga_historico.is_set():
                    self.after(0, lambda: self.finalizar_extraer_empleados([], None, cancelada=True))
                    return
                self.after(0, lambda: self.finalizar_extraer_empleados(records, None))
                return

            if self.cancelar_descarga_historico.is_set():
                self.after(0, lambda: self.finalizar_extraer_empleados([], None, cancelada=True))
                return
            conn, _modo = conectar_zk_robusto(
                info["ip"], info.get("port"), modo_preferido=info.get("modo_conexion"),
            )
            users = conn.get_users()
            atts = conn.get_attendance()
            user_dict = {u.user_id: limpiar_nombre(u.name) for u in users}
            records = []
            for a in atts:
                nombre = user_dict.get(a.user_id, "Desconocido")
                if "admin" not in nombre.lower():
                    records.append({"nombre": nombre, "fecha_hora": a.timestamp})
            if self.cancelar_descarga_historico.is_set():
                self.after(0, lambda: self.finalizar_extraer_empleados([], None, cancelada=True))
                return
            self.after(0, lambda: self.finalizar_extraer_empleados(records, None))
        except hikvision_connector.HikvisionDescargaCancelada:
            self.after(0, lambda: self.finalizar_extraer_empleados([], None, cancelada=True))
        except Exception as e:
            self.after(0, lambda error=str(e): self.finalizar_extraer_empleados([], error))
        finally:
            if conn:
                try:
                    conn.disconnect()
                except Exception:
                    pass

    def finalizar_extraer_empleados(self, records, error, cancelada=False, persist=True):
        self.descargando_historico = False
        self.btn_extraer.config(state="normal", text="Extraer", bg=ACCENT)
        self.btn_cancelar_descarga.config(state="disabled")
        if cancelada:
            self.progreso_historico.stop()
            self.progreso_historico.config(mode="determinate")
            self.progreso_historico.config(value=0)
            self.lbl_historico_estado.config(text="Descarga cancelada. No se modificaron los registros existentes.", fg=FG_DIM)
            return
        if error:
            self.progreso_historico.stop()
            self.progreso_historico.config(mode="determinate")
            self.progreso_historico.config(value=0)
            self.lbl_historico_estado.config(text="Error de conexión", fg=DANGER)
            return messagebox.showerror("Error de Conexión", f"Fallo al comunicar con el biométrico:\n{error}")
        pd = cargar_pandas()
        self.df_dispositivo = pd.DataFrame(records)
        if self.df_dispositivo.empty:
            self.nombres_completos = []
            self.cb_empleados.config(values=[])
            self.lbl_historico_estado.config(text="Sin registros descargados", fg=FG_DIM)
            return messagebox.showinfo("Sin datos", "No se encontraron registros en el biométrico.")
        if persist:
            append_attendance(
                records,
                device=self.cb_checadores.get(),
                source="consulta_historica",
            )
        self.nombres_completos = sorted(self.df_dispositivo["nombre"].unique())
        self.cb_empleados.config(values=["Todos los empleados"] + self.nombres_completos)
        self.cb_empleados.set("Todos los empleados")
        self.progreso_historico.stop()
        self.progreso_historico.config(mode="determinate")
        self.progreso_historico.config(value=100)
        self.lbl_historico_estado.config(text=f"{len(self.nombres_completos)} empleados cargados", fg=SUCCESS)
        messagebox.showinfo("Extracción Completada", "Base de datos descargada correctamente.\nYa puedes buscar o ver todos los empleados.")
    def filtrar_teclado(self, event):
        tecleado = self.cb_empleados.get()
        if tecleado == "":
            self.cb_empleados.config(values=["Todos los empleados"] + self.nombres_completos)
        else:
            opciones = ["Todos los empleados"] + self.nombres_completos
            filtrados = [n for n in opciones if tecleado.lower() in n.lower()]
            self.cb_empleados.config(values=filtrados)

    def cargar_todo_empleado(self, event=None):
        self.procesar_y_mostrar(aplicar_fechas=self.fechas_historico_aplicadas)

    def filtrar_fechas(self):
        self.fechas_historico_aplicadas = True
        self.procesar_y_mostrar(aplicar_fechas=True)

    def ver_todos_empleados(self):
        if self.df_dispositivo is None or self.df_dispositivo.empty:
            return
        self.cb_empleados.set("Todos los empleados")
        self.fechas_historico_aplicadas = False
        self.procesar_y_mostrar(aplicar_fechas=False)

    def limpiar_tabla(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        for item in self.tree_movimientos.get_children():
            self.tree_movimientos.delete(item)
        pd = cargar_pandas()
        self.df_historico_actual = pd.DataFrame()
        self.lbl_historico_estado.config(text="Tabla limpia", fg=FG_DIM)

    def procesar_y_mostrar(self, aplicar_fechas=False):
        emp_sel = self.cb_empleados.get()
        if self.df_dispositivo is None or self.df_dispositivo.empty:
            return

        for item in self.tree.get_children(): self.tree.delete(item)
        for item in self.tree_movimientos.get_children(): self.tree_movimientos.delete(item)

        if not emp_sel:
            emp_sel = "Todos los empleados"

        if emp_sel == "Todos los empleados":
            df_emp = self.df_dispositivo.copy()
        else:
            df_emp = self.df_dispositivo[self.df_dispositivo["nombre"] == emp_sel].copy()
        df_emp["Fecha_D"] = df_emp["fecha_hora"].dt.date
        df_emp["Hora_D"] = df_emp["fecha_hora"].dt.time
        
        if aplicar_fechas:
            f_ini = self.cal_desde.get_date()
            f_fin = self.cal_hasta.get_date()
            df_emp = df_emp[(df_emp["Fecha_D"] >= f_ini) & (df_emp["Fecha_D"] <= f_fin)]
            if df_emp.empty:
                self.lbl_historico_estado.config(text="Sin resultados en el rango", fg=DANGER)
                return messagebox.showinfo("Sin Resultados", "No se encontraron registros para este empleado en el rango de fechas seleccionado.")

        if self.vista_historico.get() == "movimientos":
            df_movimientos = df_emp.sort_values(
                by=["fecha_hora", "nombre"],
                ascending=[False, True]
            ).reset_index(drop=True)
            self.tree_movimientos.tag_configure("impar", background=BG_CARD)
            self.tree_movimientos.tag_configure("par", background=BG_CARD_ALT)
            for indice, (_, fila) in enumerate(df_movimientos.iterrows()):
                fecha_hora = fila["fecha_hora"]
                tag_color = "impar" if indice % 2 != 0 else "par"
                self.tree_movimientos.insert(
                    "",
                    "end",
                    values=(
                        f"  {fila['nombre']}",
                        fecha_hora.strftime("%d/%m/%Y"),
                        dias_sem[fecha_hora.weekday()],
                        formato_hora(fecha_hora)
                    ),
                    tags=(tag_color,)
                )
            pd = cargar_pandas()
            self.df_historico_actual = pd.DataFrame({
                "Empleado": df_movimientos["nombre"].tolist(),
                "Fecha": df_movimientos["fecha_hora"].apply(lambda f: f.strftime("%d/%m/%Y")).tolist(),
                "Dia": df_movimientos["fecha_hora"].apply(lambda f: dias_sem[f.weekday()]).tolist(),
                "Hora": df_movimientos["fecha_hora"].apply(formato_hora).tolist()
            })
            self.lbl_historico_estado.config(text=f"{len(df_movimientos)} movimientos mostrados", fg=SUCCESS)
            return

        reporte = agrupar_entrada_salida(df_emp, usar_primera_ultima=self.usar_primera_ultima_actual)
        reporte = reporte.sort_values(by=["Fecha_D", "nombre"], ascending=[False, True])
        
        for i, r in reporte.iterrows():
            dia_txt = dias_sem[r["Fecha_D"].weekday()]
            str_ent = formato_entrada_reporte(r["Entrada"])
            str_sal = formato_salida_reporte(r["Entrada"], r["Salida"])
            
            # Colores de cebra para la tabla
            tag_color = 'impar' if i % 2 != 0 else 'par'
            tags = ('incompleto', tag_color) if str_sal == "SIN SALIDA" or str_ent == "SIN ENTRADA" else (tag_color,)
            
            self.tree.insert("", "end", values=(f"  {r['nombre']}", r["Fecha_D"].strftime("%d/%m/%Y"), dia_txt, str_ent, str_sal), tags=tags)
            
        self.tree.tag_configure('impar', background=BG_CARD)
        self.tree.tag_configure('par', background=BG_CARD_ALT)
        self.tree.tag_configure('incompleto', foreground=DANGER, font=("Segoe UI", 10, "bold"))
        pd = cargar_pandas()
        self.df_historico_actual = pd.DataFrame({
            "Empleado": reporte["nombre"],
            "Fecha": reporte["Fecha_D"].apply(lambda f: f.strftime("%d/%m/%Y")),
            "Dia": reporte["Fecha_D"].apply(lambda f: dias_sem[f.weekday()]),
            "Entrada": reporte["Entrada"].apply(formato_entrada_reporte),
            "Salida": reporte.apply(lambda r: formato_salida_reporte(r["Entrada"], r["Salida"]), axis=1)
        })
        self.lbl_historico_estado.config(text=f"{len(reporte)} días mostrados", fg=SUCCESS)


# --- REGISTRO BIOMETRICO DE EMPLEADOS ---

from ui.biometria_window import VentanaBiometria
from services.biometria_service import preparar_jpeg


class VentanaAgregarUsuario(VentanaBiometria):
    """Adapter that keeps the historical menu entry compatible."""

    def __init__(self, parent):
        self.camera_review_dialog = None
        self.camera_review_photo = None
        self.camera_progress_canvas = None
        self.camera_guidance_label = None
        self.camera_face_metric = None
        self.camera_light_metric = None
        self.camera_stability_metric = None
        super().__init__(
            parent,
            cargar_configuracion=cargar_configuracion,
            guardar_configuracion=guardar_configuracion,
            conectar_zk_robusto=conectar_zk_robusto,
            aplicar_icono=aplicar_icono,
        )

    def _capture_camera(self):
        super()._capture_camera()
        dialog = self.camera_dialog
        if not dialog or not dialog.winfo_exists():
            return
        aplicar_icono(dialog)
        if self.camera_progress_canvas and self.camera_progress_canvas.winfo_exists():
            return
        self._enhance_camera_dialog()
        self._update_camera_assistance()

    def _enhance_camera_dialog(self):
        content = self.camera_canvas.master.master
        panels = content.winfo_children()
        if len(panels) < 2:
            return
        tips = panels[1]
        tk.Frame(tips, bg=BORDER_CARD, height=1).pack(fill="x", pady=12)
        tk.Label(
            tips, text="Calidad en tiempo real", bg=tips.cget("bg"), fg=FG_TEXT,
            font=("Segoe UI", 10, "bold"), anchor="w",
        ).pack(fill="x", pady=(0, 7))
        self.camera_guidance_label = tk.Label(
            tips, text="Espera a que aparezca la imagen…", bg=tips.cget("bg"), fg=ACCENT,
            font=("Segoe UI", 9, "bold"), justify="left", anchor="w", wraplength=210,
        )
        self.camera_guidance_label.pack(fill="x", pady=(0, 7))
        self.camera_progress_canvas = tk.Canvas(
            tips, width=210, height=12, bg=BG_CARD_ALT, bd=0, highlightthickness=0,
        )
        self.camera_progress_canvas.pack(fill="x", pady=(0, 8))
        self.camera_face_metric = self._camera_metric_label(tips, "Rostro: esperando")
        self.camera_light_metric = self._camera_metric_label(tips, "Iluminación: esperando")
        self.camera_stability_metric = self._camera_metric_label(tips, "Estabilidad: 0/5")

    @staticmethod
    def _camera_metric_label(parent, text):
        label = tk.Label(
            parent, text=text, bg=parent.cget("bg"), fg=FG_DIM,
            font=("Segoe UI", 8, "bold"), anchor="w",
        )
        label.pack(fill="x", pady=2)
        return label

    def _set_camera_status(self, text, foreground, background):
        super()._set_camera_status(text, foreground, background)
        self._update_camera_assistance()

    def _update_camera_assistance(self):
        widgets = (
            self.camera_progress_canvas, self.camera_guidance_label,
            self.camera_face_metric, self.camera_light_metric, self.camera_stability_metric,
        )
        if not all(widget and widget.winfo_exists() for widget in widgets):
            return
        evaluation = getattr(self, "_camera_last_evaluation", None) or {}
        streak = min(5, max(0, int(getattr(self, "_camera_valid_streak", 0))))
        stable = bool(evaluation.get("valid")) and streak >= 5
        status = evaluation.get("status", "INICIANDO CÁMARA")
        message = evaluation.get("message", "Espera a que aparezca la imagen…")
        face_detected = bool(evaluation.get("face"))
        light_ok = face_detected and status != "POCA ILUMINACIÓN"

        color = SUCCESS if stable else ("#F59E0B" if evaluation.get("level") == "warning" else DANGER)
        if not evaluation:
            color = ACCENT
        guidance = "✓ Encuadre estable. Ya puedes capturar." if stable else message
        self.camera_guidance_label.config(text=guidance, fg=color)
        self.camera_face_metric.config(
            text="Rostro: detectado" if face_detected else "Rostro: no detectado",
            fg=SUCCESS if face_detected else DANGER,
        )
        self.camera_light_metric.config(
            text="Iluminación: adecuada" if light_ok else "Iluminación: revisar",
            fg=SUCCESS if light_ok else "#F59E0B",
        )
        self.camera_stability_metric.config(
            text=f"Estabilidad: {streak}/5" if not stable else "Estabilidad: lista",
            fg=SUCCESS if stable else FG_DIM,
        )
        canvas = self.camera_progress_canvas
        canvas.delete("all")
        width = max(1, canvas.winfo_width())
        canvas.create_rectangle(0, 0, width, 12, fill=BG_CARD_ALT, outline="")
        canvas.create_rectangle(0, 0, width * streak / 5, 12, fill=color, outline="")

    def _accept_camera_frame(self):
        if self.latest_camera_image is None:
            return messagebox.showwarning(
                "Cámara", "Espera a que el encuadre permanezca estable.", parent=self.camera_dialog,
            )
        try:
            data = preparar_jpeg(image=self.latest_camera_image)
        except Exception as exc:
            return messagebox.showerror("Foto no válida", str(exc), parent=self.camera_dialog)
        if self.camera_job:
            try:
                self.camera_dialog.after_cancel(self.camera_job)
            except tk.TclError:
                pass
            self.camera_job = None
        self._show_camera_review(data)

    def _show_camera_review(self, data):
        if self.camera_review_dialog and self.camera_review_dialog.winfo_exists():
            self.camera_review_dialog.destroy()
        dialog = tk.Toplevel(self.camera_dialog)
        self.camera_review_dialog = dialog
        aplicar_icono(dialog)
        dialog.title("Revisar foto facial")
        dialog.configure(bg=BG_APP)
        dialog.resizable(False, False)
        dialog.transient(self.camera_dialog)
        dialog.protocol("WM_DELETE_WINDOW", self._retake_camera_photo)

        card = tk.Frame(dialog, bg=BG_CARD, padx=18, pady=16, highlightthickness=1, highlightbackground=BORDER_CARD)
        card.pack(padx=18, pady=18)
        tk.Label(card, text="Revisa la foto antes de usarla", bg=BG_CARD, fg=FG_TEXT, font=("Segoe UI", 14, "bold")).pack(anchor="w")
        tk.Label(
            card, text="Confirma que el rostro esté completo, centrado y sin sombras intensas.",
            bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 9), wraplength=520, justify="left",
        ).pack(anchor="w", pady=(3, 12))

        from PIL import ImageTk
        preview = data["image"].copy()
        preview.thumbnail((520, 390))
        self.camera_review_photo = ImageTk.PhotoImage(preview)
        tk.Label(card, image=self.camera_review_photo, bg="#050A13", bd=0).pack()
        tk.Label(
            card, text=f"{data['width']}×{data['height']} px · {data['size_kb']} KB",
            bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 8, "bold"),
        ).pack(anchor="w", pady=(8, 0))

        actions = tk.Frame(card, bg=BG_CARD)
        actions.pack(fill="x", pady=(12, 0))
        tk.Button(
            actions, text="Usar esta foto", command=lambda: self._confirm_camera_photo(data),
            bg=SUCCESS, fg="white", font=("Segoe UI", 10, "bold"), bd=0, pady=8, cursor="hand2",
        ).pack(side="left", fill="x", expand=True, padx=(0, 5))
        tk.Button(
            actions, text="Tomar otra", command=self._retake_camera_photo,
            bg=BG_CARD_ALT, fg=FG_TEXT, font=("Segoe UI", 10, "bold"), bd=0, pady=8, cursor="hand2",
        ).pack(side="left", fill="x", expand=True, padx=(5, 0))
        dialog.update_idletasks()
        x = self.camera_dialog.winfo_rootx() + max(0, (self.camera_dialog.winfo_width() - dialog.winfo_width()) // 2)
        y = self.camera_dialog.winfo_rooty() + max(0, (self.camera_dialog.winfo_height() - dialog.winfo_height()) // 2)
        dialog.geometry(f"+{x}+{y}")
        dialog.grab_set()
        dialog.focus_force()

    def _retake_camera_photo(self):
        if self.camera_review_dialog and self.camera_review_dialog.winfo_exists():
            try:
                self.camera_review_dialog.grab_release()
            except tk.TclError:
                pass
            self.camera_review_dialog.destroy()
        self.camera_review_dialog = None
        self.camera_review_photo = None
        if self.camera_dialog and self.camera_dialog.winfo_exists() and self.camera_capture:
            self.latest_camera_image = None
            self._camera_valid_streak = 0
            self._set_camera_capture_enabled(False)
            self.camera_job = self.camera_dialog.after(35, self._camera_tick)
            self.camera_dialog.grab_set()
            self.camera_dialog.focus_force()

    def _confirm_camera_photo(self, data):
        if self.camera_review_dialog and self.camera_review_dialog.winfo_exists():
            try:
                self.camera_review_dialog.grab_release()
            except tk.TclError:
                pass
            self.camera_review_dialog.destroy()
        self.camera_review_dialog = None
        self.camera_review_photo = None
        if self.camera_dialog and self.camera_dialog.winfo_exists():
            self.camera_dialog.configure(bg="#FFFFFF")
            self.camera_dialog.update_idletasks()
            self.camera_dialog.after(70)
        self._close_camera()
        self._set_face(data)

    def _close_camera(self):
        if self.camera_review_dialog and self.camera_review_dialog.winfo_exists():
            try:
                self.camera_review_dialog.grab_release()
            except tk.TclError:
                pass
            self.camera_review_dialog.destroy()
        self.camera_review_dialog = None
        self.camera_review_photo = None
        self.camera_progress_canvas = None
        self.camera_guidance_label = None
        self.camera_face_metric = None
        self.camera_light_metric = None
        self.camera_stability_metric = None
        super()._close_camera()

# --- FUNCIONES DE LOGICA Y PROCESAMIENTO (MAIN) ---

def actualizar_combos_globales():
    conf = cargar_configuracion()
    nombres = [c["nombre"] for c in conf["checadores"]]
    cb_checador['values'] = nombres
    if nombres and checador_var.get() not in nombres:
        cb_checador.current(0)

def actualizar_estilo_periodo():
    seleccionado = modo_reporte.get()
    for rb, frame, label in period_frame_widgets:
        if label.cget("text") == seleccionado:
            frame.config(bg=SUCCESS) 
            label.config(bg=SUCCESS, fg="#FFFFFF")
            rb.config(bg=SUCCESS, activebackground=SUCCESS, selectcolor=SUCCESS)
        else:
            frame.config(bg=BG_CARD)
            label.config(bg=BG_CARD, fg=FG_DIM)
            rb.config(bg=BG_CARD, activebackground=BG_CARD, selectcolor=BG_CARD)

def actualizar_fechas(*args):
    actualizar_estilo_periodo()
    hoy = date.today()
    modo = modo_reporte.get()
    
    if modo == "Diario":
        col_mes_mensual.pack_forget()
        col_desde.pack(side="left", expand=True, fill="x", padx=(0,5))
        cal_desde.set_date(hoy)
        cal_hasta.set_date(hoy)
        col_hasta.pack_forget() 
        lbl_desde.config(text="Día del Reporte:")
    elif modo == "Mensual":
        col_desde.pack_forget()
        col_hasta.pack_forget()
        col_mes_mensual.pack(side="left", expand=True, fill="x")
        actualizar_meses_mensual()
    else:
        col_mes_mensual.pack_forget()
        col_desde.pack(side="left", expand=True, fill="x", padx=(0,5))
        col_hasta.pack(side="right", expand=True, fill="x", padx=(5,0))
        lbl_desde.config(text="Fecha Inicio:")
        if modo == "Semanal":
            lunes = hoy - timedelta(days=hoy.weekday())
            cal_desde.set_date(lunes)
            cal_hasta.set_date(hoy)
    if "root" in globals():
        root.after(50, validar_sincronizacion_periodo)

def actualizar_meses_mensual():
    if "cb_mes_mensual" not in globals():
        return
    valores = opciones_meses_reporte()
    cb_mes_mensual["values"] = valores
    if valores and mes_reporte_var.get() not in valores:
        mes_reporte_var.set(valores[-1])

def obtener_rango_reporte():
    modo = modo_reporte.get()
    if modo == "Mensual":
        actualizar_meses_mensual()
        idx = cb_mes_mensual.current()
        if idx < 0:
            valores = opciones_meses_reporte()
            if not valores:
                return None, None
            idx = len(valores) - 1
            mes_reporte_var.set(valores[idx])
        return rango_mes_actual(idx + 1)
    return cal_desde.get_date(), cal_hasta.get_date()

def obtener_checador_seleccionado():
    checador = checador_var.get()
    conf = cargar_configuracion()
    return next((c for c in conf["checadores"] if c["nombre"] == checador), None)

def rango_actual_coincide_con_memoria():
    if datos_memoria_crudos is None or getattr(datos_memoria_crudos, "empty", True):
        return False
    f_inicio, f_fin = obtener_rango_reporte()
    if not f_inicio or not f_fin:
        return False
    return (
        checador_memoria == checador_var.get()
        and rango_memoria_inicio == f_inicio
        and rango_memoria_fin == f_fin
    )

def texto_rango(fecha_inicio, fecha_fin):
    if not fecha_inicio or not fecha_fin:
        return "sin periodo"
    return f"{fecha_inicio.strftime('%d/%m/%Y')} - {fecha_fin.strftime('%d/%m/%Y')}"

def validar_sincronizacion_periodo(*_args):
    if "btn_procesar" not in globals() or "label_archivo" not in globals():
        return
    if not tiene_permiso(auth_service.PERM_REPORTS_GENERATE):
        set_boton_estado(btn_procesar, "disabled", EXCEL_COL)
        return
    if datos_memoria_crudos is None or getattr(datos_memoria_crudos, "empty", True):
        label_archivo.config(
            text="Pendiente: descarga las asistencias de este período para habilitar Excel.",
            fg=FG_DIM,
        )
        set_boton_estado(btn_procesar, "disabled", EXCEL_COL)
        return
    if rango_actual_coincide_con_memoria():
        base = preparar_base_asistencia(datos_memoria_crudos)
        empleados = len(base["nombre"].unique()) if not base.empty else 0
        rango_txt = texto_rango(rango_memoria_inicio, rango_memoria_fin)
        origen = f"{checador_memoria} | " if checador_memoria else ""
        if not base.empty:
            f_min = base["Fecha_D"].min().strftime("%d/%m/%Y")
            f_max = base["Fecha_D"].max().strftime("%d/%m/%Y")
            label_archivo.config(text=f"{origen}{len(datos_memoria_crudos)} registros | {empleados} empleados | periodo {rango_txt} | checadas {f_min} - {f_max}", fg=SUCCESS)
        else:
            label_archivo.config(text=f"{origen}{len(datos_memoria_crudos)} registros listos | periodo {rango_txt}", fg=SUCCESS)
        set_boton_estado(btn_procesar, "normal", EXCEL_COL)
        return
    seleccionado = texto_rango(*obtener_rango_reporte())
    sincronizado = texto_rango(rango_memoria_inicio, rango_memoria_fin)
    origen = f"{checador_memoria} | " if checador_memoria else ""
    label_archivo.config(
        text=f"{origen}Datos del {sincronizado}. Cambiaste a {seleccionado}; sincroniza de nuevo.",
        fg=DANGER
    )
    set_boton_estado(btn_procesar, "disabled", EXCEL_COL)

def ejecutar_en_ui(func, *args):
    root.after(0, lambda: func(*args))

def bloquear_acciones_conexion(bloquear=True):
    estado = "disabled" if bloquear or not tiene_permiso(auth_service.PERM_ATTENDANCE_DOWNLOAD) else "normal"
    if "btn_sincronizar" in globals():
        set_boton_estado(btn_sincronizar, estado, ACCENT)

def finalizar_descarga_exitosa(records, nombre_checador="", fecha_inicio=None, fecha_fin=None):
    global datos_memoria_crudos, checador_memoria, rango_memoria_inicio, rango_memoria_fin
    bloquear_acciones_conexion(False)
    if not records:
        try:
            record_sync_success(nombre_checador, received=0)
        except OSError:
            pass
        messagebox.showwarning("Aviso", "El checador está vacío.")
        datos_memoria_crudos = None
        checador_memoria = ""
        rango_memoria_inicio = None
        rango_memoria_fin = None
        label_archivo.config(
            text="No se encontraron asistencias para el período seleccionado.", fg=FG_DIM
        )
        set_boton_estado(btn_procesar, "disabled", EXCEL_COL)
        actualizar_progreso(0, "Esperando instrucciones.")
        return

    pd = cargar_pandas()
    datos_memoria_crudos = pd.DataFrame(records)
    checador_memoria = nombre_checador or ""
    resultado_cache = append_attendance(records, device=checador_memoria, source="checador")
    rango_memoria_inicio = fecha_inicio
    rango_memoria_fin = fecha_fin
    base = preparar_base_asistencia(datos_memoria_crudos)
    empleados = len(base["nombre"].unique()) if not base.empty else 0
    rango_txt = texto_rango(rango_memoria_inicio, rango_memoria_fin)
    if not base.empty:
        f_min = base["Fecha_D"].min().strftime("%d/%m/%Y")
        f_max = base["Fecha_D"].max().strftime("%d/%m/%Y")
        origen = f"{checador_memoria} | " if checador_memoria else ""
        label_archivo.config(text=f"{origen}{len(records)} registros | {empleados} empleados | periodo {rango_txt} | checadas {f_min} - {f_max} | caché +{resultado_cache['added']}", fg=SUCCESS)
    else:
        label_archivo.config(text=f"{len(records)} registros listos | periodo {rango_txt} | caché +{resultado_cache['added']}", fg=SUCCESS)

    lbl_status.config(text="Datos descargados", fg=SUCCESS)
    badge_status.config(bg=SUCCESS)
    set_boton_estado(
        btn_procesar,
        "normal" if tiene_permiso(auth_service.PERM_REPORTS_GENERATE) else "disabled",
        EXCEL_COL,
    )
    actualizar_progreso(100, "¡Sincronización exitosa!")

def finalizar_descarga_error(error, nombre_checador=""):
    global datos_memoria_crudos, checador_memoria, rango_memoria_inicio, rango_memoria_fin
    bloquear_acciones_conexion(False)
    try:
        record_sync_error(nombre_checador, error)
    except OSError:
        pass
    datos_memoria_crudos = None
    checador_memoria = ""
    rango_memoria_inicio = None
    rango_memoria_fin = None
    lbl_status.config(text="No se pudo conectar", fg=DANGER)
    badge_status.config(bg=DANGER)
    label_archivo.config(text="No se descargaron datos. Revisa el checador e inténtalo de nuevo.", fg=DANGER)
    set_boton_estado(btn_procesar, "disabled", EXCEL_COL)
    messagebox.showerror("Error", f"Error de Red: {error}")
    actualizar_progreso(0, "Error de conexión.")

def worker_descargar(info, fecha_inicio=None, fecha_fin=None):
    conn = None
    deshabilitado = False
    try:
        if marca_checador(info) == "hikvision":
            ejecutar_en_ui(actualizar_progreso, 10, "Conectando a Hikvision...")
            hikvision_probar_conexion(info)
            records = hikvision_obtener_eventos(
                info,
                fecha_inicio,
                fecha_fin,
                progreso=lambda texto, valor: ejecutar_en_ui(actualizar_progreso, valor, texto)
            )
            ejecutar_en_ui(actualizar_progreso, 90, "Procesando datos...")
            ejecutar_en_ui(finalizar_descarga_exitosa, records, info.get("nombre", ""), fecha_inicio, fecha_fin)
            return
        ejecutar_en_ui(actualizar_progreso, 10, "Conectando al biométrico...")
        conn, modo = conectar_zk_robusto(
            info["ip"],
            info.get("port"),
            password=0,
            modo_preferido=info.get("modo_conexion"),
            progreso=lambda texto, valor: ejecutar_en_ui(actualizar_progreso, valor, texto)
        )
        guardar_preferencia_conexion(info.get("nombre", ""), modo)
        try:
            conn.disable_device()
            deshabilitado = True
        except Exception:
            deshabilitado = False

        ejecutar_en_ui(actualizar_progreso, 40, f"Conectado por {modo}. Descargando usuarios...")
        users = conn.get_users()
        user_dict = {u.user_id: u.name for u in users}

        ejecutar_en_ui(actualizar_progreso, 70, "Descargando asistencias...")
        attendances = conn.get_attendance()

        records = []
        for att in attendances:
            name = user_dict.get(att.user_id, f"Usuario {att.user_id}")
            records.append({"nombre": name, "fecha_hora": att.timestamp})

        ejecutar_en_ui(actualizar_progreso, 90, "Procesando datos...")
        ejecutar_en_ui(finalizar_descarga_exitosa, records, info.get("nombre", ""), fecha_inicio, fecha_fin)
    except Exception as e:
        ejecutar_en_ui(finalizar_descarga_error, str(e), info.get("nombre", ""))
    finally:
        if conn:
            try:
                if deshabilitado:
                    conn.enable_device()
                conn.disconnect()
            except Exception:
                pass

def conectar_y_descargar():
    global datos_memoria_crudos, checador_memoria, rango_memoria_inicio, rango_memoria_fin
    if not requerir_permiso(auth_service.PERM_ATTENDANCE_DOWNLOAD):
        return
    info = obtener_checador_seleccionado()
    if not info:
        return messagebox.showwarning("Atención", "Dispositivo no válido.")

    datos_memoria_crudos = None
    checador_memoria = ""
    rango_memoria_inicio = None
    rango_memoria_fin = None
    lbl_status.config(text="Conectando...", fg="#F59E0B")
    badge_status.config(bg="#F59E0B")
    label_archivo.config(text="Descargando datos del checador...", fg=FG_DIM)
    bloquear_acciones_conexion(True)
    set_boton_estado(btn_procesar, "disabled", EXCEL_COL)
    actualizar_progreso(5, "Preparando conexión...")
    f_inicio, f_fin = obtener_rango_reporte()
    threading.Thread(target=worker_descargar, args=(info, f_inicio, f_fin), daemon=True).start()

def obtener_horario_especial_empleado(nombre_empleado, horarios):
    nombre_normalizado = limpiar_nombre(nombre_empleado).lower()
    for especial in normalizar_horario(horarios).get("especiales", []):
        empleados = {limpiar_nombre(emp).lower() for emp in especial.get("empleados", [])}
        if nombre_normalizado in empleados:
            return especial
    return None

def resumen_horarios_especiales(horarios, empleados_reporte=None):
    empleados_validos = None
    if empleados_reporte is not None:
        empleados_validos = {limpiar_nombre(emp).lower() for emp in empleados_reporte if str(emp).strip()}
    filas = []
    for especial in normalizar_horario(horarios).get("especiales", []):
        empleados = sorted({limpiar_nombre(emp) for emp in especial.get("empleados", []) if str(emp).strip()})
        if empleados_validos is not None:
            empleados = [emp for emp in empleados if limpiar_nombre(emp).lower() in empleados_validos]
        if empleados:
            filas.append({
                "nombre": especial.get("nombre", "Horario especial"),
                "entrada": especial.get("entrada", ""),
                "salida": especial.get("salida", ""),
                "tolerancia": especial.get("tolerancia_min", 0),
                "empleados": empleados
            })
    return filas

def agregar_resumen_horarios_especiales_excel(ws, fila_inicio, start_col, total_cols, horarios, empleados_reporte, thin_border):
    filas = resumen_horarios_especiales(horarios, empleados_reporte)
    if not filas:
        return fila_inicio

    fila = fila_inicio
    end_col = start_col + max(total_cols, 5) - 1
    ws.merge_cells(start_row=fila, start_column=start_col, end_row=fila, end_column=end_col)
    titulo = ws.cell(row=fila, column=start_col, value="Horarios especiales aplicados")
    titulo.font = Font(bold=True, color="FFFFFF")
    titulo.fill = PatternFill("solid", fgColor="1D4ED8")
    titulo.alignment = Alignment(horizontal="center")
    fila += 1

    headers = ["Horario especial", "Entrada", "Salida", "Tolerancia", "Persona"]
    for j, h in enumerate(headers):
        c = ws.cell(row=fila, column=start_col + j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DBEAFE")
        c.border = thin_border
        c.alignment = Alignment(horizontal="center")
    fila += 1

    for item in filas:
        for empleado in item["empleados"]:
            vals = [item["nombre"], item["entrada"], item["salida"], f"{item['tolerancia']} min", empleado]
            for j, v in enumerate(vals):
                c = ws.cell(row=fila, column=start_col + j, value=v)
                c.border = thin_border
                c.fill = PatternFill("solid", fgColor="EFF6FF")
                c.alignment = Alignment(horizontal="center" if j < 4 else "left", wrap_text=True)
            fila += 1
    return fila + 1

def agregar_leyenda_colores_excel(ws, fila_inicio, start_col, thin_border, incluir_especial=True, incluir_no_laboral=True, incluir_no_laborado=True, incluir_sabado_extra=True):
    fila = fila_inicio
    ws.cell(row=fila, column=start_col, value="Leyenda").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=fila, column=start_col).fill = PatternFill("solid", fgColor="334155")
    ws.cell(row=fila, column=start_col).alignment = Alignment(horizontal="center")
    ws.cell(row=fila, column=start_col).border = thin_border
    fila += 1

    items = [
        ("RETARDO", "FEF3C7", "Color amarillo: retardo"),
        ("FALTA / SIN SALIDA", "FFEBEE", "Color rojo: falta o falta de salida"),
    ]
    if incluir_sabado_extra:
        items.append(("SABADO EXTRA", EXCEL_SABADO_EXTRA_FILL, "Sabado no laboral: suma asistencia extra"))
    if incluir_no_laborado:
        items.append(("NO LABORADO", "E5E7EB", "Fecha capturada como no laborada: no castiga"))
    if incluir_no_laboral:
        items.append(("NO LABORAL", "E5E7EB", "Domingo/no laboral: solo informativo, no castiga"))
    if incluir_especial:
        items.append(("HORARIO ESPECIAL", "E0F2FE", "Color azul: persona con horario especial"))

    for etiqueta, color, descripcion in items:
        c1 = ws.cell(row=fila, column=start_col, value=etiqueta)
        c2 = ws.cell(row=fila, column=start_col + 1, value=descripcion)
        c1.fill = PatternFill("solid", fgColor=color)
        c2.fill = PatternFill("solid", fgColor="FFFFFF")
        c1.font = Font(bold=True, color="0F172A")
        c1.alignment = Alignment(horizontal="center")
        c2.alignment = Alignment(horizontal="left")
        c1.border = thin_border
        c2.border = thin_border
        fila += 1
    return fila + 1

def calcular_filas_dashboard_mensual(reporte, fechas_periodo, horarios, col_emp):
    filas = []
    dias_no_laborados = obtener_mapa_dias_no_laborados()
    for emp in sorted(reporte[col_emp].unique()):
        emp_data = reporte[reporte[col_emp] == emp]
        horario_emp = obtener_horario_empleado(emp, horarios)
        especial = obtener_horario_especial_empleado(emp, horarios)
        fechas_laborales = {d for d in fechas_periodo if es_dia_laboral(d, horario_emp, dias_no_laborados)}
        fechas_con_checada_laboral = set()
        fechas_sabado_extra = set()
        fechas_domingo_info = set()
        for _, row in emp_data.iterrows():
            if bool(row.get("Falta", False)):
                continue
            fecha = row["Fecha_D"]
            if es_dia_no_laborado(fecha, dias_no_laborados):
                continue
            if es_dia_laboral(fecha, horario_emp, dias_no_laborados):
                fechas_con_checada_laboral.add(fecha)
            elif fecha.weekday() == 5:
                fechas_sabado_extra.add(fecha)
            elif fecha.weekday() == 6:
                fechas_domingo_info.add(fecha)
        faltas = int(emp_data["Falta"].fillna(False).astype(bool).sum()) if "Falta" in emp_data else 0
        trabajados = len(fechas_con_checada_laboral)
        sabados_extra = len(fechas_sabado_extra)
        domingos_info = len(fechas_domingo_info)
        dias_programados = len(fechas_laborales)
        retardos = int(emp_data["Tiene_Retardo"].fillna(False).astype(bool).sum()) if "Tiene_Retardo" in emp_data else 0
        salidas_temp = int(emp_data["Tiene_Salida_Temprana"].fillna(False).astype(bool).sum()) if "Tiene_Salida_Temprana" in emp_data else 0
        sin_salida = 0
        for _, row in emp_data.iterrows():
            if not bool(row.get("Falta", False)) and row["Fecha_D"] in fechas_laborales and not hora_vacia(row.get("Entrada")) and row.get("Entrada") == row.get("Salida"):
                sin_salida += 1
        asistencia_pct = ((trabajados + sabados_extra) / dias_programados) if dias_programados else 0
        filas.append({
            "empleado": emp,
            "horario": especial.get("nombre", "General") if especial else "General",
            "programados": dias_programados,
            "trabajados": trabajados,
            "sabados_extra": sabados_extra,
            "domingos_info": domingos_info,
            "asistencia_pct": asistencia_pct,
            "retardos": retardos,
            "faltas": faltas,
            "sin_salida": sin_salida,
            "salidas_temp": salidas_temp
        })
    return filas

def crear_dashboard_mensual(wb, reporte, fechas_periodo, horarios, f_inicio, f_fin, thin_border, col_emp):
    ws_dash = wb.create_sheet("Dashboard Mensual", 0)
    ws_dash.sheet_view.showGridLines = False
    ws_dash.protection.set_password(PASS_EXCEL)
    ws_dash.protection.sheet = True

    filas = calcular_filas_dashboard_mensual(reporte, fechas_periodo, horarios, col_emp)
    dias_no_laborados = obtener_mapa_dias_no_laborados()
    dias_laborables_mes = {d for d in fechas_periodo if es_dia_laboral(d, horarios, dias_no_laborados)}
    dias_no_laborados_mes = {d for d in fechas_periodo if d in dias_no_laborados}
    total_programados = sum(f["programados"] for f in filas)
    total_trabajados = sum(f["trabajados"] for f in filas)
    total_sabados_extra = sum(f["sabados_extra"] for f in filas)
    total_domingos_info = sum(f["domingos_info"] for f in filas)
    total_faltas = sum(f["faltas"] for f in filas)
    total_retardos = sum(f["retardos"] for f in filas)
    total_sin_salida = sum(f["sin_salida"] for f in filas)
    asistencia_pct = (sum(f["asistencia_pct"] for f in filas) / len(filas)) if filas else 0

    ws_dash.merge_cells("A1:K2")
    titulo = ws_dash["A1"]
    titulo.value = "DASHBOARD MENSUAL DE ASISTENCIA"
    titulo.font = Font(size=20, bold=True, color="FFFFFF")
    titulo.alignment = Alignment(horizontal="center", vertical="center")
    titulo.fill = PatternFill("solid", fgColor="1F4E78")
    ws_dash.merge_cells("A3:K3")
    subtitulo = ws_dash["A3"]
    subtitulo.value = f"Periodo: {f_inicio.strftime('%d/%m/%Y')} al {f_fin.strftime('%d/%m/%Y')} | Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    subtitulo.font = Font(italic=True, color="475569")
    subtitulo.alignment = Alignment(horizontal="center")

    cards = [
        ("Personas", len(filas), "DBEAFE"),
        ("Asistencia", f"{asistencia_pct:.0%}", "DCFCE7" if asistencia_pct >= 0.9 else "FEF3C7"),
        ("Dias laborables", len(dias_laborables_mes), "E0F2FE"),
        ("No laborados", len(dias_no_laborados_mes), "E5E7EB" if dias_no_laborados_mes else "F1F5F9"),
        ("Faltas", total_faltas, "FFEBEE" if total_faltas else "F1F5F9"),
        ("Retardos", total_retardos, "FEF3C7" if total_retardos else "F1F5F9"),
    ]
    for idx, (label, value, color) in enumerate(cards):
        row = 5 if idx < 3 else 8
        col = 1 + (idx % 3) * 3
        ws_dash.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 2)
        ws_dash.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 2)
        c_label = ws_dash.cell(row=row, column=col, value=label)
        c_value = ws_dash.cell(row=row + 1, column=col, value=value)
        for c in [c_label, c_value]:
            c.fill = PatternFill("solid", fgColor=color)
            c.alignment = Alignment(horizontal="center", vertical="center")
        c_label.font = Font(bold=True, color="334155")
        c_value.font = Font(size=18, bold=True, color="0F172A")

    fila = 12
    headers = ["Empleado", "Horario", "Dias laborables", "Dias trabajados", "Sabados extra", "% asistencia", "Retardos", "Faltas", "Sin salida", "Salidas temp.", "Domingos info"]
    for j, h in enumerate(headers, start=1):
        c = ws_dash.cell(row=fila, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="366092")
        c.border = thin_border
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    fila += 1

    for item in filas:
        vals = [
            item["empleado"], item["horario"], item["programados"], item["trabajados"],
            item["sabados_extra"], item["asistencia_pct"], item["retardos"], item["faltas"],
            item["sin_salida"], item["salidas_temp"], item["domingos_info"]
        ]
        bg = "EFF6FF" if item["horario"] != "General" else ("FFFFFF" if fila % 2 == 0 else "F8FAFC")
        for j, v in enumerate(vals, start=1):
            c = ws_dash.cell(row=fila, column=j, value=v)
            c.border = thin_border
            c.fill = PatternFill("solid", fgColor=bg)
            c.alignment = Alignment(horizontal="center" if j != 1 else "left", wrap_text=True)
            if j == 6:
                c.number_format = "0%"
                if item["asistencia_pct"] > 1:
                    c.fill = PatternFill("solid", fgColor="BBF7D0")
                    c.font = Font(color="166534", bold=True)
                elif item["asistencia_pct"] < 0.8:
                    c.fill = PatternFill("solid", fgColor="FFEBEE")
                    c.font = Font(color="C62828", bold=True)
                elif item["asistencia_pct"] < 0.9:
                    c.fill = PatternFill("solid", fgColor="FEF3C7")
                    c.font = Font(color="B45309", bold=True)
            if j == 5 and int(v or 0) > 0:
                c.font = Font(color="166534", bold=True)
            if j == 11 and int(v or 0) > 0:
                c.font = Font(color="475569", bold=True)
            if j in [7, 8, 9, 10] and int(v or 0) > 0:
                c.font = Font(color="B45309" if j == 7 else "C62828", bold=True)
        fila += 1

    agregar_resumen_horarios_especiales_excel(ws_dash, fila + 2, 1, 11, horarios, reporte[col_emp].unique(), thin_border)
    widths = [32, 18, 15, 15, 13, 13, 11, 11, 12, 14, 13]
    for idx, width in enumerate(widths, start=1):
        ws_dash.column_dimensions[get_column_letter(idx)].width = width
    ws_dash.freeze_panes = "A13"
    ws_dash.page_setup.orientation = ws_dash.ORIENTATION_LANDSCAPE
    ws_dash.page_setup.paperSize = ws_dash.PAPERSIZE_LETTER
    ws_dash.page_setup.fitToWidth = 1
    ws_dash.page_setup.fitToHeight = 0
    return ws_dash

def procesar():
    global ultimo_reporte, titulo_actual, datos_memoria_crudos, checador_memoria
    if not requerir_permiso(auth_service.PERM_REPORTS_GENERATE):
        return
    if datos_memoria_crudos is None or datos_memoria_crudos.empty: 
        messagebox.showwarning("Atención", "No hay datos en memoria. Conecta primero.")
        return
        
    carpeta = filedialog.askdirectory(title="Guardar Excel en...")
    if not carpeta: return
    
    try:
        cargar_openpyxl()
        actualizar_progreso(20, "Aplicando fechas y horarios...")
        f_inicio, f_fin = obtener_rango_reporte()
        if not f_inicio or not f_fin:
            messagebox.showwarning("Mes requerido", "Selecciona el mes del reporte.")
            actualizar_progreso(0, "Listo")
            return
        checador_info = obtener_checador_seleccionado()
        if f_inicio > f_fin:
            messagebox.showwarning("Fechas invalidas", "La fecha inicio no puede ser mayor que la fecha fin.")
            actualizar_progreso(0, "Listo")
            return
        if checador_info and checador_memoria and checador_memoria != checador_info.get("nombre", ""):
            messagebox.showwarning(
                "Sincronizacion requerida",
                f"Los datos cargados son de '{checador_memoria}', pero seleccionaste '{checador_info.get('nombre', '')}'.\n\nSincroniza el checador seleccionado antes de generar el reporte."
            )
            actualizar_progreso(0, "Listo")
            return
        if not rango_actual_coincide_con_memoria():
            messagebox.showwarning(
                "Sincronizacion requerida",
                f"Los datos cargados son del {texto_rango(rango_memoria_inicio, rango_memoria_fin)}, pero el reporte seleccionado es {texto_rango(f_inicio, f_fin)}.\n\nSincroniza este periodo antes de generar el reporte."
            )
            validar_sincronizacion_periodo()
            actualizar_progreso(0, "Listo")
            return
        horarios = obtener_horarios_config(checador_var.get())
        empleados_asignados = checador_info.get("empleados", []) if checador_info else []
        usar_primera_ultima = usa_primera_ultima_checador(checador_info)
        reporte, base_asistencia, fechas_periodo = construir_reporte_asistencia(
            datos_memoria_crudos,
            f_inicio,
            f_fin,
            horarios,
            incluir_faltas=True,
            empleados_asignados=empleados_asignados,
            usar_primera_ultima=usar_primera_ultima
        )

        if reporte.empty:
            messagebox.showinfo("Reporte Vacío", "No hay registros ni faltas laborales en ese rango.")
            actualizar_progreso(0, "Listo")
            return

        col_emp = 'nombre'
        f_min, f_max, modo = f_inicio, f_fin, modo_reporte.get()
        
        if modo == "Diario":
            nombre_fin = f"ASISTENCIA_{f_min.strftime('%d_%m_%Y')}.xlsx"
            titulo_actual = f"ASISTENCIA DEL DÍA: {f_min.strftime('%d/%m/%Y')}"
        else:
            nombre_fin = f"REPORTE_{modo.upper()}_{f_min.day}_AL_{f_max.day}.xlsx"
            titulo_actual = f"REPORTE {modo.upper()}: {f_min.day} AL {f_max.day} DE {meses[f_max.month].upper()}"

        salida = os.path.join(carpeta, nombre_fin)
        if os.path.exists(salida) and not messagebox.askyesno("Confirmar", f"'{nombre_fin}' ya existe. ¿Reemplazar?"): return

        wb = Workbook(); ws = wb.active
        ws.title = "Detalle Mensual" if modo == "Mensual" else "Reporte"
        ws.protection.set_password(PASS_EXCEL)
        ws.protection.sheet = True

        thin_border = Border(left=Side(style='thin', color="D9D9D9"), right=Side(style='thin', color="D9D9D9"), top=Side(style='thin', color="D9D9D9"), bottom=Side(style='thin', color="D9D9D9"))
        unique_dates = fechas_periodo if modo in ["Semanal", "Mensual"] else [f_inicio]
        
        if modo == "Diario": total_cols = 7
        elif modo == "Semanal": total_cols = max(1 + (len(unique_dates) * 2), 5)
        else: total_cols = 10
            
        start_col = 2 

        f_gen = ws.cell(row=1, column=start_col + total_cols - 1, value=f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
        f_gen.font = Font(italic=True, size=9, color="808080"); f_gen.alignment = Alignment(horizontal="right")

        ws.merge_cells(start_row=2, start_column=start_col, end_row=4, end_column=start_col + total_cols - 1)
        t = ws.cell(row=2, column=start_col, value="ASISTENCIA")
        t.font = Font(size=26, bold=True, color="1F4E78"); t.alignment = Alignment(horizontal="center", vertical="center"); t.fill = PatternFill("solid", fgColor="DCE6F1")
        
        ws.merge_cells(start_row=5, start_column=start_col, end_row=5, end_column=start_col + total_cols - 1)
        s = ws.cell(row=5, column=start_col, value=titulo_actual); s.font = Font(bold=True, color="FFFFFF"); s.alignment = Alignment(horizontal="center"); s.fill = PatternFill("solid", fgColor="366092")

        current_row = 7 

        if modo == "Diario":
            headers = ["Nombre Completo", "Horario", "Fecha", "Entrada", "Salida", "Retardo", "Tiempo Retardo"]
            for j, h in enumerate(headers):
                c = ws.cell(row=6, column=start_col + j, value=h)
                c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="F2F2F2"); c.border = thin_border; c.alignment = Alignment(horizontal="center")
            data_d = reporte[reporte["Fecha_D"] == f_min].sort_values(by=col_emp)
            for i, (_, row) in enumerate(data_d.iterrows()):
                current_row = 7 + i
                especial = obtener_horario_especial_empleado(row[col_emp], horarios)
                horario_txt = especial.get("nombre", "General") if especial else "General"
                bg = "E0F2FE" if especial else ("F9F9F9" if i % 2 != 0 else "FFFFFF")
                if row["Falta"]:
                    vals = [row[col_emp], horario_txt, row["Fecha_D"].strftime("%d/%m/%Y"), "FALTA", "FALTA", "FALTA", "FALTA"]
                else:
                    if bool(row.get("Dia_No_Laboral", False)):
                        estado_no_laboral = "NO"
                        tiempo_no_laboral = "0 min"
                        salida_txt = formato_salida_reporte(row["Entrada"], row["Salida"])
                        if salida_txt == "SIN SALIDA":
                            salida_txt = "NO APLICA"
                        vals = [row[col_emp], horario_txt, row["Fecha_D"].strftime("%d/%m/%Y"), formato_entrada_reporte(row["Entrada"]), salida_txt, estado_no_laboral, tiempo_no_laboral]
                    else:
                        salida_txt = formato_salida_reporte(row["Entrada"], row["Salida"])
                        vals = [row[col_emp], horario_txt, row["Fecha_D"].strftime("%d/%m/%Y"), formato_entrada_reporte(row["Entrada"]), salida_txt, row["Retardo_SN"], row["Tiempo_Retardo"]]
                for j, v in enumerate(vals):
                    c = ws.cell(row=current_row, column=start_col + j, value=v)
                    c.border = thin_border; c.alignment = Alignment(horizontal="center"); c.fill = PatternFill("solid", fgColor=bg)
                    if j == 1 and especial: c.font = Font(color="1D4ED8", bold=True)
                    if v == "SIN ENTRADA": c.font = Font(color="C62828", bold=True); c.fill = PatternFill("solid", fgColor="FFEBEE")
                    if v == "SIN SALIDA": c.font = Font(color="C62828", bold=True); c.fill = PatternFill("solid", fgColor="FFEBEE")
                    if v == "FALTA": c.font = Font(color="C62828", bold=True); c.fill = PatternFill("solid", fgColor="FFEBEE")
                    if v == "SI": c.font = Font(color="B45309", bold=True); c.fill = PatternFill("solid", fgColor="FEF3C7")
                    if v == "SABADO EXTRA": c.font = Font(color=EXCEL_SABADO_EXTRA_FONT, bold=True); c.fill = PatternFill("solid", fgColor=EXCEL_SABADO_EXTRA_FILL)
                    if v == "NO LABORADO": c.font = Font(color="475569", bold=True); c.fill = PatternFill("solid", fgColor="E5E7EB")
            current_row += 1
            current_row = agregar_leyenda_colores_excel(ws, current_row + 1, start_col, thin_border, incluir_especial=True, incluir_no_laboral=False, incluir_no_laborado=False, incluir_sabado_extra=False)
            current_row = agregar_resumen_horarios_especiales_excel(ws, current_row, start_col, total_cols, horarios, data_d[col_emp].unique(), thin_border)

        elif modo == "Semanal":
            ws.merge_cells(start_row=6, start_column=start_col, end_row=7, end_column=start_col); ws.cell(row=6, column=start_col, value="Nombre Completo")
            c_idx = start_col + 1
            for d in unique_dates:
                ws.merge_cells(start_row=6, start_column=c_idx, end_row=6, end_column=c_idx+1)
                ws.cell(row=6, column=c_idx, value=f"{dias_sem[d.weekday()]} {d.day}")
                ws.cell(row=7, column=c_idx, value="Entrada"); ws.cell(row=7, column=c_idx+1, value="Salida")
                c_idx += 2
            for r in [6,7]:
                for c in range(start_col, start_col + total_cols):
                    cell = ws.cell(row=r, column=c); cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="F2F2F2"); cell.border = thin_border; cell.alignment = Alignment(horizontal="center")
            
            empleados = sorted(reporte[col_emp].unique())
            current_row = 8
            for i, emp in enumerate(empleados):
                especial = obtener_horario_especial_empleado(emp, horarios)
                bg = "E0F2FE" if especial else ("F9F9F9" if i % 2 != 0 else "FFFFFF")
                nombre_txt = emp
                c_nombre = ws.cell(row=current_row, column=start_col, value=nombre_txt)
                c_nombre.fill = PatternFill("solid", fgColor=bg)
                c_nombre.border = thin_border
                c_nombre.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                if especial:
                    c_nombre.font = Font(color="1D4ED8", bold=True)
                c_idx = start_col + 1
                for d in unique_dates:
                    res = reporte[(reporte[col_emp]==emp) & (reporte["Fecha_D"]==d)]
                    if res.empty:
                        ev, sv = ("", "")
                    elif res.iloc[0]["Falta"]:
                        ev, sv = ("FALTA", "FALTA")
                    elif bool(res.iloc[0].get("Dia_No_Laboral", False)):
                        if bool(res.iloc[0].get("No_Laborado", False)):
                            marca = "NO LABORADO"
                        elif bool(res.iloc[0].get("Extra_Sabado", False)):
                            marca = "SABADO EXTRA"
                        else:
                            marca = ""
                        entrada_txt = formato_entrada_reporte(res.iloc[0]["Entrada"])
                        ev = f"{entrada_txt}\n{marca}" if marca else entrada_txt
                        salida_txt = formato_salida_reporte(res.iloc[0]["Entrada"], res.iloc[0]["Salida"])
                        sv = "" if salida_txt == "SIN SALIDA" else salida_txt
                    else:
                        ev = formato_entrada_reporte(res.iloc[0]["Entrada"])
                        sv = formato_salida_reporte(res.iloc[0]["Entrada"], res.iloc[0]["Salida"])
                        if str(res.iloc[0]["Retardo"]).startswith("Si"):
                            ev = f"{ev}\nRETARDO"
                        if str(res.iloc[0]["Salida_Temprana"]).startswith("Si"):
                            sv = f"{sv}\nSALIDA TEMP."
                    for j, v in enumerate([ev, sv]):
                        c = ws.cell(row=current_row, column=c_idx+j, value=v)
                        c.fill, c.border, c.alignment = PatternFill("solid", fgColor=bg), thin_border, Alignment(horizontal="center", wrap_text=True)
                        if v == "SIN ENTRADA" or (isinstance(v, str) and "SIN ENTRADA" in v): c.font = Font(color="C62828", bold=True); c.fill = PatternFill("solid", fgColor="FFEBEE")
                        if v == "SIN SALIDA": c.font = Font(color="C62828", bold=True); c.fill = PatternFill("solid", fgColor="FFEBEE")
                        if v == "FALTA": c.font = Font(color="C62828", bold=True); c.fill = PatternFill("solid", fgColor="FFEBEE")
                        if isinstance(v, str) and "RETARDO" in v: c.font = Font(color="B45309", bold=True); c.fill = PatternFill("solid", fgColor="FEF3C7")
                        if isinstance(v, str) and "SALIDA TEMP." in v: c.font = Font(color="C2410C", bold=True); c.fill = PatternFill("solid", fgColor="FFEDD5")
                        if isinstance(v, str) and "SABADO EXTRA" in v: c.font = Font(color=EXCEL_SABADO_EXTRA_FONT, bold=True); c.fill = PatternFill("solid", fgColor=EXCEL_SABADO_EXTRA_FILL)
                        if isinstance(v, str) and "NO LABORADO" in v: c.font = Font(color="475569", bold=True); c.fill = PatternFill("solid", fgColor="E5E7EB")
                    c_idx += 2
                current_row += 1
            current_row = agregar_leyenda_colores_excel(ws, current_row + 1, start_col, thin_border, incluir_especial=True, incluir_no_laboral=False)
            current_row = agregar_resumen_horarios_especiales_excel(ws, current_row, start_col, total_cols, horarios, empleados, thin_border)

        elif modo == "Mensual":
            headers = ["Nombre Completo", "Horario", "Semana", "Lunes", "Martes", "Miercoles", "Jueves", "Viernes", "Sabado", "Domingo"]
            for j, h in enumerate(headers):
                c = ws.cell(row=6, column=start_col + j, value=h)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="366092")
                c.border = thin_border
                c.alignment = Alignment(horizontal="center")
            ws.row_dimensions[6].height = 26
            dias_no_laborados = obtener_mapa_dias_no_laborados()

            semanas = []
            if unique_dates:
                curr_week = []
                curr_monday = unique_dates[0] - timedelta(days=unique_dates[0].weekday())
                for d in unique_dates:
                    monday = d - timedelta(days=d.weekday())
                    if monday == curr_monday:
                        curr_week.append(d)
                    else:
                        semanas.append((curr_monday, curr_week))
                        curr_week = [d]
                        curr_monday = monday
                if curr_week:
                    semanas.append((curr_monday, curr_week))

            empleados = sorted(reporte[col_emp].unique())
            current_row = 7
            for emp in empleados:
                especial = obtener_horario_especial_empleado(emp, horarios)
                horario_txt = especial.get("nombre", "General") if especial else "General"
                emp_data = reporte[reporte[col_emp] == emp].sort_values(by="Fecha_D")
                emp_por_fecha = {row["Fecha_D"]: row for _, row in emp_data.iterrows()}
                start_row_emp = current_row
                primera_semana_emp = True
                for num_sem, (monday, dates_in_week) in enumerate(semanas, start=1):
                    if not any(d in emp_por_fecha or d in dias_no_laborados for d in dates_in_week):
                        continue
                    start_str = dates_in_week[0].strftime("%d/%m")
                    end_str = dates_in_week[-1].strftime("%d/%m")
                    vals_base = [emp if primera_semana_emp else "", horario_txt if primera_semana_emp else "", f"Sem {num_sem}\n{start_str}-{end_str}"]
                    bg_base = "E0F2FE" if especial else ("FFFFFF" if current_row % 2 == 0 else "F8FAFC")
                    for j, v in enumerate(vals_base):
                        c = ws.cell(row=current_row, column=start_col + j, value=v)
                        c.border = thin_border
                        c.fill = PatternFill("solid", fgColor=bg_base)
                        c.alignment = Alignment(horizontal="left" if j == 0 else "center", vertical="center", wrap_text=True)
                        if j == 1 and especial:
                            c.font = Font(color="1D4ED8", bold=True)

                    for dia_idx in range(7):
                        d = monday + timedelta(days=dia_idx)
                        c = ws.cell(row=current_row, column=start_col + 3 + dia_idx)
                        c.border = thin_border
                        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                        c.fill = PatternFill("solid", fgColor=bg_base)
                        if d not in unique_dates:
                            c.value = ""
                            continue
                        row = emp_por_fecha.get(d)
                        if row is None:
                            if d in dias_no_laborados:
                                motivo = dias_no_laborados.get(d, "")
                                c.value = f"{d.day}\nNO LABORADO" + (f"\n{motivo}" if motivo else "")
                                c.font = Font(color="475569", bold=True)
                                c.fill = PatternFill("solid", fgColor="E5E7EB")
                            else:
                                c.value = ""
                            continue
                        if row.get("Falta", False):
                            c.value = f"{d.day}\nFALTA"
                            c.font = Font(color="C62828", bold=True)
                            c.fill = PatternFill("solid", fgColor="FFEBEE")
                        elif bool(row.get("Dia_No_Laboral", False)):
                            entrada_txt = formato_hora(row["Entrada"])
                            salida_txt = "" if row["Entrada"] == row["Salida"] else formato_hora(row["Salida"])
                            horas_txt = entrada_txt if not salida_txt else f"{entrada_txt} - {salida_txt}"
                            if bool(row.get("Extra_Sabado", False)):
                                c.value = f"{d.day}\nSABADO EXTRA\n{horas_txt}"
                                c.font = Font(color=EXCEL_SABADO_EXTRA_FONT, bold=True)
                                c.fill = PatternFill("solid", fgColor=EXCEL_SABADO_EXTRA_FILL)
                            elif bool(row.get("No_Laborado", False)):
                                motivo = row.get("Motivo_No_Laborado", "")
                                detalle = motivo if motivo else horas_txt
                                c.value = f"{d.day}\nNO LABORADO\n{detalle}"
                                c.font = Font(color="475569", bold=True)
                                c.fill = PatternFill("solid", fgColor="E5E7EB")
                            else:
                                c.value = f"{d.day}\nNO LABORAL\n{horas_txt}"
                                c.font = Font(color="475569", bold=True)
                                c.fill = PatternFill("solid", fgColor="E5E7EB")
                        else:
                            entrada_txt = formato_entrada_reporte(row["Entrada"])
                            salida_txt = formato_salida_reporte(row["Entrada"], row["Salida"])
                            marcas = []
                            if entrada_txt == "SIN ENTRADA":
                                marcas.append("SIN ENTRADA")
                            if salida_txt == "SIN SALIDA":
                                marcas.append("SIN SALIDA")
                            if bool(row.get("Tiene_Retardo", False)):
                                marcas.append("RETARDO")
                            if bool(row.get("Tiene_Salida_Temprana", False)):
                                marcas.append("SALIDA TEMP.")
                            estado_txt = " / ".join(marcas) if marcas else "ASISTENCIA"
                            c.value = f"{d.day}\n{estado_txt}\n{entrada_txt} - {salida_txt}"
                            if "SIN ENTRADA" in estado_txt or "SIN SALIDA" in estado_txt:
                                c.font = Font(color="C62828", bold=True)
                                c.fill = PatternFill("solid", fgColor="FFEBEE")
                            elif "RETARDO" in estado_txt:
                                c.font = Font(color="B45309", bold=True)
                                c.fill = PatternFill("solid", fgColor="FEF3C7")
                            elif "SALIDA TEMP." in estado_txt:
                                c.font = Font(color="C2410C", bold=True)
                                c.fill = PatternFill("solid", fgColor="FFEDD5")
                    ws.row_dimensions[current_row].height = 55
                    current_row += 1
                    primera_semana_emp = False
                if current_row > start_row_emp:
                    end_row_emp = current_row - 1
                    if end_row_emp > start_row_emp:
                        ws.merge_cells(start_row=start_row_emp, start_column=start_col, end_row=end_row_emp, end_column=start_col)
                        ws.merge_cells(start_row=start_row_emp, start_column=start_col + 1, end_row=end_row_emp, end_column=start_col + 1)
                    ws.cell(row=start_row_emp, column=start_col).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    ws.cell(row=start_row_emp, column=start_col + 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            current_row = agregar_leyenda_colores_excel(ws, current_row + 1, start_col, thin_border, incluir_especial=True)
            current_row = agregar_resumen_horarios_especiales_excel(ws, current_row, start_col, total_cols, horarios, empleados, thin_border)
            crear_dashboard_mensual(wb, reporte, unique_dates, horarios, f_inicio, f_fin, thin_border, col_emp)

        ws.freeze_panes = "A7"
        actualizar_progreso(90, "Dando formato...")
        for col in range(start_col, start_col + total_cols):
            max_length = 0
            column_letter = get_column_letter(col)
            for row in range(6, current_row):
                cell_value = ws[f"{column_letter}{row}"].value
                if cell_value:
                    try:
                        if len(str(cell_value)) > max_length: max_length = len(str(cell_value))
                    except: pass
            offset = col - start_col
            if modo == "Diario":
                anchos = [32, 18, 12, 13, 13, 12, 16]
                ws.column_dimensions[column_letter].width = anchos[offset] if offset < len(anchos) else 14
            elif modo == "Mensual":
                anchos = [30, 18, 12, 16, 16, 16, 16, 16, 16, 16]
                ws.column_dimensions[column_letter].width = anchos[offset] if offset < len(anchos) else 14
            elif col == start_col:
                ws.column_dimensions[column_letter].width = min(max_length + 2, 32)
            else:
                ws.column_dimensions[column_letter].width = 14

        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0 
        ws.sheet_view.showGridLines = False
        if modo == "Mensual":
            wb.active = 0
        
        guardado_ok = False
        while not guardado_ok:
            try:
                wb.save(salida)
                guardado_ok = True
            except PermissionError:
                if not messagebox.askretrycancel("Archivo Abierto", "Excel abierto. Ciérrelo y reintente."): return
        
        ultimo_reporte = os.path.abspath(salida)
        actualizar_progreso(100, "¡Reporte Generado! ✔")
        root.after(1200, lambda: os.startfile(salida))

    except Exception as e: messagebox.showerror("Error", f"Error: {e}")

def actualizar_progreso(valor, texto):
    barra['value'] = valor
    label_estado.config(text=texto); root.update()

def salir_aplicacion():
    if messagebox.askyesno("Salir", "¿Deseas cerrar el sistema de asistencia?"):
        if bloqueo_instancia is not None:
            bloqueo_instancia.release()
        root.destroy()


def cerrar_sesion():
    if messagebox.askyesno("Cerrar sesión", "Se cerrará la sesión actual y la aplicación."):
        if bloqueo_instancia is not None:
            bloqueo_instancia.release()
        root.destroy()


def cambiar_contrasena_actual():
    if usuario_actual and PasswordChangeWindow.change(root, usuario_actual["username"]):
        messagebox.showinfo("Contraseña actualizada", "Tu contraseña se actualizó correctamente.", parent=root)

# --- INTERFAZ GRÁFICA MAIN ---

configurar_app_windows()
root = tk.Tk()
aplicar_icono(root)
root.withdraw()

try:
    bloqueo_instancia = InstanceLock(APP_DIR).acquire()
except InstanceLockedError as exc:
    messagebox.showerror("Sistema en uso", f"{exc}\n\nCierra la sesión abierta antes de intentar acceder de nuevo.", parent=root)
    root.destroy()
    raise SystemExit(1)

usuario_actual = LoginWindow.ask(root)
if usuario_actual is None:
    bloqueo_instancia.release()
    root.destroy()
    raise SystemExit(0)

root.title(f"Sistema Asistencia {VERSION}")
root.geometry("500x760")
root.configure(bg=BG_APP)
root.resizable(False, False)
root.protocol("WM_DELETE_WINDOW", salir_aplicacion)
root.deiconify()

# --- MENÚ NATIVO ---
barra_menu = tk.Menu(root)
root.config(menu=barra_menu)

menu_config = tk.Menu(barra_menu, tearoff=0)
hay_configuracion = False
if tiene_permiso(auth_service.PERM_DEVICES_MANAGE):
    menu_config.add_command(label="Gestionar dispositivos", command=lambda: ejecutar_accion_protegida(auth_service.PERM_DEVICES_MANAGE, lambda: VentanaConfiguracion(root)))
    hay_configuracion = True
if tiene_permiso(auth_service.PERM_SCHEDULES_MANAGE):
    menu_config.add_command(label="Horarios por checador", command=lambda: ejecutar_accion_protegida(auth_service.PERM_SCHEDULES_MANAGE, lambda: VentanaHorarios(root)))
    hay_configuracion = True
if tiene_permiso(auth_service.PERM_NON_WORKING_DAYS_MANAGE):
    menu_config.add_command(label="Días no laborados", command=lambda: ejecutar_accion_protegida(auth_service.PERM_NON_WORKING_DAYS_MANAGE, lambda: VentanaDiasNoLaborados(root)))
    hay_configuracion = True
if tiene_permiso(auth_service.PERM_DEVICE_HEALTH_VIEW):
    if hay_configuracion:
        menu_config.add_separator()
    menu_config.add_command(label="Estado de checadores...", command=lambda: ejecutar_accion_protegida(auth_service.PERM_DEVICE_HEALTH_VIEW, lambda: VentanaEstadoChecadores(root)))
    hay_configuracion = True
if tiene_permiso(auth_service.PERM_BACKUPS_CREATE):
    menu_config.add_command(label="Crear respaldo completo...", command=lambda: ejecutar_accion_protegida(auth_service.PERM_BACKUPS_CREATE, lambda: create_backup_dialog(root)))
    hay_configuracion = True
if tiene_permiso(auth_service.PERM_BACKUPS_RESTORE):
    menu_config.add_command(label="Restaurar respaldo...", command=lambda: ejecutar_accion_protegida(auth_service.PERM_BACKUPS_RESTORE, lambda: restore_backup_dialog(root, actualizar_combos_globales)))
    hay_configuracion = True
if tiene_permiso(auth_service.PERM_USERS_MANAGE):
    if hay_configuracion:
        menu_config.add_separator()
    menu_config.add_command(label="Usuarios y permisos...", command=lambda: ejecutar_accion_protegida(auth_service.PERM_USERS_MANAGE, lambda: UserManagementWindow(root)))
    hay_configuracion = True
if hay_configuracion:
    barra_menu.add_cascade(label="Configuración", menu=menu_config)

menu_empleados = tk.Menu(barra_menu, tearoff=0)
hay_empleados = False
if tiene_permiso(auth_service.PERM_BIOMETRICS_MANAGE):
    menu_empleados.add_command(label="Registro biométrico", command=lambda: ejecutar_accion_protegida(auth_service.PERM_BIOMETRICS_MANAGE, lambda: VentanaAgregarUsuario(root)))
    hay_empleados = True
if tiene_permiso(auth_service.PERM_PEOPLE_MANAGE):
    if hay_empleados:
        menu_empleados.add_separator()
    menu_empleados.add_command(label="Asignar personas a checadores", command=lambda: ejecutar_accion_protegida(auth_service.PERM_PEOPLE_MANAGE, lambda: VentanaPersonasChecador(root)))
    hay_empleados = True
if tiene_permiso(auth_service.PERM_HISTORY_VIEW):
    menu_empleados.add_command(label="Consultar historial", command=lambda: ejecutar_accion_protegida(auth_service.PERM_HISTORY_VIEW, lambda: VentanaHistorico(root)))
    hay_empleados = True
if hay_empleados:
    barra_menu.add_cascade(label="Empleados", menu=menu_empleados)

menu_sesion = tk.Menu(barra_menu, tearoff=0)
barra_menu.add_cascade(label="Sesión", menu=menu_sesion)
menu_sesion.add_command(label="Cambiar mi contraseña...", command=cambiar_contrasena_actual)
menu_sesion.add_command(label="Cerrar sesión", command=cerrar_sesion)

modo_reporte = tk.StringVar(value="Diario")
checador_var = tk.StringVar()
mes_reporte_var = tk.StringVar()
modo_reporte.trace_add("write", actualizar_fechas)
mes_reporte_var.trace_add("write", lambda *_args: root.after(50, validar_sincronizacion_periodo))

style = Style()
style.theme_use('clam')
style.configure("TProgressbar", thickness=6, troughcolor=BG_CARD, background=SUCCESS, borderwidth=0)
style.configure("TCombobox", fieldbackground=BG_CARD, background=BG_CARD, foreground=FG_TEXT, darkcolor=BG_CARD, lightcolor=BG_CARD, bordercolor=BORDER_CARD, arrowcolor=FG_TEXT)
style.map("TCombobox", fieldbackground=[('readonly', BG_CARD)], selectbackground=[('readonly', ACCENT)], selectforeground=[('readonly', FG_TEXT)], foreground=[('readonly', FG_TEXT)])
root.option_add('*TCombobox*Listbox.background', BG_CARD)
root.option_add('*TCombobox*Listbox.foreground', FG_TEXT)
root.option_add('*TCombobox*Listbox.selectBackground', ACCENT)
root.option_add('*TCombobox*Listbox.selectForeground', FG_TEXT)

# --- INICIO: FLUJO GUIADO DE OPERACIÓN ---
header = tk.Frame(root, bg=BG_APP, pady=15)
header.pack(fill="x")
tk.Label(header, text="SISTEMA ASISTENCIA", bg=BG_APP, fg=FG_TEXT, font=("Segoe UI", 16, "bold")).pack()
tk.Label(
    header,
    text="Descarga asistencias y genera reportes Excel",
    bg=BG_APP,
    fg=ACCENT,
    font=("Segoe UI", 10),
).pack()
label_usuario_actual = tk.Label(header, bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 8))
label_usuario_actual.pack(pady=(5, 0))
actualizar_etiqueta_usuario()

main_container = tk.Frame(root, bg=BG_APP, padx=25)
main_container.pack(fill="both", expand=True)

# --- PASO 1: CHECADOR ---
tk.Label(main_container, text="1. SELECCIONA EL CHECADOR", bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(0,5))
card_checador = tk.Frame(main_container, bg=BG_CARD, pady=13, padx=20, highlightthickness=1, highlightbackground=BORDER_CARD)
card_checador.pack(fill="x", pady=(0, 16))

row_disp = tk.Frame(card_checador, bg=BG_CARD)
row_disp.pack(fill="x")

tk.Label(row_disp, text="Checador:", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 10)).pack(side="left")
cb_checador = ttk.Combobox(row_disp, textvariable=checador_var, width=19, state="readonly", font=("Segoe UI", 10))
cb_checador.pack(side="left", padx=10, fill="x", expand=True)
cb_checador.bind("<<ComboboxSelected>>", lambda _e: root.after(50, validar_sincronizacion_periodo))

status_frame = tk.Frame(card_checador, bg=BG_CARD)
status_frame.pack(anchor="w", pady=(10, 0))
badge_status = tk.Frame(status_frame, bg=ACCENT, width=10, height=10, relief="flat")
badge_status.pack(side="left", padx=(0,5))
lbl_status = tk.Label(status_frame, text="Listo para consultar", bg=BG_CARD, fg=ACCENT, font=("Segoe UI", 9, "bold"))
lbl_status.pack(side="left")

# --- PASO 2: PERÍODO ---
tk.Label(main_container, text="2. ELIGE EL PERÍODO", bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(0,5))
card_filtros = tk.Frame(main_container, bg=BG_CARD, pady=15, padx=20, highlightthickness=1, highlightbackground=BORDER_CARD)
card_filtros.pack(fill="x", pady=(0, 16))

p_container = tk.Frame(card_filtros, bg=BG_CARD)
p_container.pack(fill="x", pady=(0, 15))

for m in ["Diario", "Semanal", "Mensual"]:
    rb_f = tk.Frame(p_container, bg=BG_CARD)
    rb_f.pack(side="left", expand=True, fill="both", padx=2)
    rb = tk.Radiobutton(rb_f, text=m, variable=modo_reporte, value=m, bg=BG_CARD, fg=FG_DIM, indicatoron=0, bd=0, padx=5, pady=6, font=("Segoe UI", 10, "bold"), cursor="hand2")
    rb.pack(fill="both")
    period_frame_widgets.append((rb, rb_f, rb))

row_fechas = tk.Frame(card_filtros, bg=BG_CARD)
row_fechas.pack(fill="x")

cal_style = {
    'background': BG_CARD, 'foreground': FG_TEXT, 'headersbackground': BG_APP, 'headersforeground': FG_DIM,        
    'normalbackground': BG_CARD, 'normalforeground': FG_TEXT, 'weekendbackground': BG_CARD, 'weekendforeground': FG_DIM,        
    'othermonthbackground': BG_APP, 'othermonthwebackground': BG_APP, 'othermonthforeground': "#334155",  
    'othermonthweforeground': "#334155", 'selectbackground': SUCCESS, 'selectforeground': FG_TEXT,        
    'bottombackground': BG_APP, 'bordercolor': BORDER_CARD, 'borderwidth': 0, 'font': ("Segoe UI", 10)
}

col_desde = tk.Frame(row_fechas, bg=BG_CARD)
col_desde.pack(side="left", expand=True, fill="x", padx=(0,5))
lbl_desde = tk.Label(col_desde, text="Fecha Inicio:", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 9))
lbl_desde.pack(anchor="w", pady=(0,2))
cal_desde = DateEntry(col_desde, locale='es_MX', date_pattern='dd/MM/yyyy', **cal_style)
cal_desde.pack(fill="x")
cal_desde.bind("<<DateEntrySelected>>", lambda _e: root.after(50, validar_sincronizacion_periodo))

col_hasta = tk.Frame(row_fechas, bg=BG_CARD)
col_hasta.pack(side="right", expand=True, fill="x", padx=(5,0))
tk.Label(col_hasta, text="Fecha Fin:", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 9)).pack(anchor="w", pady=(0,2))
cal_hasta = DateEntry(col_hasta, locale='es_MX', date_pattern='dd/MM/yyyy', **cal_style)
cal_hasta.pack(fill="x")
cal_hasta.bind("<<DateEntrySelected>>", lambda _e: root.after(50, validar_sincronizacion_periodo))

col_mes_mensual = tk.Frame(row_fechas, bg=BG_CARD)
tk.Label(col_mes_mensual, text="Mes del Reporte:", bg=BG_CARD, fg=FG_DIM, font=("Segoe UI", 9)).pack(anchor="w", pady=(0,2))
cb_mes_mensual = ttk.Combobox(col_mes_mensual, textvariable=mes_reporte_var, values=opciones_meses_reporte(), state="readonly", font=("Segoe UI", 10))
cb_mes_mensual.pack(fill="x", ipady=2)
cb_mes_mensual.bind("<<ComboboxSelected>>", lambda _e: root.after(50, validar_sincronizacion_periodo))
actualizar_meses_mensual()

# --- PASO 3: DESCARGA ---
tk.Label(main_container, text="3. DESCARGA ASISTENCIAS", bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(0,5))
card_descarga = tk.Frame(main_container, bg=BG_CARD, pady=13, padx=20, highlightthickness=1, highlightbackground=BORDER_CARD)
card_descarga.pack(fill="x", pady=(0, 16))

btn_sincronizar = tk.Button(
    card_descarga,
    text="Descargar asistencias del período",
    command=conectar_y_descargar,
    bg=ACCENT,
    fg=FG_TEXT,
    font=("Segoe UI", 11, "bold"),
    bd=0,
    pady=8,
    cursor="hand2",
)
btn_sincronizar.pack(fill="x")
if not tiene_permiso(auth_service.PERM_ATTENDANCE_DOWNLOAD):
    set_boton_estado(btn_sincronizar, "disabled", ACCENT)

label_archivo = tk.Label(
    card_descarga,
    text="Pendiente: descarga las asistencias de este período para habilitar Excel.",
    bg=BG_CARD,
    fg=FG_DIM,
    font=("Segoe UI", 8),
    wraplength=430,
    justify="center",
)
label_archivo.pack(pady=(7, 0))

actualizar_combos_globales()
actualizar_fechas()

# --- PASO 4: EXCEL ---
tk.Label(main_container, text="4. GENERA EL REPORTE", bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 9, "bold")).pack(anchor="w", pady=(0,5))
row_acciones = tk.Frame(main_container, bg=BG_APP)
row_acciones.pack(fill="x")

btn_procesar = tk.Button(row_acciones, text=" Generar Excel", command=procesar, font=("Segoe UI", 11, "bold"), bd=0, pady=10)
btn_procesar.pack(side="left", expand=True, fill="x", padx=(0, 5))
set_boton_estado(btn_procesar, "disabled", EXCEL_COL) 

btn_salir = tk.Button(row_acciones, text=" Salir", command=salir_aplicacion, font=("Segoe UI", 11, "bold"), bd=0, pady=10)
btn_salir.pack(side="right", expand=True, fill="x", padx=(5, 0))
set_boton_estado(btn_salir, "normal", DANGER) 

# --- FOOTER ANCLADO A ROOT ---
footer = tk.Frame(root, bg=BG_APP, pady=15)
footer.pack(fill="x", side="bottom")

barra = Progressbar(footer, style="TProgressbar", mode='determinate')
barra.pack(fill="x", padx=25, pady=(0, 5)) 

label_estado = tk.Label(footer, text="Esperando instrucciones", bg=BG_APP, fg=FG_DIM, font=("Segoe UI", 9))
label_estado.pack()

tk.Label(footer, text=f"Departamento de Sistemas | Asistencia Agropecuaria {VERSION}", bg=BG_APP, fg="#475569", font=("Segoe UI", 8)).pack(pady=(10, 0))

root.mainloop()
